"""单条组价路由.

当前阶段完成“套定额闭环”：任务入库、单条运行入库、AI 结果候选校验、
人工确认/拒绝、人工对比工程评测。不在本阶段计算综合单价。
"""

from __future__ import annotations

import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from io import BytesIO
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from threading import BoundedSemaphore, Lock
from time import perf_counter, sleep
from typing import Any, Iterable, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openai import OpenAI
from pydantic import BaseModel, Field
from psycopg2.extras import Json

from db.pricing_kb_versions import apply_version_schema, resolve_version_id
from db.schema_lock import acquire_schema_transaction_lock
from api.auth import CurrentUser, current_user, require_project_owner, require_task_owner
from db.connection import DatabasePoolBusyError, suspend_connections_for_model_call
from api.services.model_profiles import active_profile, bind_default_profile_iterator, client_for, use_default_profile
from api.services.ai_usage import use_usage_context

router = APIRouter()
_SCHEMA_LOCK = Lock()
_SCHEMA_READY = False
_MODEL_GATE_LOCK = Lock()
_MODEL_GATE: BoundedSemaphore | None = None
_MODEL_GATE_LIMIT = 0
_MODEL_GATE_RUNNING = 0
_MODEL_GATE_WAITING = 0


def _json_dumps(data: Any) -> str:
    return json.dumps(data, ensure_ascii=False, default=str)


def _parse_json_content(content: str | None) -> dict[str, Any]:
    text = (content or "").strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[1] if "\n" in text else ""
        text = text.rsplit("```", 1)[0].strip()
    parsed = _parse_json_object(text)
    if not isinstance(parsed, dict):
        raise ValueError("JSON result must be an object")
    return parsed


def _parse_json_object(text: str | None) -> dict[str, Any]:
    value = (text or "").strip()
    if not value:
        raise ValueError("empty JSON content")
    decoder = json.JSONDecoder()
    parsed, _ = decoder.raw_decode(value)
    if not isinstance(parsed, dict):
        raise ValueError("JSON result must be an object")
    return parsed


def _sse(data: dict[str, Any]) -> str:
    return f"data: {_json_dumps(data)}\n\n"


def _base_code(item_code: str | None) -> str:
    code = (item_code or "").strip().replace(" ", "")
    return code[:-3] if len(code) > 9 else code


def _apply_schema(conn):
    """Idempotent schema for the pricing-task workflow."""
    with conn.cursor() as cur:
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS pricing_tasks (
                id                  SERIAL PRIMARY KEY,
                name                TEXT NOT NULL,
                boq_project_id      INTEGER NOT NULL REFERENCES boq_projects(id) ON DELETE CASCADE,
                quota_library_ids   JSONB NOT NULL DEFAULT '[]'::jsonb,
                manual_project_id   INTEGER REFERENCES manual_boq_projects(id) ON DELETE SET NULL,
                kb_version_id       BIGINT REFERENCES pricing_kb_versions(id),
                owner_user_id       INTEGER REFERENCES users(id),
                legacy_local_id     TEXT UNIQUE,
                accuracy_report     JSONB,
                status              VARCHAR(16) NOT NULL DEFAULT 'active',
                created_at          TIMESTAMP DEFAULT NOW(),
                updated_at          TIMESTAMP DEFAULT NOW()
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS pricing_task_runs (
                id                  SERIAL PRIMARY KEY,
                task_id             INTEGER REFERENCES pricing_tasks(id) ON DELETE CASCADE,
                boq_item_id         INTEGER NOT NULL REFERENCES boq_items(id) ON DELETE CASCADE,
                boq_project_id      INTEGER NOT NULL REFERENCES boq_projects(id) ON DELETE CASCADE,
                kb_version_id       BIGINT REFERENCES pricing_kb_versions(id),
                status              VARCHAR(16) NOT NULL DEFAULT 'running',
                reasoning_text      TEXT,
                code_check          JSONB,
                feature_check       JSONB,
                work_procedures     JSONB,
                quota_candidates    JSONB,
                quota_match         JSONB,
                evaluation          JSONB,
                conversion_check    JSONB,
                coefficient_check   JSONB,
                accuracy_report     JSONB,
                step_timings        JSONB,
                error_message       TEXT,
                created_at          TIMESTAMP DEFAULT NOW(),
                finished_at         TIMESTAMP
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS pricing_task_results (
                id              SERIAL PRIMARY KEY,
                boq_item_id     INTEGER NOT NULL,
                boq_project_id  INTEGER NOT NULL,
                dezmid          BIGINT NOT NULL,
                dekid           BIGINT NOT NULL,
                subitem_code    VARCHAR(64),
                subitem_name    TEXT,
                qty_factor      NUMERIC(20, 6) DEFAULT 1,
                work_procedure  TEXT,
                confidence      VARCHAR(16),
                missing_info    TEXT,
                ai_reasoning    TEXT,
                status          VARCHAR(16) DEFAULT 'pending',
                created_at      TIMESTAMP DEFAULT NOW(),
                updated_at      TIMESTAMP DEFAULT NOW()
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS pricing_task_batches (
                id                  SERIAL PRIMARY KEY,
                name                TEXT NOT NULL,
                boq_project_id      INTEGER NOT NULL REFERENCES boq_projects(id) ON DELETE CASCADE,
                quota_library_ids   JSONB NOT NULL DEFAULT '[]'::jsonb,
                manual_project_id   INTEGER REFERENCES manual_boq_projects(id) ON DELETE SET NULL,
                kb_version_id       BIGINT REFERENCES pricing_kb_versions(id),
                workflow_version     VARCHAR(16) NOT NULL DEFAULT 'legacy',
                concurrency_limit    INTEGER NOT NULL DEFAULT 20,
                owner_user_id       INTEGER REFERENCES users(id),
                status              VARCHAR(16) NOT NULL DEFAULT 'active',
                selected_count      INTEGER NOT NULL DEFAULT 0,
                completed_count     INTEGER NOT NULL DEFAULT 0,
                failed_count        INTEGER NOT NULL DEFAULT 0,
                started_at          TIMESTAMP,
                finished_at         TIMESTAMP,
                created_at          TIMESTAMP DEFAULT NOW(),
                updated_at          TIMESTAMP DEFAULT NOW()
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS pricing_task_batch_item_runs (
                id                  SERIAL PRIMARY KEY,
                batch_id            INTEGER NOT NULL REFERENCES pricing_task_batches(id) ON DELETE CASCADE,
                boq_item_id         INTEGER NOT NULL REFERENCES boq_items(id) ON DELETE CASCADE,
                boq_project_id      INTEGER NOT NULL REFERENCES boq_projects(id) ON DELETE CASCADE,
                kb_version_id       BIGINT REFERENCES pricing_kb_versions(id),
                status              VARCHAR(16) NOT NULL DEFAULT 'idle',
                reasoning_text      TEXT,
                code_check          JSONB,
                feature_check       JSONB,
                work_procedures     JSONB,
                quota_candidates    JSONB,
                quota_match         JSONB,
                evaluation          JSONB,
                confirmed_results   JSONB,
                conversion_check    JSONB,
                coefficient_check   JSONB,
                step_timings        JSONB,
                error_message       TEXT,
                created_at          TIMESTAMP DEFAULT NOW(),
                started_at          TIMESTAMP,
                finished_at         TIMESTAMP,
                execution_id       BIGINT,
                attempt_count      INTEGER NOT NULL DEFAULT 0,
                current_tool       VARCHAR(32),
                current_tool_status VARCHAR(16),
                current_tool_output TEXT,
                next_attempt_at    TIMESTAMP,
                lease_owner        TEXT,
                lease_expires_at   TIMESTAMP,
                updated_at         TIMESTAMP DEFAULT NOW(),
                UNIQUE(batch_id, boq_item_id)
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS pricing_task_manual_comparison_reviews (
                id                          BIGSERIAL PRIMARY KEY,
                source_type                 VARCHAR(16) NOT NULL,
                source_run_id               BIGINT NOT NULL,
                manual_project_id           INTEGER NOT NULL,
                manual_item_id              INTEGER NOT NULL,
                before_manual_quotas         JSONB NOT NULL,
                after_manual_quotas          JSONB NOT NULL,
                before_evaluation            JSONB NOT NULL,
                after_evaluation             JSONB NOT NULL,
                retained_manual_quota_ids    JSONB NOT NULL DEFAULT '[]'::jsonb,
                accepted_ai_quotas           JSONB NOT NULL DEFAULT '[]'::jsonb,
                operator_name                TEXT NOT NULL DEFAULT '系统',
                created_at                   TIMESTAMP NOT NULL DEFAULT NOW(),
                CONSTRAINT pricing_task_manual_comparison_reviews_source_type_check
                    CHECK (source_type IN ('single', 'batch'))
            )
            """
        )
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS task_id INTEGER")
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS run_id INTEGER")
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS match_reason TEXT")
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS evaluation JSONB")
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS conversion_confirmed BOOLEAN NOT NULL DEFAULT FALSE")
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS conversion_note TEXT")
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS conversion_confirmed_at TIMESTAMP")
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS conversion_resources JSONB")
        cur.execute("ALTER TABLE pricing_task_results ADD COLUMN IF NOT EXISTS conversion_resource_changes JSONB")
        cur.execute("ALTER TABLE pricing_tasks ADD COLUMN IF NOT EXISTS accuracy_report JSONB")
        cur.execute("ALTER TABLE pricing_tasks ADD COLUMN IF NOT EXISTS kb_version_id BIGINT")
        cur.execute("ALTER TABLE pricing_task_runs ADD COLUMN IF NOT EXISTS kb_version_id BIGINT")
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS kb_version_id BIGINT")
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS workflow_version VARCHAR(16) NOT NULL DEFAULT 'legacy'")
        cur.execute("ALTER TABLE pricing_tasks ADD COLUMN IF NOT EXISTS owner_user_id INT REFERENCES users(id)")
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS owner_user_id INT REFERENCES users(id)")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS kb_version_id BIGINT")
        cur.execute("ALTER TABLE pricing_task_runs ADD COLUMN IF NOT EXISTS conversion_check JSONB")
        cur.execute("ALTER TABLE pricing_task_runs ADD COLUMN IF NOT EXISTS coefficient_check JSONB")
        cur.execute("ALTER TABLE pricing_task_runs ADD COLUMN IF NOT EXISTS chapter_rule_check JSONB")
        cur.execute("ALTER TABLE pricing_task_runs ADD COLUMN IF NOT EXISTS accuracy_report JSONB")
        cur.execute("ALTER TABLE pricing_task_runs ADD COLUMN IF NOT EXISTS step_timings JSONB")
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS selected_count INTEGER NOT NULL DEFAULT 0")
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS completed_count INTEGER NOT NULL DEFAULT 0")
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS failed_count INTEGER NOT NULL DEFAULT 0")
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS started_at TIMESTAMP")
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS finished_at TIMESTAMP")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS confirmed_results JSONB")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS coefficient_check JSONB")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS chapter_rule_check JSONB")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS step_timings JSONB")
        cur.execute(
            """
            UPDATE pricing_tasks t SET kb_version_id=a.kb_version_id
            FROM pricing_kb_active_version a WHERE t.kb_version_id IS NULL
            """
        )
        cur.execute(
            """
            UPDATE pricing_task_runs r
            SET kb_version_id=COALESCE(
                (SELECT t.kb_version_id FROM pricing_tasks t WHERE t.id=r.task_id),
                a.kb_version_id
            )
            FROM pricing_kb_active_version a
            WHERE r.kb_version_id IS NULL
            """
        )
        cur.execute(
            """
            UPDATE pricing_task_batches b SET kb_version_id=a.kb_version_id
            FROM pricing_kb_active_version a WHERE b.kb_version_id IS NULL
            """
        )
        cur.execute(
            """
            UPDATE pricing_task_batch_item_runs r SET kb_version_id=b.kb_version_id
            FROM pricing_task_batches b
            WHERE r.batch_id=b.id AND r.kb_version_id IS NULL
            """
        )
        cur.execute("CREATE INDEX IF NOT EXISTS idx_pricing_tasks_project ON pricing_tasks(boq_project_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ptr_boq_item ON pricing_task_results(boq_item_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ptr_run ON pricing_task_results(run_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ptr_task ON pricing_task_results(task_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_pricing_task_runs_task_item ON pricing_task_runs(task_id, boq_item_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_pricing_tasks_kb_version ON pricing_tasks(kb_version_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_pricing_task_runs_kb_version ON pricing_task_runs(kb_version_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_pricing_task_batches_project ON pricing_task_batches(boq_project_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_pricing_tasks_owner ON pricing_tasks(owner_user_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_pricing_task_batches_owner ON pricing_task_batches(owner_user_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ptbir_batch ON pricing_task_batch_item_runs(batch_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_ptbir_batch_item ON pricing_task_batch_item_runs(batch_id, boq_item_id)")
        cur.execute(
            "CREATE INDEX IF NOT EXISTS idx_ptmcr_source "
            "ON pricing_task_manual_comparison_reviews(source_type, source_run_id, created_at DESC)"
        )
        cur.execute("ALTER TABLE pricing_task_batches ADD COLUMN IF NOT EXISTS concurrency_limit INTEGER NOT NULL DEFAULT 20")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS execution_id BIGINT")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS attempt_count INTEGER NOT NULL DEFAULT 0")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS current_tool VARCHAR(32)")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS current_tool_status VARCHAR(16)")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS current_tool_output TEXT")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS next_attempt_at TIMESTAMP")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS lease_owner TEXT")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS lease_expires_at TIMESTAMP")
        cur.execute("ALTER TABLE pricing_task_batch_item_runs ADD COLUMN IF NOT EXISTS updated_at TIMESTAMP DEFAULT NOW()")
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS pricing_task_batch_executions (
                id BIGSERIAL PRIMARY KEY,
                batch_id INTEGER NOT NULL REFERENCES pricing_task_batches(id) ON DELETE CASCADE,
                status VARCHAR(20) NOT NULL DEFAULT 'queued',
                concurrency_limit INTEGER NOT NULL,
                selected_count INTEGER NOT NULL DEFAULT 0,
                completed_count INTEGER NOT NULL DEFAULT 0,
                failed_count INTEGER NOT NULL DEFAULT 0,
                stop_requested_at TIMESTAMP,
                created_at TIMESTAMP NOT NULL DEFAULT NOW(),
                started_at TIMESTAMP,
                finished_at TIMESTAMP,
                updated_at TIMESTAMP NOT NULL DEFAULT NOW()
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS pricing_task_batch_events (
                id BIGSERIAL PRIMARY KEY,
                batch_id INTEGER NOT NULL REFERENCES pricing_task_batches(id) ON DELETE CASCADE,
                execution_id BIGINT REFERENCES pricing_task_batch_executions(id) ON DELETE CASCADE,
                item_run_id INTEGER REFERENCES pricing_task_batch_item_runs(id) ON DELETE CASCADE,
                boq_item_id INTEGER,
                event_type VARCHAR(32) NOT NULL,
                payload JSONB NOT NULL DEFAULT '{}'::jsonb,
                created_at TIMESTAMP NOT NULL DEFAULT NOW()
            )
            """
        )
        cur.execute("CREATE INDEX IF NOT EXISTS idx_background_execution_claim ON pricing_task_batch_executions(status, created_at)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_background_item_claim ON pricing_task_batch_item_runs(execution_id, status, next_attempt_at)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_background_events_batch ON pricing_task_batch_events(batch_id, id)")
    conn.commit()


def _ensure_schema(conn):
    """Initialize once per API process; never run DDL concurrently in requests."""
    global _SCHEMA_READY
    if _SCHEMA_READY:
        return
    with _SCHEMA_LOCK:
        if _SCHEMA_READY:
            return
        try:
            # This initializer commits independently and uses the same advisory
            # lock, so finish it before opening the pricing-task transaction.
            apply_version_schema(conn)
            acquire_schema_transaction_lock(conn)
            _apply_schema(conn)
            _SCHEMA_READY = True
        except Exception:
            conn.rollback()
            raise


def initialize_schema() -> None:
    """Run pricing-task migrations before the API starts accepting traffic."""
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
    finally:
        conn.close()


class PricingTaskCreate(BaseModel):
    name: str
    boq_project_id: int
    quota_library_ids: list[int] = Field(default_factory=list)
    manual_project_id: Optional[int] = None
    legacy_local_id: Optional[str] = None
    kb_version_id: Optional[int] = None


class PricingTaskImportItem(BaseModel):
    id: str
    name: str
    project_id: int
    manual_project_id: Optional[int] = None
    chapter_ids: list[int] = Field(default_factory=list)


class PricingTaskImportRequest(BaseModel):
    tasks: list[PricingTaskImportItem]


class PricingTaskBatchCreate(BaseModel):
    name: str
    boq_project_id: int
    quota_library_ids: list[int] = Field(default_factory=list)
    manual_project_id: Optional[int] = None
    kb_version_id: Optional[int] = None


class NewPricingTaskBatchCreate(BaseModel):
    name: str
    boq_project_id: int
    quota_library_ids: list[int] = Field(default_factory=list)
    manual_project_id: int
    kb_version_id: Optional[int] = None


class BackgroundBatchExecutionCreate(BaseModel):
    boq_item_ids: list[int] = Field(min_length=1)


class RunRequest(BaseModel):
    boq_item_id: int


class ConfirmRunRequest(BaseModel):
    results: Optional[list[dict[str, Any]]] = None


class ManualComparisonAiQuota(BaseModel):
    dekid: int
    dezmid: int


class ManualComparisonUpdateRequest(BaseModel):
    retained_manual_quota_ids: list[int] = Field(default_factory=list)
    accepted_ai_quotas: list[ManualComparisonAiQuota] = Field(default_factory=list)


class ConversionConfirmRequest(BaseModel):
    items: list[dict[str, Any]] = Field(default_factory=list)


_TOOL_SUBMIT_FEATURE_ANALYSIS = {
    "type": "function",
    "function": {
        "name": "submit_feature_analysis",
        "description": "提交项目特征完整性分析结果，说明当前项目特征描述是否足以支撑套定额，并列出缺失的必要特征。",
        "parameters": {
            "type": "object",
            "properties": {
                "is_complete": {"type": "boolean", "description": "项目特征描述是否完整充分，可满足套定额要求"},
                "missing_features": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "缺少的必要特征信息列表，若完整则为空数组",
                },
                "analysis": {"type": "string", "description": "简短分析说明"},
                "default_fills": {
                    "type": "array",
                    "description": "Choose only supplied TQDK_TQDXMTZ native default candidates; do not create or rewrite a default value.",
                    "items": {
                        "type": "object",
                        "properties": {
                            "candidate_id": {"type": "string"},
                            "target_feature_name": {"type": "string"},
                            "match_state": {"type": "string", "enum": ["comprehensive", "vague", "missing"]},
                            "original_feature_text": {"type": "string"},
                            "reason": {"type": "string"},
                            "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
                        },
                        "required": ["candidate_id", "target_feature_name", "match_state", "original_feature_text", "reason", "confidence"],
                        "additionalProperties": False,
                    },
                },
            },
            "required": ["is_complete", "missing_features", "analysis", "default_fills"],
            "additionalProperties": False,
        },
    },
}

_TOOL_SUBMIT_CHAPTER_RULE_CHECK = {
    "type": "function",
    "function": {
        "name": "submit_chapter_rule_check",
        "description": "逐条判断章节说明规则是否命中当前清单，并说明必须执行的组价动作。只评估提供的规则，不得编造规则。",
        "strict": True,
        "parameters": {
            "type": "object",
            "properties": {
                "rules": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "chapter_id": {"type": "integer"},
                            "rule_reference": {"type": "string", "description": "章节内条款序号或可定位的标题"},
                            "rule_text": {"type": "string", "description": "命中的规则原文摘要，不得改变规则含义"},
                            "matched": {"type": "boolean"},
                            "matched_keywords": {"type": "array", "items": {"type": "string"}},
                            "evidence": {"type": "string"},
                            "action": {"type": "string", "description": "命中后必须执行或核查的组价动作；未命中时写不适用原因"},
                            "requires_project_check": {"type": "boolean"},
                            "requires_manual_review": {"type": "boolean"},
                        },
                        "required": ["chapter_id", "rule_reference", "rule_text", "matched", "matched_keywords", "evidence", "action", "requires_project_check", "requires_manual_review"],
                        "additionalProperties": False,
                    },
                },
                "issues": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["rules", "issues"],
            "additionalProperties": False,
        },
    },
}

_TOOL_SUBMIT_CHAPTER_RULE_VALIDATION = {
    "type": "function",
    "function": {
        "name": "submit_chapter_rule_validation",
        "description": "校验最终套取的定额是否落实所有已命中的章节规则。",
        "strict": True,
        "parameters": {
            "type": "object",
            "properties": {
                "validations": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "rule_index": {"type": "integer"},
                            "status": {"type": "string", "enum": ["passed", "warning", "failed", "manual_review"]},
                            "evidence": {"type": "string"},
                            "related_quota_codes": {"type": "array", "items": {"type": "string"}},
                            "message": {"type": "string"},
                        },
                        "required": ["rule_index", "status", "evidence", "related_quota_codes", "message"],
                        "additionalProperties": False,
                    },
                },
                "issues": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["validations", "issues"],
            "additionalProperties": False,
        },
    },
}

_TOOL_SUBMIT_QUOTA_MATCH = {
    "type": "function",
    "function": {
        "name": "submit_quota_match",
        "description": "提交套定额结果。只能从候选定额中按 dekid/dezmid 选择，不允许编造候选外子目。",
        "strict": True,
        "parameters": {
            "type": "object",
            "properties": {
                "matches": {
                    "type": "array",
                    "description": "匹配的定额子目列表",
                    "items": {
                        "type": "object",
                        "properties": {
                            "dekid": {"type": "integer", "description": "候选定额库 ID"},
                            "dezmid": {"type": "integer", "description": "候选定额子目 ID"},
                            "qty_factor": {"type": "number", "description": "工程量系数，一般为1"},
                            "confidence": {"type": "string", "enum": ["high", "medium", "low"], "description": "匹配置信度"},
                            "match_reason": {"type": "string", "description": "选取该定额的简要理由"},
                        },
                        "required": ["dekid", "dezmid", "qty_factor", "confidence", "match_reason"],
                        "additionalProperties": False,
                    },
                },
                "issues": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "影响套定额的模糊或不清楚的问题，若无则为空数组",
                },
            },
            "required": ["matches", "issues"],
            "additionalProperties": False,
        },
    },
}

_TOOL_CHECK_ITEM_CODE = {
    "type": "function",
    "function": {
        "name": "check_item_code",
        "description": "根据工程量清单编码和清单名称，查询标准清单名称并判断编码与名称是否一致。",
        "strict": True,
        "parameters": {
            "type": "object",
            "properties": {
                "item_code": {"type": "string", "description": "工程量清单编码"},
                "item_name": {"type": "string", "description": "工程量清单名称"},
            },
            "required": ["item_code", "item_name"],
            "additionalProperties": False,
        },
    },
}


_TOOL_SUBMIT_CONVERSION_CHECK = {
    "type": "function",
    "function": {
        "name": "submit_conversion_check",
        "description": "提交确认定额的组合定额换算结果。只给组合定额次数建议，不修改已确认定额和工程量系数。",
        "strict": True,
        "parameters": {
            "type": "object",
            "properties": {
                "items": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "dekid": {"type": "integer"},
                            "dezmid": {"type": "integer"},
                            "quota_code": {"type": "string"},
                            "quota_name": {"type": "string"},
                            "needs_conversion": {"type": "boolean"},
                            "reason": {"type": "string"},
                            "requires_manual_review": {"type": "boolean"},
                            "adjustment_rules": {
                                "type": "array",
                                "description": "按 tdek_tzhhs 查询到的组合定额换算规则及项目特征匹配结果。每条规则单独判断，不合并。",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "rule_index": {"type": "integer"},
                                        "prompt": {"type": "string"},
                                        "base_value": {"type": "number"},
                                        "increment_unit": {"type": "number"},
                                        "combo_dezmid": {"type": "integer"},
                                        "combo_code": {"type": "string"},
                                        "combo_name": {"type": "string"},
                                        "combo_unit": {"type": "string"},
                                        "combo_work_content": {"type": "string"},
                                        "matched": {"type": "boolean"},
                                        "matched_feature": {"type": "string"},
                                        "feature_value": {"type": "number"},
                                        "calculated_times": {"type": "number"},
                                        "reason": {"type": "string"},
                                        "requires_manual_review": {"type": "boolean"},
                                        "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
                                    },
                                    "required": [
                                        "rule_index",
                                        "prompt",
                                        "base_value",
                                        "increment_unit",
                                        "combo_dezmid",
                                        "combo_code",
                                        "combo_name",
                                        "combo_unit",
                                        "combo_work_content",
                                        "matched",
                                        "matched_feature",
                                        "feature_value",
                                        "calculated_times",
                                        "reason",
                                        "requires_manual_review",
                                        "confidence",
                                    ],
                                    "additionalProperties": False,
                                },
                            },
                            "missing_inputs": {"type": "array", "items": {"type": "string"}},
                            "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
                        },
                        "required": [
                            "dekid",
                            "dezmid",
                            "quota_code",
                            "quota_name",
                            "needs_conversion",
                            "reason",
                            "requires_manual_review",
                            "adjustment_rules",
                            "missing_inputs",
                            "confidence",
                        ],
                        "additionalProperties": False,
                    },
                },
                "issues": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["items", "issues"],
            "additionalProperties": False,
        },
    },
}


_TOOL_SUBMIT_COEFFICIENT_CHECK = {
    "type": "function",
    "function": {
        "name": "submit_coefficient_check",
        "description": "提交第七步系数换算判断结果。只保存和展示系数，不改写工料机含量。",
        "strict": True,
        "parameters": {
            "type": "object",
            "properties": {
                "items": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "quota_key": {"type": "string"},
                            "source_type": {"type": "string", "enum": ["base", "combo"]},
                            "dekid": {"type": "integer"},
                            "dezmid": {"type": "integer"},
                            "quota_code": {"type": "string"},
                            "quota_name": {"type": "string"},
                            "coefficient_rules": {
                                "type": "array",
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "rule_index": {"type": "integer"},
                                        "tsxx": {"type": "string"},
                                        "hssm": {"type": "string"},
                                        "group_no": {"type": "integer"},
                                        "matched": {"type": "boolean"},
                                        "matched_feature": {"type": "string"},
                                        "feature_value": {"type": "string"},
                                        "factor": {"type": "number"},
                                        "target_resource_types": {
                                            "type": "array",
                                            "items": {"type": "string", "enum": ["1", "2", "3", "all"]},
                                        },
                                        "reason": {"type": "string"},
                                        "requires_manual_review": {"type": "boolean"},
                                        "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
                                    },
                                    "required": [
                                        "rule_index",
                                        "tsxx",
                                        "hssm",
                                        "group_no",
                                        "matched",
                                        "matched_feature",
                                        "feature_value",
                                        "factor",
                                        "target_resource_types",
                                        "reason",
                                        "requires_manual_review",
                                        "confidence",
                                    ],
                                    "additionalProperties": False,
                                },
                            },
                            "missing_inputs": {"type": "array", "items": {"type": "string"}},
                            "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
                        },
                        "required": [
                            "quota_key",
                            "source_type",
                            "dekid",
                            "dezmid",
                            "quota_code",
                            "quota_name",
                            "coefficient_rules",
                            "missing_inputs",
                            "confidence",
                        ],
                        "additionalProperties": False,
                    },
                },
                "issues": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["items", "issues"],
            "additionalProperties": False,
        },
    },
}


class ModelRateLimitError(RuntimeError):
    """A provider-side capacity failure that must not enter item retry loops."""


def _model_gate_limit() -> int:
    try:
        configured = int(os.getenv("PRICING_MODEL_CONCURRENCY", "99"))
    except ValueError:
        configured = 99
    # This is deliberately a global safety ceiling, not a per-batch setting.
    return max(1, min(configured, 99))


def _get_model_gate() -> BoundedSemaphore:
    global _MODEL_GATE, _MODEL_GATE_LIMIT
    limit = _model_gate_limit()
    with _MODEL_GATE_LOCK:
        if _MODEL_GATE is None or _MODEL_GATE_LIMIT != limit:
            # The limit is configured before process startup. Do not replace a
            # live semaphore while calls are active.
            if _MODEL_GATE is None or _MODEL_GATE_RUNNING == 0:
                _MODEL_GATE = BoundedSemaphore(limit)
                _MODEL_GATE_LIMIT = limit
    return _MODEL_GATE or BoundedSemaphore(limit)


def get_model_gate_metrics() -> dict[str, int]:
    with _MODEL_GATE_LOCK:
        limit = _MODEL_GATE_LIMIT or _model_gate_limit()
        return {
            "limit": limit,
            "running": _MODEL_GATE_RUNNING,
            "waiting": _MODEL_GATE_WAITING,
        }


@contextmanager
def _model_call_slot():
    global _MODEL_GATE_RUNNING, _MODEL_GATE_WAITING
    gate = _get_model_gate()
    with _MODEL_GATE_LOCK:
        _MODEL_GATE_WAITING += 1
    gate.acquire()
    with _MODEL_GATE_LOCK:
        _MODEL_GATE_WAITING -= 1
        _MODEL_GATE_RUNNING += 1
    try:
        suspend_connections_for_model_call()
        yield
    finally:
        with _MODEL_GATE_LOCK:
            _MODEL_GATE_RUNNING -= 1
        gate.release()


def _is_model_rate_limited(exc: Exception) -> bool:
    text = f"{type(exc).__name__} {exc}".lower()
    return any(marker in text for marker in (
        "429", "rate limit", "rate_limit", "too many requests", "insufficient quota",
        "quota exceeded", "insufficient balance", "service overloaded", "server overload",
        "模型限流", "服务过载", "配额不足", "余额不足",
    ))


class _GatedStream:
    def __init__(self, stream: Any, slot: Any) -> None:
        self._stream = stream
        self._slot = slot
        self._released = False

    def _release(self) -> None:
        if not self._released:
            self._released = True
            self._slot.__exit__(None, None, None)

    def __iter__(self):
        try:
            yield from self._stream
        except Exception as exc:
            if _is_model_rate_limited(exc):
                raise ModelRateLimitError("模型服务限流或配额不足，请稍后手动重新执行该清单") from exc
            raise
        finally:
            self._release()

    def close(self) -> None:
        try:
            close = getattr(self._stream, "close", None)
            if close:
                close()
        finally:
            self._release()


class _GatedCompletions:
    def __init__(self, client: OpenAI) -> None:
        self._client = client

    def create(self, *args, **kwargs):
        slot = _model_call_slot()
        slot.__enter__()
        try:
            response = self._client.chat.completions.create(*args, **kwargs)
        except Exception as exc:
            slot.__exit__(type(exc), exc, exc.__traceback__)
            if _is_model_rate_limited(exc):
                raise ModelRateLimitError("模型服务限流或配额不足，请稍后手动重新执行该清单") from exc
            raise
        if kwargs.get("stream"):
            return _GatedStream(response, slot)
        slot.__exit__(None, None, None)
        return response


class _GatedChat:
    def __init__(self, client: OpenAI) -> None:
        self.completions = _GatedCompletions(client)


class _GatedOpenAIClient:
    def __init__(self, client: OpenAI) -> None:
        self.chat = _GatedChat(client)


def _client(thinking: bool = True) -> _GatedOpenAIClient:
    # A process-wide client cache would leak private credentials between users.
    # The global gate still protects provider capacity for every fresh client.
    return _GatedOpenAIClient(client_for(active_profile(), timeout=120.0))


def _model() -> str:
    return active_profile().model

_MODEL_CALL_ATTEMPTS = 3
_MODEL_RETRY_DELAYS = (1.5, 4.0)


def _wait_before_model_retry(attempt: int, stage: str, exc: Exception) -> None:
    if attempt >= _MODEL_CALL_ATTEMPTS - 1:
        return
    delay = _MODEL_RETRY_DELAYS[min(attempt, len(_MODEL_RETRY_DELAYS) - 1)]
    print(
        f"[pricing-task] {stage} retrying in {delay:.1f}s after attempt={attempt + 1}: {exc}",
        file=sys.stderr,
        flush=True,
    )
    sleep(delay)


def _user_facing_model_error(exc: Exception) -> str:
    message = str(exc).strip() or type(exc).__name__
    lowered = f"{type(exc).__name__} {message}".lower()
    if "connection" in lowered:
        return "模型服务连接失败，已自动重试 3 次，请稍后重新执行。"
    if "timeout" in lowered or "timed out" in lowered:
        return "模型服务响应超时，已自动重试 3 次，请稍后重新执行。"
    return message


def exec_check_item_code(conn, item_code: str, item_name: str, kb_version_id: int) -> dict[str, Any]:
    base_code = _base_code(item_code)
    with conn.cursor() as cur:
        cur.execute(
            "SELECT zmmc FROM tqdk_tqdzm WHERE kb_version_id=pricing_kb_data_version(%s,'TQDK_TQDZM') AND zmbh=%s LIMIT 5",
            (kb_version_id, base_code),
        )
        rows = cur.fetchall()
    standard_names = sorted({r[0] for r in rows if r[0]})
    standard_name = standard_names[0] if standard_names else ""
    boq = (item_name or "").strip()
    std = standard_name.strip()
    return {
        "item_code": item_code,
        "item_name": item_name,
        "base_code": base_code,
        "standard_name": standard_name,
        "standard_names": standard_names,
        "found": bool(standard_names),
        "is_consistent": bool(std) and (boq == std or boq in std or std in boq),
    }


def _load_feature_default_context(conn, base_code: str, kb_version_id: int) -> dict[str, Any]:
    """Load the versioned feature schema and its native default candidates."""
    context: dict[str, Any] = {
        "feature_schema": [],
        "default_candidates": [],
        "schema_kb_version_id": kb_version_id,
    }
    if not base_code:
        return context

    with conn.cursor() as cur:
        cur.execute("SELECT EXISTS(SELECT 1 FROM tqdk_tqdxmtz WHERE kb_version_id=pricing_kb_data_version(%s,'TQDK_TQDXMTZ'))", (kb_version_id,))
        schema_available = bool(cur.fetchone()[0])
        schema_kb_version_id = kb_version_id
        if not schema_available:
            cur.execute("SELECT kb_version_id FROM pricing_kb_active_version WHERE singleton=TRUE")
            active_row = cur.fetchone()
            if active_row:
                schema_kb_version_id = int(active_row[0])
        context["schema_kb_version_id"] = schema_kb_version_id
        cur.execute(
            """
            SELECT d.tzmc, d.defaulttzms, d.source_rowid
            FROM tqdk_tqdzm zm
            JOIN tqdk_tqdxmtz d
              ON d.qdkid=zm.qdkid
             AND d.qdzmid=zm.id
             AND d.kb_version_id=pricing_kb_data_version(%s,'TQDK_TQDXMTZ')
            WHERE zm.kb_version_id=pricing_kb_data_version(%s,'TQDK_TQDZM')
              AND zm.zmbh=%s
            ORDER BY d.source_rowid
            """,
            (schema_kb_version_id, kb_version_id, base_code),
        )
        schema_rows = cur.fetchall()

    schema_seen: set[str] = set()
    candidate_seen: set[tuple[str, str, str]] = set()
    for feature_name_raw, native_default_raw, source_rowid in schema_rows:
        feature_name = str(feature_name_raw or "").strip()
        native_default = str(native_default_raw or "").strip()
        if not feature_name:
            continue
        if feature_name not in schema_seen:
            schema_seen.add(feature_name)
            context["feature_schema"].append(
                {
                    "feature_name": feature_name,
                    "native_default_value": native_default,
                    "source": "TQDK_TQDXMTZ",
                    "source_rowid": source_rowid,
                }
            )
        if not native_default:
            continue
        key = ("TQDK_TQDXMTZ", feature_name, native_default)
        if key in candidate_seen:
            continue
        candidate_seen.add(key)
        context["default_candidates"].append(
            {
                "candidate_id": f"tqdxmtz:{source_rowid}",
                "source": "TQDK_TQDXMTZ",
                "priority": 1,
                "source_code": base_code,
                "feature_name": feature_name,
                "target_feature_name": feature_name,
                "feature_value": "",
                "default_value": native_default,
                "source_rowid": source_rowid,
            }
        )
    return context


def _replace_comprehensive_feature(text: str, labels: list[str], default_value: str) -> tuple[str, bool]:
    for label in labels:
        normalized_label = str(label or "").strip()
        if not normalized_label:
            continue
        pattern = re.compile(rf"({re.escape(normalized_label)}\s*[:：]\s*)综合考虑")
        if pattern.search(text):
            return pattern.sub(lambda match: f"{match.group(1)}{default_value}", text), True
    return text, False


def _normalize_feature_analysis_result(
    raw: dict[str, Any],
    original_description: str | None,
    base_code: str,
    feature_context: dict[str, list[dict[str, Any]]],
) -> dict[str, Any]:
    original_text = (original_description or "").strip()
    candidates = feature_context.get("default_candidates", [])
    candidate_by_id = {
        str(candidate.get("candidate_id") or ""): candidate
        for candidate in candidates
        if str(candidate.get("candidate_id") or "")
    }
    schema_names = {
        str(item.get("feature_name") or "").strip()
        for item in feature_context.get("feature_schema", [])
        if str(item.get("feature_name") or "").strip()
    }
    fills: list[dict[str, Any]] = []
    review_items: list[dict[str, Any]] = []
    filled_targets: set[str] = set()
    raw_fills = raw.get("default_fills", []) if isinstance(raw, dict) else []
    for raw_fill in raw_fills:
        if not isinstance(raw_fill, dict):
            continue
        candidate = candidate_by_id.get(str(raw_fill.get("candidate_id") or "").strip())
        if not candidate:
            continue
        target_feature_name = str(
            raw_fill.get("target_feature_name")
            or candidate.get("target_feature_name")
            or candidate.get("feature_name")
            or ""
        ).strip()
        if not target_feature_name or target_feature_name in filled_targets:
            continue
        if schema_names and target_feature_name not in schema_names:
            continue
        default_value = str(candidate.get("default_value") or "").strip()
        if not default_value or default_value in original_text:
            continue
        match_state = raw_fill.get("match_state")
        if match_state not in {"comprehensive", "vague", "missing"}:
            match_state = "comprehensive" if "综合考虑" in original_text else "vague"
        original_feature_text = str(raw_fill.get("original_feature_text") or "").strip()
        if match_state == "comprehensive" and "综合考虑" not in original_text:
            continue
        if match_state == "vague" and (not original_feature_text or original_feature_text not in original_text):
            continue
        if match_state == "missing" and original_feature_text:
            continue
        confidence = raw_fill.get("confidence") if raw_fill.get("confidence") in {"high", "medium", "low"} else "low"
        normalized_fill = {
            "candidate_id": candidate["candidate_id"],
            "feature_name": candidate["feature_name"],
            "target_feature_name": target_feature_name,
            "match_state": match_state,
            "original_feature_text": original_feature_text,
            "original_value": original_feature_text or ("综合考虑" if match_state == "comprehensive" else "未明确"),
            "default_value": default_value,
            "source": candidate["source"],
            "source_code": candidate.get("source_code") or base_code,
            "source_rowid": candidate.get("source_rowid"),
            "reason": str(raw_fill.get("reason") or "").strip(),
            "confidence": confidence,
        }
        filled_targets.add(target_feature_name)
        if confidence == "low":
            review_items.append(normalized_fill)
        else:
            fills.append(normalized_fill)

    effective_description = original_text
    appended_fills: list[str] = []
    for fill in fills:
        replaced = False
        if fill["match_state"] == "comprehensive":
            effective_description, replaced = _replace_comprehensive_feature(
                effective_description,
                [fill["target_feature_name"], fill["feature_name"]],
                fill["default_value"],
            )
        if not replaced:
            appended_fills.append(f"{fill['target_feature_name']}：{fill['default_value']}")
    if appended_fills:
        suffix = "\n".join(appended_fills)
        effective_description = f"{effective_description}\n【智能补全项目特征】\n{suffix}".strip()

    unresolved_features = []
    if "综合考虑" in original_text:
        unresolved_features = [
            item["feature_name"]
            for item in feature_context.get("feature_schema", [])
            if item.get("feature_name") not in filled_targets
            and not str(item.get("native_default_value") or "").strip()
        ]
    unresolved_features.extend(
        item["target_feature_name"]
        for item in review_items
        if item["target_feature_name"] not in unresolved_features
    )
    return {
        "is_complete": bool(raw.get("is_complete")) if isinstance(raw, dict) else False,
        "missing_features": [str(v) for v in raw.get("missing_features", []) if str(v).strip()] if isinstance(raw, dict) else [],
        "analysis": str(raw.get("analysis") or "").strip() if isinstance(raw, dict) else "",
        "original_description": original_text,
        "normalized_description": effective_description,
        "effective_description": effective_description,
        "default_fills": fills,
        "default_review_items": review_items,
        "description_updated": False,
        "feature_schema": feature_context.get("feature_schema", []),
        "schema_kb_version_id": feature_context.get("schema_kb_version_id"),
        "default_candidates": candidates,
        "unresolved_features": unresolved_features,
    }


def _load_chapter_rule_context(
    conn, base_code: str, kb_version_id: int, boq_project_id: int
) -> dict[str, Any]:
    """Resolve the BOQ chapter path and its non-empty rule descriptions for one pinned KB version."""
    context: dict[str, Any] = {
        "base_code": base_code,
        "kb_version_id": kb_version_id,
        "chapters": [],
        "project_items": [],
        "available": False,
    }
    if not base_code:
        return context

    with conn.cursor() as cur:
        cur.execute(
            """
            WITH RECURSIVE chapter_path AS (
                SELECT c.qdkid, c.id, c.pid, c.zjmc, c.zjsm, 0 AS depth
                FROM tqdk_tqdzm item
                JOIN tqdk_tzjmc c ON c.qdkid=item.qdkid AND c.id=item.zjh
                WHERE item.kb_version_id=pricing_kb_data_version(%s,'TQDK_TQDZM')
                  AND c.kb_version_id=pricing_kb_data_version(%s,'TQDK_TZJMC')
                  AND item.zmbh=%s
                UNION ALL
                SELECT parent.qdkid, parent.id, parent.pid, parent.zjmc, parent.zjsm, child.depth + 1
                FROM tqdk_tzjmc parent
                JOIN chapter_path child ON child.qdkid=parent.qdkid AND child.pid=parent.id
                WHERE parent.kb_version_id=pricing_kb_data_version(%s,'TQDK_TZJMC')
            )
            SELECT qdkid, id, pid, zjmc, zjsm, depth
            FROM chapter_path
            WHERE LENGTH(REGEXP_REPLACE(COALESCE(zjsm, ''), '\\s+', '', 'g')) > 0
            ORDER BY depth DESC, id
            """,
            (kb_version_id, kb_version_id, base_code, kb_version_id),
        )
        rows = cur.fetchall()
        context["chapters"] = [
            {
                "qdkid": int(row[0]),
                "chapter_id": int(row[1]),
                "parent_id": int(row[2]) if row[2] is not None else None,
                "chapter_name": row[3] or "",
                "zjsm": row[4] or "",
                "depth": int(row[5]),
            }
            for row in rows
        ]
        context["available"] = bool(context["chapters"])
        needs_project_index = any(
            re.search(r"全(?:工)?程|全项目|全.*清单", str(row[4] or ""))
            for row in rows
        )
        if needs_project_index:
            cur.execute(
                """
                SELECT item_code, item_name
                FROM boq_items
                WHERE project_id=%s
                ORDER BY item_seq NULLS LAST, id
                LIMIT 500
                """,
                (boq_project_id,),
            )
            context["project_items"] = [
                {"item_code": row[0] or "", "item_name": row[1] or ""}
                for row in cur.fetchall()
            ]
    return context


def _normalize_chapter_rule_check(raw: dict[str, Any], context: dict[str, Any]) -> dict[str, Any]:
    chapters = context.get("chapters", [])
    allowed_ids = {int(chapter["chapter_id"]) for chapter in chapters}
    chapter_names = {int(chapter["chapter_id"]): chapter.get("chapter_name") or "" for chapter in chapters}
    rules: list[dict[str, Any]] = []
    for item in raw.get("rules", []) if isinstance(raw, dict) else []:
        if not isinstance(item, dict):
            continue
        try:
            chapter_id = int(item.get("chapter_id"))
        except (TypeError, ValueError):
            continue
        if chapter_id not in allowed_ids:
            continue
        rules.append(
            {
                "chapter_id": chapter_id,
                "chapter_name": chapter_names[chapter_id],
                "rule_reference": str(item.get("rule_reference") or "").strip(),
                "rule_text": str(item.get("rule_text") or "").strip(),
                "matched": bool(item.get("matched")),
                "matched_keywords": [str(v).strip() for v in item.get("matched_keywords", []) if str(v).strip()],
                "evidence": str(item.get("evidence") or "").strip(),
                "action": str(item.get("action") or "").strip(),
                "requires_project_check": bool(item.get("requires_project_check")),
                "requires_manual_review": bool(item.get("requires_manual_review")),
            }
        )
    return {
        "available": bool(chapters),
        "base_code": context.get("base_code"),
        "kb_version_id": context.get("kb_version_id"),
        "chapters": chapters,
        "project_items_checked": len(context.get("project_items") or []),
        "rules": rules,
        "issues": [str(v).strip() for v in raw.get("issues", []) if str(v).strip()] if isinstance(raw, dict) else [],
        "validation": {"validations": [], "issues": [], "status": "pending" if rules else "not_applicable"},
    }


def _validate_chapter_rules(
    rule_check: dict[str, Any], matches: list[dict[str, Any]]
) -> dict[str, Any]:
    matched_rules = [rule for rule in rule_check.get("rules", []) if rule.get("matched")]
    if not matched_rules:
        return {"validations": [], "issues": [], "status": "not_applicable"}
    quota_summary = [
        {"code": item.get("zmbh") or "", "name": item.get("zmmc") or "", "reason": item.get("match_reason") or ""}
        for item in matches
    ]
    messages = [
        {"role": "system", "content": build_system_prompt()},
        {
            "role": "user",
            "content": (
                "请校验最终套取的定额是否落实每条已命中的章节规则。"
                "不能根据现有信息确认时标记 manual_review；发现冲突标记 failed。\n\n"
                f"【已命中章节规则】\n{_json_dumps(matched_rules)}\n\n"
                f"【最终定额】\n{_json_dumps(quota_summary)}"
            ),
        },
    ]
    raw = _run_tool_fallback(messages, _TOOL_SUBMIT_CHAPTER_RULE_VALIDATION, 3000)
    validations = []
    for item in raw.get("validations", []) if isinstance(raw, dict) else []:
        if not isinstance(item, dict):
            continue
        rule_index = item.get("rule_index")
        if not isinstance(rule_index, int) or not 0 <= rule_index < len(matched_rules):
            continue
        status = item.get("status")
        if status not in {"passed", "warning", "failed", "manual_review"}:
            status = "manual_review"
        validations.append(
            {
                "rule_index": rule_index,
                "status": status,
                "evidence": str(item.get("evidence") or "").strip(),
                "related_quota_codes": [str(v).strip() for v in item.get("related_quota_codes", []) if str(v).strip()],
                "message": str(item.get("message") or "").strip(),
            }
        )
    issues = [str(v).strip() for v in raw.get("issues", []) if str(v).strip()] if isinstance(raw, dict) else []
    status = "passed"
    if any(item["status"] == "failed" for item in validations):
        status = "failed"
    elif len(validations) < len(matched_rules) or any(item["status"] in {"warning", "manual_review"} for item in validations):
        status = "manual_review"
    return {"validations": validations, "issues": issues, "status": status}

def exec_fetch_quota_candidates(
    conn,
    item_code: str,
    kb_version_id: int,
    quota_library_ids: list[int] | None = None,
) -> dict[str, Any]:
    base_code = _base_code(item_code)
    params: list[Any] = [
        kb_version_id, kb_version_id, kb_version_id,
        kb_version_id, kb_version_id, kb_version_id,
        base_code, base_code,
    ]
    library_filter = ""
    if quota_library_ids:
        library_filter = "AND q.dekid = ANY(%s)"
        params.append(quota_library_ids)
    with conn.cursor() as cur:
        cur.execute(
            f"""
            WITH data_versions AS MATERIALIZED (
                SELECT pricing_kb_data_version(%s,'TQDK_TQDZM') AS tqdzm,
                       pricing_kb_data_version(%s,'TQDK_TQDZY') AS tqdzy,
                       pricing_kb_data_version(%s,'TQDK_TQDZY_SPECIAL') AS tqdzy_special,
                       pricing_kb_data_version(%s,'TDEK_TDEZM') AS tdezm,
                       pricing_kb_data_version(%s,'TLibs') AS tlibs,
                       pricing_kb_data_version(%s,'TDEK_TZJMC') AS tdek_tzjmc
            ), candidate_links AS (
                SELECT cand.dekid, cand.dezmid,
                       'TQDK_TQDZY'::TEXT AS source_table,
                       NULL::BIGINT AS typical_group_id,
                       NULL::TEXT AS typical_group_name
                FROM data_versions versions
                JOIN tqdk_tqdzm zm ON zm.kb_version_id=versions.tqdzm
                JOIN tqdk_tqdzy cand
                  ON cand.kb_version_id=versions.tqdzy
                 AND cand.qdkid=zm.qdkid AND cand.qdzmid=zm.id
                WHERE zm.zmbh=%s

                UNION ALL

                SELECT cand.dekid, cand.dezmid,
                       'TQDK_TQDZY_SPECIAL'::TEXT AS source_table,
                       cand.pid AS typical_group_id,
                       parent.zmmc AS typical_group_name
                FROM data_versions versions
                JOIN tqdk_tqdzm zm ON zm.kb_version_id=versions.tqdzm
                JOIN tqdk_tqdzy_special cand
                  ON cand.kb_version_id=versions.tqdzy_special
                 AND cand.qdkid=zm.qdkid AND cand.qdzmid=zm.id
                LEFT JOIN tqdk_tqdzy_special parent
                  ON parent.kb_version_id=cand.kb_version_id
                 AND parent.qdkid=cand.qdkid
                 AND parent.qdzmid=cand.qdzmid
                 AND parent.id=cand.pid
                WHERE zm.zmbh=%s
                  AND cand.dekid IS NOT NULL
                  AND cand.dezmid IS NOT NULL
            ), merged_candidates AS (
                SELECT dekid, dezmid,
                       ARRAY_AGG(DISTINCT source_table ORDER BY source_table) AS source_tables,
                       ARRAY_AGG(DISTINCT typical_group_id ORDER BY typical_group_id)
                           FILTER (WHERE typical_group_id IS NOT NULL) AS typical_group_ids,
                       ARRAY_AGG(DISTINCT typical_group_name ORDER BY typical_group_name)
                           FILTER (WHERE typical_group_name IS NOT NULL AND typical_group_name <> '') AS typical_group_names
                FROM candidate_links
                GROUP BY dekid, dezmid
            )
            SELECT q.id, q.dekid, l.mc, q.zmbh, q.zmmc, q.dw, q.gznr, c.zjmc,
                   merged.source_tables, merged.typical_group_ids, merged.typical_group_names
            FROM data_versions versions
            JOIN merged_candidates merged ON TRUE
            JOIN tdek_tdezm q
              ON q.kb_version_id=versions.tdezm
             AND q.dekid=merged.dekid AND q.id=merged.dezmid
            JOIN tlibs l ON l.kb_version_id=versions.tlibs AND l.id=q.dekid
            LEFT JOIN tdek_tzjmc c
              ON c.kb_version_id=versions.tdek_tzjmc
             AND c.dekid=q.dekid AND c.id=q.zjh
            WHERE TRUE
              {library_filter}
            ORDER BY q.dekid, q.zmbh NULLS LAST, q.id
            """,
            params,
        )
        rows = cur.fetchall()
    candidates = [
        {
            "id": int(r[0]),
            "dezmid": int(r[0]),
            "dekid": int(r[1]),
            "library_name": r[2],
            "zmbh": r[3],
            "zmmc": r[4],
            "dw": r[5],
            "gznr": r[6] or "",
            "chapter_name": r[7],
            "source_tables": list(r[8] or []),
            "typical_group_ids": [int(value) for value in (r[9] or [])],
            "typical_group_names": list(r[10] or []),
        }
        for r in rows
    ]
    return {
        "item_code": item_code,
        "base_code": base_code,
        "kb_version_id": kb_version_id,
        "quota_library_ids": quota_library_ids or [],
        "candidates": candidates,
        "total": len(candidates),
    }


def build_system_prompt() -> str:
    return (
        "你是专业的建筑工程造价工程师，任务是将招标工程量清单中的清单项与定额子目进行匹配（套定额）。\n"
        "必须基于后端提供的候选定额进行选择，不得编造候选外定额。\n"
        "工程量系数只表示清单工程量与定额单位之间的纯换算关系。\n"
        "置信度：high 表示特征与定额充分匹配；medium 表示主要特征匹配但仍有疑问；low 表示关键特征缺失。\n"
        "请全程使用中文分析。"
    )


COMBO_ADJUSTMENT_RULE_GUIDE = """
第七步组合定额换算规则：
1. 规则来源只使用 tdek_tzhhs。
2. 每条已确认定额可能有 0-N 条组合规则；每条组合规则必须单独判断，不允许合并。
3. 只判断项目特征中的数量特征是否与 tdek_tzhhs.tsxx、组合定额 zmmc 的增减指标匹配。
4. 匹配时提取项目特征值，由后端按 (特征值 - jcz) / zjdw 计算组合定额次数；小数保留，负数截为 0。
5. 缺少数量特征、单位明显不一致或 zjdw 无效时，不计算次数，标记人工复核。
6. 组合定额费用为 0 的说明型规则也要展示和解释，但不得臆造工料机调整。
""".strip()


def _collect_stream_tool_args(stream, reasoning_parts: list[str]) -> str:
    args = ""
    for chunk in stream:
        if not chunk.choices:
            continue
        delta = chunk.choices[0].delta
        if hasattr(delta, "reasoning_content") and delta.reasoning_content:
            reasoning_parts.append(delta.reasoning_content)
        if getattr(delta, "content", None):
            reasoning_parts.append(delta.content)
        if delta.tool_calls:
            for tc in delta.tool_calls:
                if tc.function and tc.function.arguments:
                    args += tc.function.arguments
    return args


def _required_tool_choice(tool: dict[str, Any]) -> dict[str, Any]:
    tool_name = str(tool.get("function", {}).get("name") or "").strip()
    if not tool_name:
        raise ValueError("tool function name is required")
    return {"type": "function", "function": {"name": tool_name}}


def _run_tool_fallback(
    messages: list[dict[str, Any]], tool: dict[str, Any], max_tokens: int
) -> dict[str, Any]:
    last_error: Exception | None = None
    for attempt in range(_MODEL_CALL_ATTEMPTS):
        try:
            response = _client(thinking=False).chat.completions.create(
                model=_model(),
                messages=messages,
                tools=[tool],
                tool_choice=_required_tool_choice(tool),
                extra_body={"thinking": {"type": "disabled"}},
                max_tokens=max_tokens,
                stream=False,
            )
            message = response.choices[0].message
            if not message.tool_calls:
                raise ValueError("model returned no tool call")
            return _parse_json_object(message.tool_calls[0].function.arguments)
        except Exception as exc:
            if isinstance(exc, ModelRateLimitError):
                raise
            last_error = exc
            print(
                f"[pricing-task] non-stream tool error attempt={attempt + 1}: {exc}",
                file=sys.stderr,
                flush=True,
            )
            _wait_before_model_retry(attempt, "non-stream tool", exc)
    raise last_error or RuntimeError("non-stream tool failed")


def _run_stream_tool(messages: list[dict[str, Any]], tool: dict[str, Any], max_tokens: int, reasoning_parts: list[str]) -> dict[str, Any]:
    last_error: Exception | None = None
    for attempt in range(_MODEL_CALL_ATTEMPTS):
        try:
            stream = _client(thinking=True).chat.completions.create(
                model=_model(),
                messages=messages,
                tools=[tool],
                reasoning_effort="high",
                extra_body={"thinking": {"type": "enabled"}},
                max_tokens=max_tokens,
                stream=True,
            )
            raw = _collect_stream_tool_args(stream, reasoning_parts)
            return _parse_json_object(raw)
        except Exception as exc:  # keep the current stage retry-local
            if isinstance(exc, ModelRateLimitError):
                raise
            last_error = exc
            print(f"[pricing-task] stream tool error attempt={attempt + 1}: {exc}", file=sys.stderr, flush=True)
            _wait_before_model_retry(attempt, "stream tool", exc)
    raise last_error or RuntimeError("stream tool failed")


def _stream_tool_call(messages: list[dict[str, Any]], tool: dict[str, Any], max_tokens: int) -> Iterable[tuple[str, Any]]:
    last_error: Exception | None = None
    for attempt in range(_MODEL_CALL_ATTEMPTS):
        raw = ""
        try:
            stream = _client(thinking=True).chat.completions.create(
                model=_model(),
                messages=messages,
                tools=[tool],
                reasoning_effort="high",
                extra_body={"thinking": {"type": "enabled"}},
                max_tokens=max_tokens,
                stream=True,
            )
            for chunk in stream:
                if not chunk.choices:
                    continue
                delta = chunk.choices[0].delta
                token = ""
                if hasattr(delta, "reasoning_content") and delta.reasoning_content:
                    token += delta.reasoning_content
                if getattr(delta, "content", None):
                    token += delta.content
                if token:
                    yield ("reasoning_token", token)
                if delta.tool_calls:
                    for tc in delta.tool_calls:
                        if tc.function and tc.function.arguments:
                            raw += tc.function.arguments
            try:
                tool_result = _parse_json_object(raw)
            except ValueError as exc:
                last_error = exc
                print(
                    f"[pricing-task] stream returned invalid tool arguments attempt={attempt + 1}: {exc}",
                    file=sys.stderr,
                    flush=True,
                )
                _wait_before_model_retry(attempt, "stream tool arguments", exc)
                continue
            yield ("tool_result", tool_result)
            return
        except Exception as exc:
            if isinstance(exc, ModelRateLimitError):
                raise
            last_error = exc
            print(f"[pricing-task] stream tool error attempt={attempt + 1}: {exc}", file=sys.stderr, flush=True)
            _wait_before_model_retry(attempt, "stream tool call", exc)
    try:
        print("[pricing-task] stream tool exhausted; using non-stream fallback", file=sys.stderr, flush=True)
        yield ("tool_result", _run_tool_fallback(messages, tool, max_tokens))
        return
    except Exception as fallback_error:
        print(f"[pricing-task] tool fallback error: {fallback_error}", file=sys.stderr, flush=True)
        raise fallback_error from last_error


def _stream_text_completion(messages: list[dict[str, Any]], max_tokens: int) -> Iterable[tuple[str, Any]]:
    last_error: Exception | None = None
    for attempt in range(_MODEL_CALL_ATTEMPTS):
        parts: list[str] = []
        try:
            stream = _client(thinking=True).chat.completions.create(
                model=_model(),
                messages=messages,
                reasoning_effort="high",
                extra_body={"thinking": {"type": "enabled"}},
                max_tokens=max_tokens,
                stream=True,
            )
            for chunk in stream:
                if not chunk.choices:
                    continue
                delta = chunk.choices[0].delta
                token = ""
                if hasattr(delta, "reasoning_content") and delta.reasoning_content:
                    token += delta.reasoning_content
                if getattr(delta, "content", None):
                    token += delta.content
                if token:
                    parts.append(token)
                    yield ("reasoning_token", token)
            yield ("text_result", "".join(parts))
            return
        except Exception as exc:
            if isinstance(exc, ModelRateLimitError):
                raise
            last_error = exc
            print(f"[pricing-task] stream text error attempt={attempt + 1}: {exc}", file=sys.stderr, flush=True)
            _wait_before_model_retry(attempt, "stream text", exc)
    raise last_error or RuntimeError("stream text failed")


def _run_submit_match(messages: list[dict[str, Any]]) -> dict[str, Any]:
    last_error: Exception | None = None
    for attempt in range(_MODEL_CALL_ATTEMPTS):
        try:
            resp = _client(thinking=False).chat.completions.create(
                model=_model(),
                messages=messages,
                tools=[_TOOL_SUBMIT_QUOTA_MATCH],
                tool_choice={"type": "function", "function": {"name": "submit_quota_match"}},
                extra_body={"thinking": {"type": "disabled"}},
                max_tokens=8000,
                stream=False,
            )
            msg = resp.choices[0].message
            if not msg.tool_calls:
                raise ValueError(f"AI did not call submit_quota_match: {msg.content!r}")
            call = msg.tool_calls[0]
            if call.function.name != "submit_quota_match":
                raise ValueError(f"unexpected tool call {call.function.name!r}")
            return _parse_json_object(call.function.arguments)
        except Exception as exc:
            if isinstance(exc, ModelRateLimitError):
                raise
            last_error = exc
            print(f"[pricing-task] submit match error attempt={attempt + 1}: {exc}", file=sys.stderr, flush=True)
            _wait_before_model_retry(attempt, "submit match", exc)
    raise last_error or RuntimeError("submit match failed")


def _run_submit_conversion_check(messages: list[dict[str, Any]]) -> dict[str, Any]:
    last_error: Exception | None = None
    for attempt in range(_MODEL_CALL_ATTEMPTS):
        try:
            resp = _client(thinking=False).chat.completions.create(
                model=_model(),
                messages=messages,
                tools=[_TOOL_SUBMIT_CONVERSION_CHECK],
                tool_choice={"type": "function", "function": {"name": "submit_conversion_check"}},
                extra_body={"thinking": {"type": "disabled"}},
                max_tokens=8000,
                stream=False,
            )
            msg = resp.choices[0].message
            if not msg.tool_calls:
                raise ValueError(f"AI did not call submit_conversion_check: {msg.content!r}")
            call = msg.tool_calls[0]
            if call.function.name != "submit_conversion_check":
                raise ValueError(f"unexpected tool call {call.function.name!r}")
            return json.loads(call.function.arguments)
        except Exception as exc:
            if isinstance(exc, ModelRateLimitError):
                raise
            last_error = exc
            print(f"[pricing-task] submit conversion check error attempt={attempt + 1}: {exc}", file=sys.stderr, flush=True)
            _wait_before_model_retry(attempt, "submit conversion check", exc)
    try:
        json_messages = [
            *messages,
            {
                "role": "user",
                "content": (
                    "工具参数生成失败。请改用 JSON Output，仅输出一个合法 JSON 对象，不要使用 Markdown。"
                    "JSON 顶层必须包含 items 和 issues，字段结构与 submit_conversion_check 完全一致。"
                ),
            },
        ]
        resp = _client(thinking=False).chat.completions.create(
            model=_model(),
            messages=json_messages,
            response_format={"type": "json_object"},
            extra_body={"thinking": {"type": "disabled"}},
            max_tokens=8000,
            stream=False,
        )
        return _parse_json_content(resp.choices[0].message.content)
    except Exception as exc:
        if isinstance(exc, ModelRateLimitError):
            raise
        print(f"[pricing-task] conversion JSON fallback error: {exc}", file=sys.stderr, flush=True)
        raise exc from last_error


def _run_submit_coefficient_check(messages: list[dict[str, Any]]) -> dict[str, Any]:
    last_error: Exception | None = None
    for attempt in range(_MODEL_CALL_ATTEMPTS):
        try:
            resp = _client(thinking=False).chat.completions.create(
                model=_model(),
                messages=messages,
                tools=[_TOOL_SUBMIT_COEFFICIENT_CHECK],
                tool_choice={"type": "function", "function": {"name": "submit_coefficient_check"}},
                extra_body={"thinking": {"type": "disabled"}},
                max_tokens=8000,
                stream=False,
            )
            msg = resp.choices[0].message
            if not msg.tool_calls:
                raise ValueError(f"AI did not call submit_coefficient_check: {msg.content!r}")
            call = msg.tool_calls[0]
            if call.function.name != "submit_coefficient_check":
                raise ValueError(f"unexpected tool call {call.function.name!r}")
            return _parse_json_object(call.function.arguments)
        except Exception as exc:
            if isinstance(exc, ModelRateLimitError):
                raise
            last_error = exc
            print(f"[pricing-task] submit coefficient check error attempt={attempt + 1}: {exc}", file=sys.stderr, flush=True)
            _wait_before_model_retry(attempt, "submit coefficient check", exc)
    try:
        json_messages = [
            *messages,
            {
                "role": "user",
                "content": (
                    "工具参数生成失败。请改用 JSON Output，仅输出一个合法 JSON 对象，不要使用 Markdown。"
                    "JSON 顶层必须包含 items 和 issues，字段结构与 submit_coefficient_check 完全一致。"
                ),
            },
        ]
        resp = _client(thinking=False).chat.completions.create(
            model=_model(),
            messages=json_messages,
            response_format={"type": "json_object"},
            extra_body={"thinking": {"type": "disabled"}},
            max_tokens=8000,
            stream=False,
        )
        return _parse_json_content(resp.choices[0].message.content)
    except Exception as exc:
        if isinstance(exc, ModelRateLimitError):
            raise
        print(f"[pricing-task] coefficient JSON fallback error: {exc}", file=sys.stderr, flush=True)
        raise exc from last_error


def _normalize_matches(raw_match: dict[str, Any], candidates: list[dict[str, Any]]) -> dict[str, Any]:
    candidate_by_key = {(int(c["dekid"]), int(c["dezmid"])): c for c in candidates}
    normalized = []
    seen: set[tuple[int, int]] = set()
    for m in raw_match.get("matches", []):
        key = (int(m.get("dekid")), int(m.get("dezmid")))
        if key in seen:
            continue
        cand = candidate_by_key.get(key)
        if not cand:
            continue
        seen.add(key)
        confidence = m.get("confidence") if m.get("confidence") in {"high", "medium", "low"} else "low"
        try:
            qty_factor = float(m.get("qty_factor", 1) or 1)
        except Exception:
            qty_factor = 1.0
        normalized.append(
            {
                "dekid": cand["dekid"],
                "dezmid": cand["dezmid"],
                "zmbh": cand["zmbh"],
                "zmmc": cand["zmmc"],
                "dw": cand["dw"],
                "library_name": cand["library_name"],
                "chapter_name": cand.get("chapter_name"),
                "qty_factor": qty_factor,
                "confidence": confidence,
                "match_reason": m.get("match_reason") or "",
            }
        )
    return {"matches": normalized, "issues": raw_match.get("issues", [])}


def _manual_quota_payload(row: Any) -> dict[str, Any]:
    quota_code = str(row[2] or "").strip()
    return {
        "id": int(row[0]),
        "boq_item_id": int(row[1]),
        "quota_code": quota_code,
        "quota_name": row[3],
        "quota_unit": row[4],
        "quantity": float(row[5]) if row[5] is not None else None,
        "qty_factor": float(row[6]) if row[6] is not None else None,
        "quota_item_id": int(row[7]) if row[7] is not None else None,
        "is_formula": any(symbol in quota_code for symbol in ("+", "*", "/")),
    }


def _manual_quotas(conn, manual_project_id: int | None, item_code: str | None) -> list[dict[str, Any]]:
    if not manual_project_id or not item_code:
        return []
    full_code = (item_code or "").strip().replace(" ", "")
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT q.id, q.boq_item_id, q.quota_code, q.quota_name, q.quota_unit,
                   q.quantity, q.qty_factor, q.quota_item_id
            FROM manual_boq_items i
            JOIN manual_boq_quotas q ON q.boq_item_id = i.id
            WHERE i.project_id = %s
              AND replace(trim(COALESCE(i.item_code, '')), ' ', '') = %s
            ORDER BY q.id
            """,
            (manual_project_id, full_code),
        )
        rows = cur.fetchall()
    return [_manual_quota_payload(row) for row in rows]


def _manual_quotas_by_code(conn, manual_project_id: int | None) -> dict[str, list[dict[str, Any]]]:
    if not manual_project_id:
        return {}
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT replace(trim(COALESCE(i.item_code, '')), ' ', '') AS item_code,
                   q.id, q.boq_item_id, q.quota_code, q.quota_name, q.quota_unit,
                   q.quantity, q.qty_factor, q.quota_item_id
            FROM manual_boq_items i
            JOIN manual_boq_quotas q ON q.boq_item_id = i.id
            WHERE i.project_id = %s
            ORDER BY i.id, q.id
            """,
            (manual_project_id,),
        )
        rows = cur.fetchall()
    result: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        result.setdefault(str(row[0] or ""), []).append(_manual_quota_payload(row[1:]))
    return result


def _evaluate(matches: list[dict[str, Any]], manual_quotas: list[dict[str, Any]]) -> dict[str, Any]:
    ai_codes = {str(m.get("zmbh")).strip() for m in matches if m.get("zmbh")}
    manual_codes = {str(q.get("quota_code")).strip() for q in manual_quotas if q.get("quota_code")}

    # Manual exports may carry suffixes such as "换"; treat "010001-32" as hit
    # when a manual code is "010001-32换".
    hit_codes = sorted(ai_code for ai_code in ai_codes if any(ai_code in manual_code for manual_code in manual_codes))
    missed_codes = sorted(manual_code for manual_code in manual_codes if not any(ai_code in manual_code for ai_code in ai_codes))
    extra_codes = sorted(ai_code for ai_code in ai_codes if not any(ai_code in manual_code for manual_code in manual_codes))
    return {
        "manual_quotas": manual_quotas,
        "hit_codes": hit_codes,
        "missed_codes": missed_codes,
        "extra_codes": extra_codes,
        "consistent_codes": hit_codes,
        "manual_only_codes": missed_codes,
        "ai_only_codes": extra_codes,
        "hit_count": len(hit_codes),
        "missed_count": len(missed_codes),
        "extra_count": len(extra_codes),
        "consistent_count": len(hit_codes),
        "manual_only_count": len(missed_codes),
        "ai_only_count": len(extra_codes),
        "manual_count": len(manual_codes),
        "ai_count": len(ai_codes),
    }


def _normalize_accuracy_report(raw: dict[str, Any], evaluation: dict[str, Any]) -> dict[str, Any]:
    def text(key: str, default: str = "") -> str:
        return str(raw.get(key) or default).strip()

    def text_list(key: str, limit: int = 8) -> list[str]:
        value = raw.get(key)
        if not isinstance(value, list):
            return []
        return [str(item).strip() for item in value[:limit] if str(item).strip()]

    hit_count = int(evaluation.get("hit_count") or 0)
    missed_count = int(evaluation.get("missed_count") or 0)
    extra_count = int(evaluation.get("extra_count") or 0)
    manual_count = int(evaluation.get("manual_count") or 0)
    ai_count = int(evaluation.get("ai_count") or 0)
    accuracy_rate = round(hit_count / manual_count, 4) if manual_count else None

    level = text("accuracy_level", "待复核")
    if level not in {"高", "中", "低", "待复核"}:
        level = "待复核"

    return {
        "summary": text("summary", "暂无分析摘要。"),
        "accuracy_level": level,
        "accuracy_rate": accuracy_rate,
        "metrics": {
            "hit_count": hit_count,
            "missed_count": missed_count,
            "extra_count": extra_count,
            "manual_count": manual_count,
            "ai_count": ai_count,
        },
        "key_findings": text_list("key_findings"),
        "matched_analysis": text("matched_analysis"),
        "missed_analysis": text("missed_analysis"),
        "extra_analysis": text("extra_analysis"),
        "business_recommendations": text_list("business_recommendations"),
        "conclusion": text("conclusion"),
        "generated_at": datetime.now().isoformat(),
    }


def _generate_accuracy_report(
    boq_item: dict[str, Any],
    quota_match: dict[str, Any],
    evaluation: dict[str, Any],
    conversion_check: dict[str, Any] | None,
    coefficient_check: dict[str, Any] | None,
) -> dict[str, Any]:
    matches = quota_match.get("matches", []) if isinstance(quota_match, dict) else []
    payload = {
        "清单": {
            "编码": boq_item.get("item_code"),
            "名称": boq_item.get("item_name"),
            "项目特征": boq_item.get("item_description"),
            "单位": boq_item.get("unit"),
            "工程量": boq_item.get("quantity"),
        },
        "智能组价定额": [
            {
                "编码": item.get("zmbh"),
                "名称": item.get("zmmc"),
                "单位": item.get("dw"),
                "系数": item.get("qty_factor"),
                "置信度": item.get("confidence"),
                "理由": item.get("match_reason"),
            }
            for item in matches
        ],
        "人工对比工程定额": evaluation.get("manual_quotas", []),
        "对比统计": {
            "双方一致": evaluation.get("hit_codes", []),
            "仅人工": evaluation.get("missed_codes", []),
            "仅AI": evaluation.get("extra_codes", []),
            "hit_count": evaluation.get("hit_count", 0),
            "missed_count": evaluation.get("missed_count", 0),
            "extra_count": evaluation.get("extra_count", 0),
            "manual_count": evaluation.get("manual_count", 0),
            "ai_count": evaluation.get("ai_count", 0),
        },
        "组合换算": conversion_check or {},
        "系数换算": coefficient_check or {},
    }
    messages = [
        {
            "role": "system",
            "content": (
                "你是资深造价复核专家。请对智能组价结果与当前人工对比基准做一致性分析。"
                "人工对比结果是可修正的复核基准，不得把差异直接认定为AI错误。只基于输入数据，"
                "重点判断双方一致、仅人工、仅AI定额的业务原因，不编造不存在的定额。"
                "说明智能组价结果是否可采纳、哪些地方需要人工复核。请输出合法 JSON。"
            ),
        },
        {
            "role": "user",
            "content": (
                "请按以下 JSON 结构输出："
                "{summary, accuracy_level, key_findings, matched_analysis, missed_analysis, "
                "extra_analysis, business_recommendations, conclusion}。"
                "accuracy_level 只能是 高/中/低/待复核；key_findings 和 business_recommendations 为字符串数组。"
                f"\n\n输入数据：\n{_json_dumps(payload)}"
            ),
        },
    ]
    resp = _client(thinking=False).chat.completions.create(
        model=_model(),
        messages=messages,
        response_format={"type": "json_object"},
        extra_body={"thinking": {"type": "disabled"}},
        max_tokens=4000,
        stream=False,
    )
    raw = _parse_json_content(resp.choices[0].message.content)
    return _normalize_accuracy_report(raw, evaluation)


def _normalize_task_accuracy_report(raw: dict[str, Any], metrics: dict[str, Any]) -> dict[str, Any]:
    def text(key: str, default: str = "") -> str:
        value = raw.get(key)
        if isinstance(value, list):
            return "；".join(str(item).strip() for item in value if str(item).strip())
        return str(value or default).strip()

    def text_list(key: str, limit: int = 12) -> list[str]:
        value = raw.get(key)
        if not isinstance(value, list):
            return []
        return [str(item).strip() for item in value[:limit] if str(item).strip()]

    level = text("accuracy_level", "待复核")
    if level not in {"高", "中", "低", "待复核"}:
        level = "待复核"
    return {
        "summary": text("summary", "暂无分析摘要。"),
        "accuracy_level": level,
        "accuracy_rate": metrics.get("hit_rate"),
        "metrics": metrics,
        "key_findings": text_list("key_findings"),
        "matched_analysis": text("matched_analysis"),
        "missed_analysis": text("missed_analysis"),
        "extra_analysis": text("extra_analysis"),
        "risk_items": text_list("risk_items"),
        "representative_examples": text_list("representative_examples"),
        "business_recommendations": text_list("business_recommendations"),
        "conclusion": text("conclusion"),
        "generated_at": datetime.now().isoformat(),
    }


def _generate_task_accuracy_report(task: dict[str, Any], items: list[dict[str, Any]], metrics: dict[str, Any]) -> dict[str, Any]:
    compact_items = []
    for item in items:
        evaluation = item.get("evaluation") or {}
        quota_match = item.get("quota_match") or {}
        compact_items.append(
            {
                "清单编码": item.get("item_code"),
                "清单名称": item.get("item_name"),
                "项目特征": item.get("item_description"),
                "智能定额": [
                    {
                        "编码": match.get("zmbh"),
                        "名称": match.get("zmmc"),
                        "置信度": match.get("confidence"),
                        "理由": match.get("match_reason"),
                    }
                    for match in (quota_match.get("matches", []) if isinstance(quota_match, dict) else [])
                ],
                "人工定额": evaluation.get("manual_quotas", []) if isinstance(evaluation, dict) else [],
                "双方一致": evaluation.get("hit_codes", []) if isinstance(evaluation, dict) else [],
                "仅人工": evaluation.get("missed_codes", []) if isinstance(evaluation, dict) else [],
                "仅AI": evaluation.get("extra_codes", []) if isinstance(evaluation, dict) else [],
                "组合换算": item.get("conversion_check") or {},
                "系数换算": item.get("coefficient_check") or {},
            }
        )
    payload = {
        "任务": task,
        "总体指标": metrics,
        "清单明细": compact_items,
    }
    messages = [
        {
            "role": "system",
            "content": (
                "你是资深造价复核专家。请对一个组价任务下全部清单的智能组价结果与当前人工对比基准做整体一致性分析。"
                "人工对比结果是可修正的复核基准，不是真值。只基于输入数据，重点分析整体一致率、"
                "仅人工和仅AI定额的业务原因、风险清单类型，并给出可采纳范围与复核建议。请输出合法 JSON。"
            ),
        },
        {
            "role": "user",
            "content": (
                "请按以下 JSON 结构输出："
                "{summary, accuracy_level, key_findings, matched_analysis, missed_analysis, extra_analysis, "
                "risk_items, representative_examples, business_recommendations, conclusion}。"
                "accuracy_level 只能是 高/中/低/待复核；其余数组字段输出字符串数组。"
                f"\n\n输入数据：\n{_json_dumps(payload)}"
            ),
        },
    ]
    resp = _client(thinking=False).chat.completions.create(
        model=_model(),
        messages=messages,
        response_format={"type": "json_object"},
        extra_body={"thinking": {"type": "disabled"}},
        max_tokens=6000,
        stream=False,
    )
    raw = _parse_json_content(resp.choices[0].message.content)
    return _normalize_task_accuracy_report(raw, metrics)


def _quota_code(value: Any) -> str:
    return str(value or "").strip()


def _quota_hit(ai_code: str, manual_codes: list[str]) -> bool:
    return bool(ai_code) and any(ai_code in manual_code for manual_code in manual_codes)


def _manual_hit(manual_code: str, ai_codes: list[str]) -> bool:
    return bool(manual_code) and any(ai_code in manual_code for ai_code in ai_codes if ai_code)


def _build_item_consistency_report(quota_match: dict[str, Any], evaluation: dict[str, Any]) -> dict[str, Any]:
    matches = quota_match.get("matches", []) if isinstance(quota_match, dict) else []
    manual_quotas = evaluation.get("manual_quotas", []) if isinstance(evaluation, dict) else []
    ai_codes = [_quota_code(item.get("zmbh")) for item in matches if _quota_code(item.get("zmbh"))]
    manual_codes = [_quota_code(item.get("quota_code")) for item in manual_quotas if _quota_code(item.get("quota_code"))]

    ai_results = []
    for item in matches:
        code = _quota_code(item.get("zmbh"))
        hit = _quota_hit(code, manual_codes)
        ai_results.append(
            {
                "code": code,
                "name": item.get("zmmc") or "",
                "unit": item.get("dw") or "",
                "qty_factor": item.get("qty_factor"),
                "confidence": item.get("confidence") or "",
                "match_reason": item.get("match_reason") or "",
                "in_manual": hit,
                "analysis": "AI定额与当前人工基准一致。" if hit else "该定额仅存在于AI结果，需复核AI结果或修正人工基准。",
            }
        )

    manual_results = []
    for item in manual_quotas:
        code = _quota_code(item.get("quota_code"))
        hit = _manual_hit(code, ai_codes)
        manual_results.append(
            {
                "code": code,
                "name": item.get("quota_name") or "",
                "unit": item.get("quota_unit") or "",
                "quantity": item.get("quantity"),
                "qty_factor": item.get("qty_factor"),
                "in_ai": hit,
                "analysis": "人工定额与AI结果一致。" if hit else "该定额仅存在于人工基准，需复核人工结果或AI召回与匹配。",
            }
        )

    missed_count = int(evaluation.get("missed_count") or 0) if isinstance(evaluation, dict) else 0
    extra_count = int(evaluation.get("extra_count") or 0) if isinstance(evaluation, dict) else 0
    manual_count = int(evaluation.get("manual_count") or 0) if isinstance(evaluation, dict) else len(manual_codes)
    hit_count = int(evaluation.get("hit_count") or 0) if isinstance(evaluation, dict) else 0
    ai_count = int(evaluation.get("ai_count") or 0) if isinstance(evaluation, dict) else len(ai_codes)

    if manual_count == 0:
        status = "无人工对比"
        summary = "未配置或未找到人工套定额，无法判断一致性。"
    elif missed_count == 0 and extra_count == 0:
        status = "一致"
        summary = "AI套定额与人工套定额一致。"
    elif hit_count > 0:
        status = "部分一致"
        summary = "AI结果与当前人工基准部分一致，存在仅人工或仅AI定额。"
    else:
        status = "不一致"
        summary = "AI结果与当前人工基准未形成一致定额。"

    return {
        "status": status,
        "summary": summary,
        "hit_codes": evaluation.get("hit_codes", []) if isinstance(evaluation, dict) else [],
        "missed_codes": evaluation.get("missed_codes", []) if isinstance(evaluation, dict) else [],
        "extra_codes": evaluation.get("extra_codes", []) if isinstance(evaluation, dict) else [],
        "hit_count": hit_count,
        "missed_count": missed_count,
        "extra_count": extra_count,
        "manual_count": manual_count,
        "ai_count": ai_count,
        "ai_results": ai_results,
        "manual_results": manual_results,
    }


def _round_payload(name: str, step_no: int, data: Any, timing: dict[str, Any] | None = None) -> dict[str, Any]:
    return {
        "step_no": step_no,
        "name": name,
        "data": data,
        "duration_ms": timing.get("duration_ms") if isinstance(timing, dict) else None,
        "started_at": timing.get("started_at") if isinstance(timing, dict) else None,
        "finished_at": timing.get("finished_at") if isinstance(timing, dict) else None,
    }


def _build_pricing_task_detail_report(
    conn,
    task: dict[str, Any],
    rows: list[Any],
    *,
    source_type: str = "single",
) -> dict[str, Any]:
    report_items = []
    metrics = {
        "total_items": len(rows),
        "evaluated_item_count": 0,
        "consistent_item_count": 0,
        "partial_item_count": 0,
        "inconsistent_item_count": 0,
        "no_manual_item_count": 0,
        "hit_count": 0,
        "missed_count": 0,
        "extra_count": 0,
        "manual_count": 0,
        "ai_count": 0,
        "hit_rate": None,
    }

    for row in rows:
        (
            run_id,
            boq_item_id,
            status,
            code_check,
            feature_check,
            chapter_rule_check,
            quota_candidates,
            quota_match,
            evaluation,
            conversion_check,
            coefficient_check,
            step_timings,
            reasoning_text,
            created_at,
            finished_at,
            item_code,
            item_name,
            item_description,
            unit,
            quantity,
            item_seq,
        ) = row
        code_check = code_check if isinstance(code_check, dict) else {}
        feature_check = feature_check if isinstance(feature_check, dict) else {}
        chapter_rule_check = chapter_rule_check if isinstance(chapter_rule_check, dict) else {}
        quota_candidates = quota_candidates if isinstance(quota_candidates, dict) else {}
        quota_match = quota_match if isinstance(quota_match, dict) else {}
        if source_type == "batch":
            evaluation = evaluation if isinstance(evaluation, dict) else {}
        else:
            evaluation = _evaluate(
                quota_match.get("matches", []),
                _manual_quotas(conn, task.get("manual_project_id"), item_code),
            )
        if isinstance(conversion_check, dict):
            conversion_check = (
                _hydrate_conversion_for_batch_run(conn, run_id, conversion_check)
                if source_type == "batch"
                else _hydrate_conversion_for_run(conn, run_id, conversion_check)
            )
        coefficient_check = coefficient_check if isinstance(coefficient_check, dict) else {}
        step_timings = step_timings if isinstance(step_timings, dict) else {}
        consistency = _build_item_consistency_report(quota_match, evaluation)

        if "manual_quotas" in evaluation:
            metrics["evaluated_item_count"] += 1
            metrics["hit_count"] += int(consistency["hit_count"] or 0)
            metrics["missed_count"] += int(consistency["missed_count"] or 0)
            metrics["extra_count"] += int(consistency["extra_count"] or 0)
            metrics["manual_count"] += int(consistency["manual_count"] or 0)
            metrics["ai_count"] += int(consistency["ai_count"] or 0)
        if consistency["status"] == "一致":
            metrics["consistent_item_count"] += 1
        elif consistency["status"] == "部分一致":
            metrics["partial_item_count"] += 1
        elif consistency["status"] == "无人工对比":
            metrics["no_manual_item_count"] += 1
        else:
            metrics["inconsistent_item_count"] += 1

        report_items.append(
            {
                "run_id": run_id,
                "boq_item_id": boq_item_id,
                "status": status,
                "quota_candidate_count": quota_candidates.get("total"),
                "feature_completeness": {
                    "status": (
                        "完整" if feature_check.get("is_complete") is True
                        else "不完整" if "is_complete" in feature_check
                        else ""
                    ),
                    "missing_features": feature_check.get("missing_features") or [],
                    "analysis": feature_check.get("analysis") or "",
                },
                "chapter_rule_check": chapter_rule_check,
                "created_at": created_at.isoformat() if hasattr(created_at, "isoformat") else created_at,
                "finished_at": finished_at.isoformat() if hasattr(finished_at, "isoformat") else finished_at,
                "item": {
                    "id": boq_item_id,
                    "item_seq": item_seq,
                    "item_code": item_code,
                    "item_name": item_name,
                    "item_description": item_description or "",
                    "unit": unit or "",
                    "quantity": float(quantity) if quantity is not None else None,
                },
                "rounds": [
                    _round_payload("编码核查", 1, code_check, step_timings.get("1")),
                    _round_payload("项目特征", 2, feature_check, step_timings.get("2")),
                    _round_payload("章节规则校验", 3, chapter_rule_check, step_timings.get("3")),
                    _round_payload("定额候选", 4, quota_candidates, step_timings.get("4")),
                    _round_payload("AI套定额结果", 5, quota_match, step_timings.get("5")),
                    _round_payload("人工套定额对比", 6, evaluation, step_timings.get("6")),
                    _round_payload("组合换算", 7, conversion_check or {}, step_timings.get("7")),
                    _round_payload("系数换算", 8, coefficient_check or {}, step_timings.get("8")),
                ],
                "ai_quota_results": consistency["ai_results"],
                "manual_quota_results": consistency["manual_results"],
                "consistency": consistency,
                "reasoning_text": reasoning_text or "",
            }
        )

    if metrics["manual_count"]:
        metrics["hit_rate"] = round(metrics["hit_count"] / metrics["manual_count"], 4)

    return {
        "task": task,
        "metrics": metrics,
        "items": report_items,
        "generated_at": datetime.now().isoformat(),
    }


def _join_quota_values(items: list[dict[str, Any]], key: str) -> str:
    return "\n".join(str(item.get(key) or "").strip() for item in items if str(item.get(key) or "").strip())


def _append_report_summary_sheet(wb, report: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "汇总"
    task = report.get("task") or {}
    metrics = report.get("metrics") or {}
    entity_name = "批次" if task.get("report_kind") == "batch" else "任务"
    exact_item_label = "完全一致清单数" if task.get("report_kind") == "batch" else "一致清单数"
    rows = [
        (f"{entity_name}ID", task.get("id")),
        (f"{entity_name}名称", task.get("name")),
        ("组价工程", task.get("project_name")),
        ("人工对比工程", task.get("manual_project_name")),
        ("生成时间", report.get("generated_at")),
        ("清单数", metrics.get("total_items")),
        ("已对比清单数", metrics.get("evaluated_item_count")),
        (exact_item_label, metrics.get("consistent_item_count")),
        ("部分一致清单数", metrics.get("partial_item_count")),
        ("不一致清单数", metrics.get("inconsistent_item_count")),
        ("无人工对比清单数", metrics.get("no_manual_item_count")),
        ("AI定额数", metrics.get("ai_count")),
        ("人工定额数", metrics.get("manual_count")),
        ("双方一致定额数", metrics.get("hit_count")),
        ("仅人工定额数", metrics.get("missed_count")),
        ("仅AI定额数", metrics.get("extra_count")),
        ("对比一致率", metrics.get("hit_rate")),
    ]
    ws.append(["指标", "值"])
    for row in rows:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80


def _append_report_item_sheet(wb, report: dict[str, Any]) -> None:
    ws = wb.create_sheet("清单对比")
    ws.append(
        [
            "序号",
            "清单编码",
            "清单名称",
            "项目特征",
            "项目特征完整性",
            "缺失项目特征",
            "完整性分析",
            "单位",
            "工程量",
            "定额候选数量",
            "AI定额编码",
            "AI定额名称",
            "人工定额编码",
            "人工定额名称",
            "一致性状态",
            "一致性分析",
            "双方一致编码",
            "仅人工编码",
            "仅AI编码",
        ]
    )
    for index, item in enumerate(report.get("items") or [], start=1):
        boq = item.get("item") or {}
        feature_completeness = item.get("feature_completeness") or {}
        consistency = item.get("consistency") or {}
        ai_items = item.get("ai_quota_results") or []
        manual_items = item.get("manual_quota_results") or []
        ws.append(
            [
                index,
                boq.get("item_code"),
                boq.get("item_name"),
                boq.get("item_description"),
                feature_completeness.get("status"),
                "\n".join(feature_completeness.get("missing_features") or []),
                feature_completeness.get("analysis"),
                boq.get("unit"),
                boq.get("quantity"),
                item.get("quota_candidate_count"),
                _join_quota_values(ai_items, "code"),
                _join_quota_values(ai_items, "name"),
                _join_quota_values(manual_items, "code"),
                _join_quota_values(manual_items, "name"),
                consistency.get("status"),
                consistency.get("summary"),
                "\n".join(consistency.get("hit_codes") or []),
                "\n".join(consistency.get("missed_codes") or []),
                "\n".join(consistency.get("extra_codes") or []),
            ]
        )
    widths = [8, 18, 28, 48, 14, 32, 48, 10, 14, 14, 24, 42, 24, 42, 14, 42, 24, 24, 24]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width


def _append_report_quota_sheet(wb, report: dict[str, Any]) -> None:
    ws = wb.create_sheet("定额明细")
    ws.append(
        [
            "序号",
            "清单编码",
            "清单名称",
            "来源",
            "定额编码",
            "定额名称",
            "单位",
            "工程量/系数",
            "置信度",
            "是否命中",
            "分析",
            "匹配理由",
        ]
    )
    row_no = 1
    for item in report.get("items") or []:
        boq = item.get("item") or {}
        for quota in item.get("ai_quota_results") or []:
            row_no += 1
            ws.append(
                [
                    row_no - 1,
                    boq.get("item_code"),
                    boq.get("item_name"),
                    "AI",
                    quota.get("code"),
                    quota.get("name"),
                    quota.get("unit"),
                    quota.get("qty_factor"),
                    quota.get("confidence"),
                    "是" if quota.get("in_manual") else "否",
                    quota.get("analysis"),
                    quota.get("match_reason"),
                ]
            )
        for quota in item.get("manual_quota_results") or []:
            row_no += 1
            ws.append(
                [
                    row_no - 1,
                    boq.get("item_code"),
                    boq.get("item_name"),
                    "人工",
                    quota.get("code"),
                    quota.get("name"),
                    quota.get("unit"),
                    quota.get("quantity") if quota.get("quantity") is not None else quota.get("qty_factor"),
                    "",
                    "是" if quota.get("in_ai") else "否",
                    quota.get("analysis"),
                    "",
                ]
            )
    widths = [8, 18, 28, 10, 18, 42, 10, 14, 12, 10, 48, 48]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width


def _build_pricing_task_detail_report_excel(report: dict[str, Any]) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    _append_report_summary_sheet(wb, report)
    _append_report_item_sheet(wb, report)
    _append_report_quota_sheet(wb, report)

    header_fill = PatternFill("solid", fgColor="E5E7EB")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _confirmed_conversion_context(conn, run_id: int) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT r.id, r.status, i.item_code, i.item_name, i.item_description, i.unit,
                   r.kb_version_id, r.feature_check
            FROM pricing_task_runs r
            JOIN boq_items i ON i.id = r.boq_item_id
            WHERE r.id = %s
            """,
            (run_id,),
        )
        run_row = cur.fetchone()
        if not run_row:
            raise HTTPException(status_code=404, detail="run not found")
        if run_row[1] != "confirmed":
            raise HTTPException(status_code=400, detail="run must be confirmed before conversion check")
        feature_check = run_row[7] if isinstance(run_row[7], dict) else {}
        effective_description = str(feature_check.get("effective_description") or "").strip()
        boq_item = {
            "run_id": run_row[0],
            "status": run_row[1],
            "item_code": run_row[2],
            "item_name": run_row[3],
            "item_description": effective_description or run_row[4] or "",
            "original_item_description": run_row[4] or "",
            "unit": run_row[5] or "",
            "kb_version_id": int(run_row[6]),
        }
        cur.execute(
            """
            SELECT r.dekid, r.dezmid, r.subitem_code, r.subitem_name, r.qty_factor,
                   r.confidence, r.match_reason, q.dw, q.gznr, l.mc
            FROM pricing_task_results r
            LEFT JOIN tdek_tdezm q ON q.kb_version_id=pricing_kb_data_version(%s,'TDEK_TDEZM') AND q.dekid=r.dekid AND q.id=r.dezmid
            LEFT JOIN tlibs l ON l.kb_version_id=pricing_kb_data_version(%s,'TLibs') AND l.id=r.dekid
            WHERE r.run_id = %s AND r.status = 'confirmed'
            ORDER BY r.id
            """,
            (run_row[6], run_row[6], run_id),
        )
        result_rows = cur.fetchall()
        items: list[dict[str, Any]] = []
        for row in result_rows:
            dekid = int(row[0])
            dezmid = int(row[1])
            cur.execute(
                """
                SELECT h.tsxx, h.zmbh, h.jcz, h.zjdw,
                       combo.id, combo.zmmc, combo.dw, combo.gznr,
                       combo.rgf, combo.clf, combo.jxf
                FROM tdek_tzhhs h
                LEFT JOIN tdek_tdezm combo ON combo.kb_version_id=h.kb_version_id AND combo.dekid=h.dekid AND combo.zmbh=h.zmbh
                WHERE h.kb_version_id=pricing_kb_data_version(%s,'TDEK_TZHHS') AND h.dekid=%s AND h.dezmid=%s
                ORDER BY h.source_rowid
                """,
                (run_row[6], dekid, dezmid),
            )
            adjustment_rules = []
            for idx, r in enumerate(cur.fetchall(), start=1):
                combo_dezmid = int(r[4]) if r[4] is not None else 0
                combo_resources = []
                if combo_dezmid:
                    cur.execute(
                        """
                        SELECT zmbh, zmmc, dw, gcl, lx, zycl
                        FROM tdek_tzmgc
                        WHERE kb_version_id=pricing_kb_data_version(%s,'TDEK_TZMGC') AND dekid=%s AND dezmid=%s
                        ORDER BY lx NULLS LAST, source_rowid
                        """,
                        (run_row[6], dekid, combo_dezmid),
                    )
                    combo_resources = [
                        {
                            "code": rr[0] or "",
                            "name": rr[1] or "",
                            "unit": rr[2] or "",
                            "quantity": float(rr[3]) if rr[3] is not None else None,
                            "type": int(rr[4]) if rr[4] is not None else None,
                            "zycl": bool(rr[5]) if rr[5] is not None else None,
                        }
                        for rr in cur.fetchall()
                    ]
                adjustment_rules.append(
                    {
                        "rule_index": idx,
                        "prompt": r[0] or "",
                        "combo_code": r[1] or "",
                        "base_value": float(r[2]) if r[2] is not None else 0,
                        "increment_unit": float(r[3]) if r[3] is not None else 0,
                        "combo_dezmid": combo_dezmid,
                        "combo_name": r[5] or "",
                        "combo_unit": r[6] or "",
                        "combo_work_content": r[7] or "",
                        "combo_labor_cost": float(r[8]) if r[8] is not None else 0,
                        "combo_material_cost": float(r[9]) if r[9] is not None else 0,
                        "combo_machine_cost": float(r[10]) if r[10] is not None else 0,
                        "combo_resources": combo_resources,
                    }
                )
            cur.execute(
                """
                SELECT zmbh, zmmc, dw, gcl, lx, zycl
                FROM tdek_tzmgc
                WHERE kb_version_id=pricing_kb_data_version(%s,'TDEK_TZMGC') AND dekid=%s AND dezmid=%s
                ORDER BY lx NULLS LAST, source_rowid
                """,
                (run_row[6], dekid, dezmid),
            )
            resources = [
                {
                    "code": r[0] or "",
                    "name": r[1] or "",
                    "unit": r[2] or "",
                    "quantity": float(r[3]) if r[3] is not None else None,
                    "type": int(r[4]) if r[4] is not None else None,
                    "zycl": bool(r[5]) if r[5] is not None else None,
                }
                for r in cur.fetchall()
            ]
            items.append(
                {
                    "dekid": dekid,
                    "dezmid": dezmid,
                    "quota_code": row[2] or "",
                    "quota_name": row[3] or "",
                    "current_qty_factor": float(row[4]) if row[4] is not None else 1.0,
                    "confidence": row[5] or "",
                    "match_reason": row[6] or "",
                    "unit": row[7] or "",
                    "work_content": row[8] or "",
                    "library_name": row[9] or "",
                    "resources": resources,
                    "adjustment_rules": adjustment_rules,
                }
            )
    return boq_item, items


def _confirmed_items_from_matches(
    conn, matches: list[dict[str, Any]], kb_version_id: int
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    with conn.cursor() as cur:
        for match in matches:
            try:
                dekid = int(match.get("dekid") or 0)
                dezmid = int(match.get("dezmid") or 0)
            except Exception:
                continue
            if not dekid or not dezmid:
                continue
            cur.execute(
                """
                SELECT q.zmbh, q.zmmc, q.dw, q.gznr, l.mc
                FROM tdek_tdezm q
                LEFT JOIN tlibs l ON l.kb_version_id=q.kb_version_id AND l.id=q.dekid
                WHERE q.kb_version_id=pricing_kb_data_version(%s,'TDEK_TDEZM') AND q.dekid=%s AND q.id=%s
                """,
                (kb_version_id, dekid, dezmid),
            )
            qrow = cur.fetchone()
            if not qrow:
                continue
            cur.execute(
                """
                SELECT h.tsxx, h.zmbh, h.jcz, h.zjdw,
                       combo.id, combo.zmmc, combo.dw, combo.gznr,
                       combo.rgf, combo.clf, combo.jxf
                FROM tdek_tzhhs h
                LEFT JOIN tdek_tdezm combo ON combo.kb_version_id=h.kb_version_id AND combo.dekid=h.dekid AND combo.zmbh=h.zmbh
                WHERE h.kb_version_id=pricing_kb_data_version(%s,'TDEK_TZHHS') AND h.dekid=%s AND h.dezmid=%s
                ORDER BY h.source_rowid
                """,
                (kb_version_id, dekid, dezmid),
            )
            adjustment_rules = []
            for idx, rule_row in enumerate(cur.fetchall(), start=1):
                combo_dezmid = int(rule_row[4]) if rule_row[4] is not None else 0
                combo_resources = _load_combo_resources(conn, kb_version_id, dekid, combo_dezmid, rule_row[1] or "") if combo_dezmid or rule_row[1] else []
                adjustment_rules.append(
                    {
                        "rule_index": idx,
                        "prompt": rule_row[0] or "",
                        "combo_code": rule_row[1] or "",
                        "base_value": float(rule_row[2]) if rule_row[2] is not None else 0,
                        "increment_unit": float(rule_row[3]) if rule_row[3] is not None else 0,
                        "combo_dezmid": combo_dezmid,
                        "combo_name": rule_row[5] or "",
                        "combo_unit": rule_row[6] or "",
                        "combo_work_content": rule_row[7] or "",
                        "combo_labor_cost": float(rule_row[8]) if rule_row[8] is not None else 0,
                        "combo_material_cost": float(rule_row[9]) if rule_row[9] is not None else 0,
                        "combo_machine_cost": float(rule_row[10]) if rule_row[10] is not None else 0,
                        "combo_resources": combo_resources,
                    }
                )
            cur.execute(
                """
                SELECT zmbh, zmmc, dw, gcl, lx, zycl
                FROM tdek_tzmgc
                WHERE kb_version_id=pricing_kb_data_version(%s,'TDEK_TZMGC') AND dekid=%s AND dezmid=%s
                ORDER BY lx NULLS LAST, source_rowid
                """,
                (kb_version_id, dekid, dezmid),
            )
            resources = [
                {
                    "code": r[0] or "",
                    "name": r[1] or "",
                    "unit": r[2] or "",
                    "quantity": float(r[3]) if r[3] is not None else None,
                    "type": int(r[4]) if r[4] is not None else None,
                    "zycl": bool(r[5]) if r[5] is not None else None,
                }
                for r in cur.fetchall()
            ]
            items.append(
                {
                    "dekid": dekid,
                    "dezmid": dezmid,
                    "quota_code": qrow[0] or match.get("zmbh") or "",
                    "quota_name": qrow[1] or match.get("zmmc") or "",
                    "current_qty_factor": float(match.get("qty_factor") or 1),
                    "confidence": match.get("confidence") or "",
                    "match_reason": match.get("match_reason") or "",
                    "unit": qrow[2] or "",
                    "work_content": qrow[3] or "",
                    "library_name": qrow[4] or match.get("library_name") or "",
                    "resources": resources,
                    "adjustment_rules": adjustment_rules,
                }
            )
    return items


def _confirmed_results_from_matches(
    conn, matches: list[dict[str, Any]], kb_version_id: int
) -> list[dict[str, Any]]:
    return [
        {
            "dekid": item["dekid"],
            "dezmid": item["dezmid"],
            "subitem_code": item["quota_code"],
            "subitem_name": item["quota_name"],
            "qty_factor": item.get("current_qty_factor", 1),
            "status": "confirmed",
            "conversion_confirmed": False,
            "conversion_note": "",
            "conversion_confirmed_at": None,
            "conversion_resources": item.get("resources", []),
            "conversion_resource_changes": [],
        }
        for item in _confirmed_items_from_matches(conn, matches, kb_version_id)
    ]


def _float_or_none(value: Any) -> float | None:
    try:
        if value is None or value == "":
            return None
        parsed = float(value)
        if parsed != parsed:
            return None
        return parsed
    except Exception:
        return None


_FULLWIDTH_TRANS = str.maketrans(
    "０１２３４５６７８９．，：（）＞＜＋－",
    "0123456789.,:()><+-",
)

_UNIT_ALIASES: dict[str, list[str]] = {
    "km": ["km", "公里", "千米"],
    "m": ["m", "米"],
    "mm": ["mm", "毫米"],
    "cm": ["cm", "厘米"],
    "t": ["t", "吨"],
    "kg": ["kg", "千克", "公斤"],
    "层": ["层"],
    "芯": ["芯"],
    "孔": ["孔"],
    "次": ["次"],
    "根": ["根"],
    "组": ["组"],
    "台": ["台"],
}

_FEATURE_KEYWORDS = [
    "运距",
    "厚度",
    "高度",
    "深度",
    "孔深",
    "宽度",
    "长度",
    "距离",
    "重量",
    "功率",
    "层数",
    "芯数",
]


def _normalize_feature_text(value: Any) -> str:
    return str(value or "").translate(_FULLWIDTH_TRANS).replace("ＫＭ", "KM").replace("ｋｍ", "km")


def _rule_unit_aliases(rule: dict[str, Any]) -> list[str]:
    text = _normalize_feature_text(
        " ".join(
            [
                str(rule.get("prompt") or ""),
                str(rule.get("combo_name") or ""),
                str(rule.get("combo_unit") or ""),
            ]
        )
    ).lower()
    aliases: list[str] = []
    for canonical, values in _UNIT_ALIASES.items():
        if canonical.lower() in text or any(alias.lower() in text for alias in values):
            aliases.extend(values)
    if not aliases and "每增运" in text:
        aliases.extend(_UNIT_ALIASES["km"])
    return sorted(set(aliases), key=len, reverse=True)


def _rule_keywords(rule: dict[str, Any]) -> list[str]:
    text = _normalize_feature_text(f"{rule.get('prompt') or ''} {rule.get('combo_name') or ''}")
    keywords = [keyword for keyword in _FEATURE_KEYWORDS if keyword in text]
    if not keywords and "每增运" in text:
        keywords.append("运距")
    return keywords


def _infer_feature_value_from_boq(boq_item: dict[str, Any] | None, rule: dict[str, Any]) -> dict[str, Any] | None:
    if not boq_item:
        return None
    text = _normalize_feature_text(
        "\n".join(
            [
                str(boq_item.get("item_description") or ""),
                str(boq_item.get("item_name") or ""),
            ]
        )
    )
    unit_aliases = _rule_unit_aliases(rule)
    if not text or not unit_aliases:
        return None

    alias_pattern = "|".join(re.escape(alias) for alias in unit_aliases)
    pattern = re.compile(rf"(?P<value>\d+(?:\.\d+)?)\s*(?P<unit>{alias_pattern})", re.IGNORECASE)
    keywords = _rule_keywords(rule)
    candidates: list[tuple[int, float, str, str]] = []
    for match in pattern.finditer(text):
        value = _float_or_none(match.group("value"))
        if value is None:
            continue
        start = max(0, match.start() - 24)
        end = min(len(text), match.end() + 24)
        snippet = text[start:end].strip()
        score = 10
        for keyword in keywords:
            keyword_pos = text.rfind(keyword, max(0, match.start() - 30), match.end())
            if keyword_pos >= 0:
                score += 20 - min(19, abs(match.start() - keyword_pos))
        candidates.append((score, value, match.group("unit"), snippet))

    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0], reverse=True)
    _, feature_value, unit, snippet = candidates[0]
    keyword_text = "、".join(keywords) if keywords else "数量特征"
    return {
        "feature_value": feature_value,
        "matched_feature": snippet,
        "reason": f"从项目特征识别到{keyword_text}{feature_value:g}{unit}，匹配组合定额增减指标。",
    }


def _normalize_adjustment_rules(
    raw_item: dict[str, Any] | None,
    confirmed: dict[str, Any],
    boq_item: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    raw_rules = raw_item.get("adjustment_rules", []) if isinstance(raw_item, dict) else []
    raw_by_index: dict[int, dict[str, Any]] = {}
    raw_by_code: dict[str, dict[str, Any]] = {}
    for raw_rule in raw_rules:
        if not isinstance(raw_rule, dict):
            continue
        try:
            raw_by_index[int(raw_rule.get("rule_index") or 0)] = raw_rule
        except Exception:
            pass
        combo_code = str(raw_rule.get("combo_code") or "").strip()
        if combo_code:
            raw_by_code[combo_code] = raw_rule

    normalized: list[dict[str, Any]] = []
    for rule in confirmed.get("adjustment_rules", []):
        rule_index = int(rule.get("rule_index") or 0)
        combo_code = str(rule.get("combo_code") or "")
        raw_rule = raw_by_index.get(rule_index) or raw_by_code.get(combo_code) or {}
        matched = bool(raw_rule.get("matched"))
        feature_value = _float_or_none(raw_rule.get("feature_value"))
        base_value = _float_or_none(rule.get("base_value")) or 0.0
        increment_unit = _float_or_none(rule.get("increment_unit")) or 0.0
        calculated_times: float | None = None
        reason = str(raw_rule.get("reason") or "")
        requires_manual_review = bool(raw_rule.get("requires_manual_review"))
        matched_feature = str(raw_rule.get("matched_feature") or "")
        confidence = raw_rule.get("confidence") if raw_rule.get("confidence") in {"high", "medium", "low"} else "low"

        inferred = _infer_feature_value_from_boq(boq_item, rule)
        if inferred and (feature_value is None or not matched):
            matched = True
            feature_value = inferred["feature_value"]
            matched_feature = matched_feature or str(inferred["matched_feature"])
            reason = str(inferred["reason"])
            requires_manual_review = False
            confidence = "high"

        if matched and feature_value is not None and increment_unit > 0:
            calculated_times = (feature_value - base_value) / increment_unit
            if calculated_times < 0:
                calculated_times = 0.0
                reason = reason or "特征值低于基础值，未生成追加次数。"
            else:
                formula_reason = f"按 ({feature_value:g} - {base_value:g}) / {increment_unit:g} = {calculated_times:g} 计算组合定额次数。"
                reason = f"{reason} {formula_reason}".strip() if reason else formula_reason
        elif matched:
            matched = False
            requires_manual_review = True
            reason = reason or "缺少可计算的项目特征值或增减单位无效，需人工复核。"
        else:
            reason = reason or "项目特征未匹配该组合定额的数量增减指标。"

        normalized.append(
            {
                "rule_index": rule_index,
                "prompt": str(rule.get("prompt") or raw_rule.get("prompt") or ""),
                "base_value": base_value,
                "increment_unit": increment_unit,
                "combo_dezmid": int(rule.get("combo_dezmid") or raw_rule.get("combo_dezmid") or 0),
                "combo_code": combo_code or str(raw_rule.get("combo_code") or ""),
                "combo_name": str(rule.get("combo_name") or raw_rule.get("combo_name") or ""),
                "combo_unit": str(rule.get("combo_unit") or raw_rule.get("combo_unit") or ""),
                "combo_work_content": str(rule.get("combo_work_content") or raw_rule.get("combo_work_content") or ""),
                "combo_labor_cost": _float_or_none(rule.get("combo_labor_cost")) or 0,
                "combo_material_cost": _float_or_none(rule.get("combo_material_cost")) or 0,
                "combo_machine_cost": _float_or_none(rule.get("combo_machine_cost")) or 0,
                "combo_resources": [
                    {
                        "code": str(resource.get("code") or ""),
                        "name": str(resource.get("name") or ""),
                        "unit": str(resource.get("unit") or ""),
                        "quantity": resource.get("quantity"),
                        "type": resource.get("type"),
                        "zycl": resource.get("zycl"),
                    }
                    for resource in rule.get("combo_resources", [])
                    if isinstance(resource, dict)
                ],
                "matched": matched,
                "matched_feature": matched_feature,
                "feature_value": feature_value,
                "calculated_times": calculated_times,
                "reason": reason,
                "requires_manual_review": requires_manual_review,
                "confidence": confidence,
            }
        )
    return normalized


def _default_conversion_check(
    items: list[dict[str, Any]],
    issue: str | None = None,
    boq_item: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "items": [
            {
                "dekid": item["dekid"],
                "dezmid": item["dezmid"],
                "quota_code": item["quota_code"],
                "quota_name": item["quota_name"],
                "needs_conversion": False,
                "reason": "未查询到组合定额规则，默认不建议换算。" if not item.get("adjustment_rules") else "未完成组合定额规则分析。",
                "requires_manual_review": bool(item.get("adjustment_rules")),
                "resources": item.get("resources", []),
                "adjustment_rules": _normalize_adjustment_rules(None, item, boq_item),
                "missing_inputs": [],
                "confidence": "medium",
            }
            for item in items
        ],
        "issues": [issue] if issue else [],
    }


def _normalize_conversion_check(
    raw: dict[str, Any],
    confirmed_items: list[dict[str, Any]],
    boq_item: dict[str, Any] | None = None,
) -> dict[str, Any]:
    by_key = {(item["dekid"], item["dezmid"]): item for item in confirmed_items}
    raw_by_key = {}
    for item in raw.get("items", []):
        try:
            raw_by_key[(int(item.get("dekid")), int(item.get("dezmid")))] = item
        except Exception:
            continue
    normalized = []
    for key, confirmed in by_key.items():
        item = raw_by_key.get(key)
        if not item:
            normalized.extend(_default_conversion_check([confirmed], boq_item=boq_item)["items"])
            continue
        confidence = item.get("confidence") if item.get("confidence") in {"high", "medium", "low"} else "low"
        adjustment_rules = _normalize_adjustment_rules(item, confirmed, boq_item)
        has_calculated_adjustment = any(
            bool(rule.get("matched")) and (_float_or_none(rule.get("calculated_times")) or 0) > 0
            for rule in adjustment_rules
        )
        normalized.append(
            {
                "dekid": confirmed["dekid"],
                "dezmid": confirmed["dezmid"],
                "quota_code": confirmed["quota_code"],
                "quota_name": confirmed["quota_name"],
                "needs_conversion": bool(item.get("needs_conversion")) or has_calculated_adjustment,
                "reason": str(item.get("reason") or ""),
                "requires_manual_review": bool(
                    item.get("requires_manual_review")
                    or any(rule.get("requires_manual_review") for rule in adjustment_rules)
                ),
                "resources": [
                    {
                        "code": str(resource.get("code") or ""),
                        "name": str(resource.get("name") or ""),
                        "unit": str(resource.get("unit") or ""),
                        "quantity": resource.get("quantity"),
                        "type": resource.get("type"),
                        "zycl": resource.get("zycl"),
                    }
                    for resource in confirmed.get("resources", [])
                ],
                "adjustment_rules": adjustment_rules,
                "missing_inputs": [str(v) for v in item.get("missing_inputs", []) if v],
                "confidence": confidence,
            }
        )
    return {"items": normalized, "issues": [str(v) for v in raw.get("issues", []) if v]}


def _load_combo_resources(
    conn, kb_version_id: int, dekid: int,
    combo_dezmid: int | None, combo_code: str | None,
) -> list[dict[str, Any]]:
    with conn.cursor() as cur:
        dezmid = combo_dezmid or 0
        if not dezmid and combo_code:
            cur.execute(
                """
                SELECT id
                FROM tdek_tdezm
                WHERE kb_version_id=pricing_kb_data_version(%s,'TDEK_TDEZM') AND dekid=%s AND zmbh=%s
                LIMIT 1
                """,
                (kb_version_id, dekid, combo_code),
            )
            row = cur.fetchone()
            dezmid = int(row[0]) if row else 0
        if not dezmid:
            return []
        cur.execute(
            """
            SELECT zmbh, zmmc, dw, gcl, lx, zycl
            FROM tdek_tzmgc
            WHERE kb_version_id=pricing_kb_data_version(%s,'TDEK_TZMGC') AND dekid=%s AND dezmid=%s
            ORDER BY lx NULLS LAST, source_rowid
            """,
            (kb_version_id, dekid, dezmid),
        )
        return [
            {
                "code": r[0] or "",
                "name": r[1] or "",
                "unit": r[2] or "",
                "quantity": float(r[3]) if r[3] is not None else None,
                "type": int(r[4]) if r[4] is not None else None,
                "zycl": bool(r[5]) if r[5] is not None else None,
            }
            for r in cur.fetchall()
        ]


def _hydrate_resource_main_material_flags(
    resources: Any, source_resources: list[dict[str, Any]],
) -> None:
    """Backfill absent main-material flags in legacy result snapshots."""
    if not isinstance(resources, list):
        return

    flags_by_key: dict[tuple[str, str, str, Any], list[bool | None]] = {}
    for source in source_resources:
        key = (
            str(source.get("code") or ""),
            str(source.get("name") or ""),
            str(source.get("unit") or ""),
            source.get("type"),
        )
        flags_by_key.setdefault(key, []).append(source.get("zycl"))

    for resource in resources:
        if not isinstance(resource, dict) or "zycl" in resource:
            continue
        key = (
            str(resource.get("code") or ""),
            str(resource.get("name") or ""),
            str(resource.get("unit") or ""),
            resource.get("type"),
        )
        flags = flags_by_key.get(key)
        if flags:
            resource["zycl"] = flags.pop(0)


def _hydrate_conversion_combo_resources(conn, conversion_check: Any, kb_version_id: int) -> Any:
    if not isinstance(conversion_check, dict):
        return conversion_check
    for item in conversion_check.get("items", []):
        if not isinstance(item, dict):
            continue
        try:
            dekid = int(item.get("dekid") or 0)
        except Exception:
            dekid = 0
        if not dekid:
            continue
        try:
            dezmid = int(item.get("dezmid") or 0)
        except Exception:
            dezmid = 0
        if dezmid and isinstance(item.get("resources"), list) and any(
            isinstance(resource, dict) and "zycl" not in resource
            for resource in item["resources"]
        ):
            source_resources = _load_combo_resources(conn, kb_version_id, dekid, dezmid, None)
            _hydrate_resource_main_material_flags(item["resources"], source_resources)
        for rule in item.get("adjustment_rules", []) or []:
            if not isinstance(rule, dict):
                continue
            try:
                combo_dezmid = int(rule.get("combo_dezmid") or 0)
            except Exception:
                combo_dezmid = 0
            combo_resources = rule.get("combo_resources")
            if isinstance(combo_resources, list) and not any(
                isinstance(resource, dict) and "zycl" not in resource
                for resource in combo_resources
            ):
                continue
            source_resources = _load_combo_resources(
                conn,
                kb_version_id,
                dekid,
                combo_dezmid,
                str(rule.get("combo_code") or ""),
            )
            if not combo_resources:
                rule["combo_resources"] = source_resources
            else:
                _hydrate_resource_main_material_flags(combo_resources, source_resources)
    return conversion_check


def _hydrate_conversion_for_run(conn, run_id: int, conversion_check: Any) -> Any:
    with conn.cursor() as cur:
        cur.execute("SELECT kb_version_id FROM pricing_task_runs WHERE id=%s", (run_id,))
        row = cur.fetchone()
    if not row or row[0] is None:
        return conversion_check
    return _hydrate_conversion_combo_resources(conn, conversion_check, int(row[0]))


def _hydrate_conversion_for_batch_run(conn, item_run_id: int, conversion_check: Any) -> Any:
    with conn.cursor() as cur:
        cur.execute(
            "SELECT kb_version_id FROM pricing_task_batch_item_runs WHERE id=%s",
            (item_run_id,),
        )
        row = cur.fetchone()
    if not row or row[0] is None:
        return conversion_check
    return _hydrate_conversion_combo_resources(conn, conversion_check, int(row[0]))


def _load_coefficient_rules(
    conn, kb_version_id: int, dekid: int, dezmid: int
) -> list[dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT tsxx, hssm, COALESCE(groupno, 0)
            FROM tdek_tznhs
            WHERE kb_version_id=pricing_kb_data_version(%s,'TDEK_TZNHS') AND dekid=%s AND dezmid=%s
            ORDER BY groupno NULLS LAST, source_rowid
            """,
            (kb_version_id, dekid, dezmid),
        )
        return [
            {
                "rule_index": idx,
                "tsxx": r[0] or "",
                "hssm": r[1] or "",
                "group_no": int(r[2] or 0),
            }
            for idx, r in enumerate(cur.fetchall(), start=1)
        ]


def _parse_factor_from_text(text: str) -> float:
    match = re.search(r"系数\s*([0-9]+(?:\.[0-9]+)?)", text or "")
    return float(match.group(1)) if match else 1.0


def _target_types_from_text(text: str) -> list[str]:
    value = text or ""
    if "子目乘以系数" in value or "相应子目乘以系数" in value:
        return ["all"]
    targets: list[str] = []
    if "人工" in value:
        targets.append("1")
    if "材料" in value:
        targets.append("2")
    if "机械" in value:
        targets.append("3")
    return targets or ["all"]


def _coefficient_context(conn, run_id: int) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    boq_item, confirmed_items = _confirmed_conversion_context(conn, run_id)
    with conn.cursor() as cur:
        cur.execute("SELECT conversion_check FROM pricing_task_runs WHERE id=%s", (run_id,))
        row = cur.fetchone()
    kb_version_id = int(boq_item["kb_version_id"])
    conversion_check = _hydrate_conversion_combo_resources(
        conn, row[0] if row else None, kb_version_id
    ) or {}
    return boq_item, _coefficient_items_from_context(
        conn, confirmed_items, conversion_check, kb_version_id
    )


def _coefficient_items_from_context(
    conn,
    confirmed_items: list[dict[str, Any]],
    conversion_check: dict[str, Any] | None,
    kb_version_id: int,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for item in confirmed_items:
        coefficient_rules = _load_coefficient_rules(
            conn, kb_version_id, int(item["dekid"]), int(item["dezmid"])
        )
        items.append(
            {
                "quota_key": f"base:{item['dekid']}:{item['dezmid']}",
                "source_type": "base",
                "dekid": item["dekid"],
                "dezmid": item["dezmid"],
                "quota_code": item["quota_code"],
                "quota_name": item["quota_name"],
                "resources": item.get("resources", []),
                "coefficient_rules": coefficient_rules,
            }
        )

    conversion_data = _hydrate_conversion_combo_resources(
        conn, conversion_check or {}, kb_version_id
    ) or {}
    for base_item in conversion_data.get("items", []) if isinstance(conversion_data, dict) else []:
        if not isinstance(base_item, dict):
            continue
        dekid = int(base_item.get("dekid") or 0)
        for rule in base_item.get("adjustment_rules", []) or []:
            if not isinstance(rule, dict):
                continue
            combo_dezmid = int(rule.get("combo_dezmid") or 0)
            if not dekid or not combo_dezmid:
                continue
            coefficient_rules = _load_coefficient_rules(
                conn, kb_version_id, dekid, combo_dezmid
            )
            items.append(
                {
                    "quota_key": f"combo:{dekid}:{combo_dezmid}:{rule.get('combo_code') or ''}",
                    "source_type": "combo",
                    "dekid": dekid,
                    "dezmid": combo_dezmid,
                    "quota_code": str(rule.get("combo_code") or ""),
                    "quota_name": str(rule.get("combo_name") or ""),
                    "resources": rule.get("combo_resources") or [],
                    "combo_times": rule.get("calculated_times"),
                    "base_quota_code": base_item.get("quota_code") or "",
                    "coefficient_rules": coefficient_rules,
                }
            )
    return items


def _default_coefficient_check(items: list[dict[str, Any]], issue: str | None = None) -> dict[str, Any]:
    return {
        "items": [
            {
                "quota_key": item["quota_key"],
                "source_type": item["source_type"],
                "dekid": item["dekid"],
                "dezmid": item["dezmid"],
                "quota_code": item["quota_code"],
                "quota_name": item["quota_name"],
                "resources": item.get("resources", []),
                "coefficient_rules": [
                    {
                        **rule,
                        "matched": False,
                        "matched_feature": "",
                        "feature_value": "",
                        "factor": _parse_factor_from_text(rule.get("hssm", "")),
                        "target_resource_types": _target_types_from_text(rule.get("hssm", "")),
                        "reason": "待识别项目特征是否触发该系数换算说明。",
                        "requires_manual_review": False,
                        "confidence": "medium",
                    }
                    for rule in item.get("coefficient_rules", [])
                ],
                "missing_inputs": [],
                "confidence": "medium",
            }
            for item in items
        ],
        "issues": [issue] if issue else [],
    }


def _normalize_coefficient_check(raw: dict[str, Any], items: list[dict[str, Any]]) -> dict[str, Any]:
    raw_by_key = {str(item.get("quota_key") or ""): item for item in raw.get("items", []) if isinstance(item, dict)}
    normalized_items = []
    for item in items:
        raw_item = raw_by_key.get(item["quota_key"]) or {}
        raw_rules = raw_item.get("coefficient_rules", []) if isinstance(raw_item, dict) else []
        raw_by_index: dict[int, dict[str, Any]] = {}
        for raw_rule in raw_rules:
            if not isinstance(raw_rule, dict):
                continue
            try:
                raw_by_index[int(raw_rule.get("rule_index") or 0)] = raw_rule
            except Exception:
                pass
        rules = []
        for rule in item.get("coefficient_rules", []):
            raw_rule = raw_by_index.get(int(rule.get("rule_index") or 0), {})
            hssm = str(rule.get("hssm") or raw_rule.get("hssm") or "")
            matched = bool(raw_rule.get("matched"))
            factor = _float_or_none(raw_rule.get("factor")) or _parse_factor_from_text(hssm)
            target_types = raw_rule.get("target_resource_types")
            if not isinstance(target_types, list) or not target_types:
                target_types = _target_types_from_text(hssm)
            target_types = [str(v) for v in target_types if str(v) in {"1", "2", "3", "all"}] or _target_types_from_text(hssm)
            rules.append(
                {
                    "rule_index": int(rule.get("rule_index") or 0),
                    "tsxx": str(rule.get("tsxx") or raw_rule.get("tsxx") or ""),
                    "hssm": hssm,
                    "group_no": int(rule.get("group_no") or raw_rule.get("group_no") or 0),
                    "matched": matched,
                    "matched_feature": str(raw_rule.get("matched_feature") or ""),
                    "feature_value": str(raw_rule.get("feature_value") or ""),
                    "factor": factor,
                    "target_resource_types": target_types,
                    "reason": str(raw_rule.get("reason") or ("项目特征未触发该系数换算说明。" if not matched else "")),
                    "requires_manual_review": bool(raw_rule.get("requires_manual_review")),
                    "confidence": raw_rule.get("confidence") if raw_rule.get("confidence") in {"high", "medium", "low"} else "low",
                }
            )
        normalized_items.append(
            {
                "quota_key": item["quota_key"],
                "source_type": item["source_type"],
                "dekid": item["dekid"],
                "dezmid": item["dezmid"],
                "quota_code": item["quota_code"],
                "quota_name": item["quota_name"],
                "resources": item.get("resources", []),
                "coefficient_rules": rules,
                "missing_inputs": [str(v) for v in raw_item.get("missing_inputs", []) if v] if isinstance(raw_item, dict) else [],
                "confidence": raw_item.get("confidence") if isinstance(raw_item, dict) and raw_item.get("confidence") in {"high", "medium", "low"} else "medium",
            }
        )
    return {"items": normalized_items, "issues": [str(v) for v in raw.get("issues", []) if v]}


def _create_run(conn, task_id: int | None, boq_item: dict[str, Any], kb_version_id: int) -> int:
    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO pricing_task_runs(task_id, boq_item_id, boq_project_id, kb_version_id, status)
            VALUES (%s, %s, %s, %s, 'running')
            RETURNING id
            """,
            (task_id, boq_item["id"], boq_item["project_id"], kb_version_id),
        )
        run_id = cur.fetchone()[0]
    conn.commit()
    return run_id


def _update_run(conn, run_id: int, **fields: Any) -> None:
    if not fields:
        return
    assignments = []
    values = []
    for key, value in fields.items():
        assignments.append(f"{key} = %s")
        values.append(Json(value, dumps=_json_dumps) if key in {
            "code_check",
            "feature_check",
            "chapter_rule_check",
            "work_procedures",
            "quota_candidates",
            "quota_match",
            "evaluation",
            "conversion_check",
            "coefficient_check",
            "step_timings",
        } else value)
    values.append(run_id)
    with conn.cursor() as cur:
        cur.execute(f"UPDATE pricing_task_runs SET {', '.join(assignments)} WHERE id = %s", values)
    conn.commit()


def _json_value(value: Any) -> Any:
    return Json(value, dumps=_json_dumps)


def _update_batch_item_run(conn, item_run_id: int, *, commit: bool = True, **fields: Any) -> None:
    if not fields:
        return
    json_fields = {
        "code_check",
        "feature_check",
        "chapter_rule_check",
        "work_procedures",
        "quota_candidates",
        "quota_match",
        "evaluation",
        "confirmed_results",
        "conversion_check",
        "coefficient_check",
        "step_timings",
    }
    assignments = []
    values = []
    for key, value in fields.items():
        assignments.append(f"{key} = %s")
        values.append(_json_value(value) if key in json_fields else value)
    values.append(item_run_id)
    with conn.cursor() as cur:
        cur.execute(f"UPDATE pricing_task_batch_item_runs SET {', '.join(assignments)} WHERE id = %s", values)
    if commit:
        conn.commit()


def _load_batch_step_timings(conn, item_run_id: int) -> dict[str, Any]:
    with conn.cursor() as cur:
        cur.execute("SELECT step_timings FROM pricing_task_batch_item_runs WHERE id=%s", (item_run_id,))
        row = cur.fetchone()
    return dict(row[0] or {}) if row else {}


def _finish_batch_step_timing(
    conn,
    item_run_id: int,
    timings: dict[str, Any],
    step_no: int,
    name: str,
    started_at: datetime,
    started_perf: float,
) -> dict[str, Any]:
    timing = _finish_step_timing(conn, None, timings, step_no, name, started_at, started_perf)
    _update_batch_item_run(conn, item_run_id, step_timings=timings)
    return timing


def _load_step_timings(conn, run_id: int) -> dict[str, Any]:
    with conn.cursor() as cur:
        cur.execute("SELECT step_timings FROM pricing_task_runs WHERE id=%s", (run_id,))
        row = cur.fetchone()
    return dict(row[0] or {}) if row else {}


def _finish_step_timing(
    conn,
    run_id: int | None,
    timings: dict[str, Any],
    step_no: int,
    name: str,
    started_at: datetime,
    started_perf: float,
) -> dict[str, Any]:
    finished_at = datetime.now()
    timing = {
        "step_no": step_no,
        "name": name,
        "duration_ms": int(round((perf_counter() - started_perf) * 1000)),
        "started_at": started_at.isoformat(timespec="milliseconds"),
        "finished_at": finished_at.isoformat(timespec="milliseconds"),
    }
    timings[str(step_no)] = timing
    if run_id is not None:
        _update_run(conn, run_id, step_timings=timings)
    return timing


def _load_confirmed_results(conn, run_ids: list[int]) -> dict[int, list[dict[str, Any]]]:
    if not run_ids:
        return {}
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT run_id, dekid, dezmid, subitem_code, subitem_name, qty_factor, status,
                   conversion_confirmed, conversion_note, conversion_confirmed_at,
                   conversion_resources, conversion_resource_changes
            FROM pricing_task_results
            WHERE run_id = ANY(%s::int[]) AND status = 'confirmed'
            ORDER BY id
            """,
            (run_ids,),
        )
        rows = cur.fetchall()
    by_run: dict[int, list[dict[str, Any]]] = {run_id: [] for run_id in run_ids}
    for row in rows:
        by_run.setdefault(int(row[0]), []).append(
            {
                "dekid": int(row[1]),
                "dezmid": int(row[2]),
                "subitem_code": row[3] or "",
                "subitem_name": row[4] or "",
                "qty_factor": float(row[5]) if row[5] is not None else 1.0,
                "status": row[6],
                "conversion_confirmed": bool(row[7]),
                "conversion_note": row[8] or "",
                "conversion_confirmed_at": row[9],
                "conversion_resources": row[10] or [],
                "conversion_resource_changes": row[11] or [],
            }
        )
    return by_run


def _save_pending_results(
    conn,
    task_id: int | None,
    run_id: int,
    boq_item: dict[str, Any],
    match_result: dict[str, Any],
    evaluation: dict[str, Any],
) -> None:
    with conn.cursor() as cur:
        cur.execute("DELETE FROM pricing_task_results WHERE run_id = %s", (run_id,))
        for m in match_result.get("matches", []):
            cur.execute(
                """
                INSERT INTO pricing_task_results(
                    task_id, run_id, boq_item_id, boq_project_id, dezmid, dekid,
                    subitem_code, subitem_name, qty_factor, confidence, ai_reasoning,
                    match_reason, status, evaluation
                )
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'pending',%s)
                """,
                (
                    task_id,
                    run_id,
                    boq_item["id"],
                    boq_item["project_id"],
                    m["dezmid"],
                    m["dekid"],
                    m.get("zmbh"),
                    m.get("zmmc"),
                    m.get("qty_factor", 1),
                    m.get("confidence"),
                    m.get("match_reason"),
                    m.get("match_reason"),
                    Json(evaluation, dumps=_json_dumps),
                ),
            )
    conn.commit()


def _stream_pricing_item(
    conn,
    boq_item: dict[str, Any],
    quota_library_ids: list[int],
    manual_project_id: int | None,
    task_id: int | None,
    run_id: int | None,
    kb_version_id: int,
    *,
    persist_run: bool = True,
) -> Iterable[tuple[str, Any]]:
    system_prompt = build_system_prompt()
    step_timings: dict[str, Any] = {}

    step_started_at = datetime.now()
    step_started_perf = perf_counter()
    code_check = exec_check_item_code(
        conn, boq_item["item_code"], boq_item["item_name"], kb_version_id
    )
    yield ("code_check", code_check)
    yield ("judgment", {"is_consistent": code_check["is_consistent"], "reasoning": f"标准清单名称：{code_check['standard_name'] or '未找到'}"})
    if persist_run and run_id is not None:
        _update_run(conn, run_id, code_check=code_check)
    yield ("step_timing", _finish_step_timing(conn, run_id, step_timings, 1, "编码核查", step_started_at, step_started_perf))

    step_started_at = datetime.now()
    step_started_perf = perf_counter()
    feature_context = _load_feature_default_context(
        conn,
        code_check.get("base_code") or _base_code(boq_item["item_code"]),
        kb_version_id,
    )
    feature_schema_text = _json_dumps(feature_context["feature_schema"]) if feature_context["feature_schema"] else "[]"
    feature_default_text = _json_dumps(feature_context["default_candidates"]) if feature_context["default_candidates"] else "[]"
    messages_r2 = [
        {"role": "system", "content": system_prompt},
        {
            "role": "user",
            "content": (
                "Analyze whether this BOQ item's feature description is sufficient for quota matching.\n\n"
                f"BOQ code: {boq_item['item_code']}\nBOQ name: {boq_item['item_name']}\n"
                f"Original features: {boq_item.get('item_description') or 'not provided'}\n"
                f"Unit: {boq_item.get('unit') or 'not provided'}\n\n"
                f"[Standard feature schema from TQDK_TQDXMTZ]\n{feature_schema_text}\n\n"
                f"[Native default candidates from TQDK_TQDXMTZ.DEFAULTTZMS]\n{feature_default_text}\n\n"
                "Evaluate every supplied native default candidate against the original BOQ features. Select a candidate only when its target feature is 综合考虑, missing, or vague/incomplete (for example 详见设计, 未注明, or a category without the required grade/count/method). "
                "Use match_state=comprehensive, vague, or missing and quote the relevant original text in original_feature_text; use an empty string when the feature is missing. "
                "Do not select a candidate when the original feature already contains an explicit concrete value, even if it differs from the default. "
                "All candidates come from TQDK_TQDXMTZ.DEFAULTTZMS. Preserve conditional default expressions verbatim; do not resolve or rewrite them. "
                "Do not overwrite an explicit original value and do not invent defaults. Use low confidence for uncertain cases; low-confidence items will require manual review and will not enter downstream pricing context."
            ),
        },
    ]
    feature_result: dict[str, Any] = {}
    for event_type, data in _stream_tool_call(messages_r2, _TOOL_SUBMIT_FEATURE_ANALYSIS, 4000):
        if event_type == "reasoning_token":
            yield ("reasoning_token", data)
        else:
            feature_result = data
    feature_result = _normalize_feature_analysis_result(
        feature_result,
        boq_item.get("item_description"),
        code_check.get("base_code") or _base_code(boq_item["item_code"]),
        feature_context,
    )
    effective_description = str(feature_result.get("effective_description") or "").strip()
    if feature_result.get("default_fills") and effective_description:
        # Preserve the source BOQ; downstream steps consume this run's enriched description only.
        boq_item["item_description"] = effective_description
    yield ("feature_check", feature_result)
    if persist_run and run_id is not None:
        _update_run(conn, run_id, feature_check=feature_result)
    yield ("step_timing", _finish_step_timing(conn, run_id, step_timings, 2, "项目特征", step_started_at, step_started_perf))

    step_started_at = datetime.now()
    step_started_perf = perf_counter()
    chapter_context = _load_chapter_rule_context(
        conn,
        code_check.get("base_code") or _base_code(boq_item["item_code"]),
        kb_version_id,
        int(boq_item["project_id"]),
    )
    if chapter_context["available"]:
        rule_messages = [
            {"role": "system", "content": system_prompt},
            {
                "role": "user",
                "content": (
                    "请按章节说明逐条识别当前清单命中的规则，并给出必须执行的组价动作。"
                    "规则来自知识库，命中后必须在后续套定额中执行；不得编造规则。"
                    "若章节说明要求核查全工程清单，只能依据提供的工程清单索引判断。\n\n"
                    f"【清单项】\n编码：{boq_item['item_code']}\n名称：{boq_item['item_name']}\n"
                    f"项目特征：{boq_item.get('item_description') or '（未填写）'}\n单位：{boq_item.get('unit') or '无'}\n\n"
                    f"【章节说明规则（父章节在前，所属章节优先）】\n{_json_dumps(chapter_context['chapters'])}\n\n"
                    f"【全工程清单索引】\n{_json_dumps(chapter_context['project_items'])}"
                ),
            },
        ]
        raw_rule_check: dict[str, Any] = {}
        for event_type, data in _stream_tool_call(rule_messages, _TOOL_SUBMIT_CHAPTER_RULE_CHECK, 5000):
            if event_type == "reasoning_token":
                yield ("reasoning_token", data)
            else:
                raw_rule_check = data
        chapter_rule_check = _normalize_chapter_rule_check(raw_rule_check, chapter_context)
    else:
        chapter_rule_check = _normalize_chapter_rule_check({}, chapter_context)
        chapter_rule_check["issues"] = ["未找到该清单对应的章节说明规则"]
    yield ("chapter_rule_check", chapter_rule_check)
    if persist_run and run_id is not None:
        _update_run(conn, run_id, chapter_rule_check=chapter_rule_check)
    yield ("step_timing", _finish_step_timing(conn, run_id, step_timings, 3, "章节规则校验", step_started_at, step_started_perf))

    step_started_at = datetime.now()
    step_started_perf = perf_counter()
    candidates_data = exec_fetch_quota_candidates(
        conn, boq_item["item_code"], kb_version_id, quota_library_ids
    )
    yield ("quota_candidates", candidates_data)
    if persist_run and run_id is not None:
        _update_run(conn, run_id, quota_candidates=candidates_data)
    yield ("step_timing", _finish_step_timing(conn, run_id, step_timings, 4, "定额候选", step_started_at, step_started_perf))

    step_started_at = datetime.now()
    step_started_perf = perf_counter()
    candidates = candidates_data["candidates"]
    candidate_text = "\n".join(
        f"[{i + 1}] dekid={c['dekid']} dezmid={c['dezmid']} 编码={c['zmbh']} 名称={c['zmmc']} 单位={c['dw']} "
        f"库={c['library_name']} 章节={c.get('chapter_name') or ''} 来源={'、'.join(c.get('source_tables') or [])}"
        + (f" 典型组价={'、'.join(c.get('typical_group_names') or [])}" if c.get("typical_group_names") else "")
        + (f"\n    工作内容：{c['gznr']}" if c.get("gznr") else "")
        for i, c in enumerate(candidates)
    ) or "（无候选定额）"
    messages_r5_analysis = [
        {"role": "system", "content": system_prompt},
        {
            "role": "user",
            "content": (
                f"请先对候选定额进行套定额分析，不要调用工具，不要输出 JSON。\n"
                f"请说明候选取舍、工程量系数、置信度依据和模糊问题，最后给出建议选择的 dekid/dezmid。\n\n"
                f"套定额原则：基于清单项目特征、候选定额名称和工作内容进行匹配分析和判断。\n\n"
                f"【清单项】\n编码：{boq_item['item_code']}\n名称：{boq_item['item_name']}\n"
                f"项目特征：{boq_item.get('item_description') or '（未填写）'}\n单位：{boq_item.get('unit') or '无'}\n"
                f"【编码核查】{_json_dumps(code_check)}\n"
                f"【项目特征分析】{_json_dumps(feature_result)}\n"
                f"【章节规则校验】{_json_dumps(chapter_rule_check)}\n"
                f"【候选定额子目（共 {candidates_data['total']} 条）】\n{candidate_text}\n"
            ),
        },
    ]
    quota_analysis = ""
    if candidates:
        yield ("reasoning_token", "\n\n[第五轮A：套定额分析]\n")
        for event_type, data in _stream_text_completion(messages_r5_analysis, 6000):
            if event_type == "reasoning_token":
                yield ("reasoning_token", data)
            else:
                quota_analysis = data

    messages_r5_submit = [
        {"role": "system", "content": system_prompt},
        {
            "role": "user",
            "content": (
                f"请根据以下套定额分析和候选定额，严格调用 submit_quota_match 提交最终结构化结果。\n"
                f"只能提交候选中的 dekid/dezmid，不允许编造候选外子目。\n\n"
                f"套定额原则：用清单项目特征信息、候选定额名称和工作内容进行匹配分析和判断。\n\n"
                f"【清单项】\n编码：{boq_item['item_code']}\n名称：{boq_item['item_name']}\n"
                f"项目特征：{boq_item.get('item_description') or '（未填写）'}\n单位：{boq_item.get('unit') or '无'}\n"
                f"【编码核查】{_json_dumps(code_check)}\n"
                f"【项目特征分析】{_json_dumps(feature_result)}\n"
                f"【章节规则校验】{_json_dumps(chapter_rule_check)}\n"
                f"【第五轮A套定额分析】\n{quota_analysis or '（无分析文本）'}\n\n"
                f"【候选定额子目（共 {candidates_data['total']} 条）】\n{candidate_text}\n"
            ),
        },
    ]
    raw_match = _run_submit_match(messages_r5_submit) if candidates else {"matches": [], "issues": ["未找到候选定额子目"]}
    match_result = _normalize_matches(raw_match, candidates)
    try:
        chapter_rule_check["validation"] = _validate_chapter_rules(chapter_rule_check, match_result["matches"])
    except ModelRateLimitError:
        raise
    except Exception as exc:
        chapter_rule_check["validation"] = {
            "validations": [],
            "issues": [f"章节规则最终校验失败：{exc}"],
            "status": "manual_review",
        }
    yield ("chapter_rule_check", chapter_rule_check)
    if persist_run and run_id is not None:
        _update_run(conn, run_id, quota_match=match_result)
        _update_run(conn, run_id, chapter_rule_check=chapter_rule_check)
    yield ("quota_match", match_result)
    yield ("step_timing", _finish_step_timing(conn, run_id, step_timings, 5, "套定额结果", step_started_at, step_started_perf))

    step_started_at = datetime.now()
    step_started_perf = perf_counter()
    manual = _manual_quotas(conn, manual_project_id, boq_item.get("item_code"))
    evaluation = _evaluate(match_result["matches"], manual)
    if persist_run and run_id is not None:
        _update_run(conn, run_id, evaluation=evaluation)
        _save_pending_results(conn, task_id, run_id, boq_item, match_result, evaluation)
    yield ("evaluation", evaluation)
    yield ("step_timing", _finish_step_timing(conn, run_id, step_timings, 6, "人工对比", step_started_at, step_started_perf))


def _row_to_task(row) -> dict[str, Any]:
    ids = row[6] or []
    return {
        "id": row[0],
        "name": row[1],
        "boq_project_id": row[2],
        "project_id": row[2],
        "project_name": row[3],
        "manual_project_id": row[4],
        "manual_project_name": row[5],
        "quota_library_ids": ids,
        "quota_library_names": row[7] or [],
        "legacy_local_id": row[8],
        "created_at": row[9],
        "latest_run_count": row[10],
        "accuracy_report": row[11] if len(row) > 11 else None,
        "kb_version_id": int(row[12]) if len(row) > 12 and row[12] is not None else None,
        "consistency_rate": float(row[13]) if len(row) > 13 and row[13] is not None else None,
    }


def _row_to_batch(row) -> dict[str, Any]:
    return {
        "id": row[0],
        "name": row[1],
        "boq_project_id": row[2],
        "project_id": row[2],
        "project_name": row[3],
        "manual_project_id": row[4],
        "manual_project_name": row[5],
        "quota_library_ids": row[6] or [],
        "quota_library_names": row[7] or [],
        "status": row[8],
        "selected_count": row[9] or 0,
        "completed_count": row[10] or 0,
        "failed_count": row[11] or 0,
        "created_at": row[12],
        "started_at": row[13],
        "finished_at": row[14],
        "kb_version_id": int(row[15]) if len(row) > 15 and row[15] is not None else None,
        "consistency_rate": float(row[16]) if len(row) > 16 and row[16] is not None else None,
        "concurrency_limit": int(row[17]) if len(row) > 17 and row[17] is not None else 20,
    }


_BACKGROUND_COMPLETED_STATUSES = {"confirmed", "completed", "no_match"}


def _background_execution_status(status: str | None) -> str:
    if not status or status in {"idle", "queued"}:
        return "waiting"
    if status == "running":
        return "processing"
    if status == "retrying":
        return "retrying"
    if status == "failed":
        return "failed"
    if status in _BACKGROUND_COMPLETED_STATUSES:
        return "completed"
    return status


def _background_consistency_status(
    evaluation: dict[str, Any] | None = None,
    *,
    hit_count: int | None = None,
    missed_count: int | None = None,
    extra_count: int | None = None,
    manual_count: int | None = None,
    ai_count: int | None = None,
) -> str:
    if evaluation is not None:
        hit_count = int(evaluation.get("hit_count") or 0)
        missed_count = int(evaluation.get("missed_count") or 0)
        extra_count = int(evaluation.get("extra_count") or 0)
        manual_count = int(evaluation.get("manual_count") or 0)
        ai_count = int(evaluation.get("ai_count") or 0)
    if None in {hit_count, missed_count, extra_count, manual_count, ai_count}:
        return "unassessed"
    if manual_count == 0 and ai_count == 0:
        return "no_comparable"
    if missed_count == 0 and extra_count == 0:
        return "exact"
    if hit_count and hit_count > 0:
        return "partial"
    return "inconsistent"


def _require_task_run_owner(conn, user: CurrentUser, run_id: int) -> None:
    with conn.cursor() as cur:
        cur.execute("SELECT task_id FROM pricing_task_runs WHERE id=%s", (run_id,))
        row = cur.fetchone()
    if not row or not row[0]:
        raise HTTPException(status_code=404, detail="任务运行不存在")
    require_task_owner(conn, user, int(row[0]))


def _require_batch_item_run_owner(conn, user: CurrentUser, run_id: int) -> None:
    with conn.cursor() as cur:
        cur.execute("SELECT batch_id FROM pricing_task_batch_item_runs WHERE id=%s", (run_id,))
        row = cur.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="批次运行不存在")
    require_task_owner(conn, user, int(row[0]), batch=True)


@router.get("/pricing-tasks")
def list_pricing_tasks(user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT t.id, t.name, t.boq_project_id, p.project_name, t.manual_project_id, mp.project_name AS manual_project_name,
                       t.quota_library_ids, COALESCE(array_agg(l.mc ORDER BY l.id) FILTER (WHERE l.id IS NOT NULL), '{}') AS library_names,
                       t.legacy_local_id, t.created_at,
                       (SELECT COUNT(*) FROM pricing_task_runs r WHERE r.task_id=t.id) AS run_count
                       , t.accuracy_report, t.kb_version_id,
                       (
                           SELECT ROUND(
                               SUM(COALESCE((latest.evaluation->>'hit_count')::numeric, 0))
                               / NULLIF(SUM(COALESCE((latest.evaluation->>'manual_count')::numeric, 0)), 0),
                               4
                           )
                           FROM (
                               SELECT DISTINCT ON (r.boq_item_id) r.evaluation
                               FROM pricing_task_runs r
                               WHERE r.task_id=t.id AND r.evaluation IS NOT NULL
                               ORDER BY r.boq_item_id, r.created_at DESC, r.id DESC
                           ) latest
                       ) AS consistency_rate
                FROM pricing_tasks t
                JOIN boq_projects p ON p.id = t.boq_project_id
                LEFT JOIN manual_boq_projects mp ON mp.id = t.manual_project_id
                LEFT JOIN LATERAL jsonb_array_elements_text(t.quota_library_ids) lib_id(value) ON TRUE
                LEFT JOIN tlibs l ON l.kb_version_id=pricing_kb_data_version(t.kb_version_id,'TLibs') AND l.id=lib_id.value::bigint
                WHERE t.status <> 'deleted' AND (%s OR t.owner_user_id=%s)
                GROUP BY t.id, p.project_name, mp.project_name
                ORDER BY t.created_at DESC
                """,
                (user.is_admin, user.id),
            )
            return [_row_to_task(row) for row in cur.fetchall()]
    finally:
        conn.close()


@router.post("/pricing-tasks")
def create_pricing_task(body: PricingTaskCreate, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            try:
                kb_version_id = resolve_version_id(conn, body.kb_version_id)
            except (ValueError, RuntimeError) as exc:
                raise HTTPException(status_code=409, detail=str(exc)) from exc
            require_project_owner(conn, user, body.boq_project_id)
            cur.execute("SELECT project_name FROM boq_projects WHERE id=%s", (body.boq_project_id,))
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="BOQ project not found")
            if body.manual_project_id:
                cur.execute("SELECT id FROM manual_boq_projects WHERE id=%s", (body.manual_project_id,))
                if not cur.fetchone():
                    raise HTTPException(status_code=404, detail="manual project not found")
            cur.execute(
                """
                INSERT INTO pricing_tasks(name, boq_project_id, quota_library_ids, manual_project_id, legacy_local_id, kb_version_id, owner_user_id)
                VALUES (%s, %s, %s::jsonb, %s, %s, %s, %s)
                ON CONFLICT (legacy_local_id) DO UPDATE SET updated_at=NOW()
                RETURNING id
                """,
                (
                    body.name.strip(),
                    body.boq_project_id,
                    json.dumps(body.quota_library_ids or []),
                    body.manual_project_id,
                    body.legacy_local_id,
                    kb_version_id,
                    user.id,
                ),
            )
            task_id = cur.fetchone()[0]
        conn.commit()
        return {"id": task_id, "kb_version_id": kb_version_id}
    finally:
        conn.close()


@router.post("/pricing-tasks/import-local")
def import_local_tasks(body: PricingTaskImportRequest, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    imported = []
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            try:
                kb_version_id = resolve_version_id(conn, None)
            except (ValueError, RuntimeError) as exc:
                raise HTTPException(status_code=409, detail=str(exc)) from exc
            for item in body.tasks:
                require_project_owner(conn, user, item.project_id)
                cur.execute(
                    """
                    INSERT INTO pricing_tasks(name, boq_project_id, quota_library_ids, manual_project_id, legacy_local_id, kb_version_id, owner_user_id)
                    VALUES (%s, %s, '[]'::jsonb, %s, %s, %s, %s)
                    ON CONFLICT (legacy_local_id) DO UPDATE SET updated_at=NOW()
                    RETURNING id
                    """,
                    (item.name, item.project_id, item.manual_project_id, item.id, kb_version_id, user.id),
                )
                imported.append({"legacy_local_id": item.id, "id": cur.fetchone()[0]})
        conn.commit()
        return {"imported": imported}
    finally:
        conn.close()


@router.get("/pricing-tasks/{task_id}")
def get_pricing_task(task_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, task_id)
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT t.id, t.name, t.boq_project_id, p.project_name, t.manual_project_id, mp.project_name AS manual_project_name,
                       t.quota_library_ids, COALESCE(array_agg(l.mc ORDER BY l.id) FILTER (WHERE l.id IS NOT NULL), '{}') AS library_names,
                       t.legacy_local_id, t.created_at,
                       (SELECT COUNT(*) FROM pricing_task_runs r WHERE r.task_id=t.id) AS run_count
                       , t.accuracy_report, t.kb_version_id,
                       (
                           SELECT ROUND(
                               SUM(COALESCE((latest.evaluation->>'hit_count')::numeric, 0))
                               / NULLIF(SUM(COALESCE((latest.evaluation->>'manual_count')::numeric, 0)), 0),
                               4
                           )
                           FROM (
                               SELECT DISTINCT ON (r.boq_item_id) r.evaluation
                               FROM pricing_task_runs r
                               WHERE r.task_id=t.id AND r.evaluation IS NOT NULL
                               ORDER BY r.boq_item_id, r.created_at DESC, r.id DESC
                           ) latest
                       ) AS consistency_rate
                FROM pricing_tasks t
                JOIN boq_projects p ON p.id = t.boq_project_id
                LEFT JOIN manual_boq_projects mp ON mp.id = t.manual_project_id
                LEFT JOIN LATERAL jsonb_array_elements_text(t.quota_library_ids) lib_id(value) ON TRUE
                LEFT JOIN tlibs l ON l.kb_version_id=pricing_kb_data_version(t.kb_version_id,'TLibs') AND l.id=lib_id.value::bigint
                WHERE t.id=%s AND t.status <> 'deleted'
                GROUP BY t.id, p.project_name, mp.project_name
                """,
                (task_id,),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="task not found")
            return _row_to_task(row)
    finally:
        conn.close()


@router.delete("/pricing-tasks/{task_id}", status_code=204)
def delete_pricing_task(task_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, task_id)
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE pricing_tasks SET status='deleted', updated_at=NOW() WHERE id=%s",
                (task_id,),
            )
        conn.commit()
    finally:
        conn.close()


def _batch_select_sql() -> str:
    return """
        SELECT b.id, b.name, b.boq_project_id, p.project_name,
               b.manual_project_id, mp.project_name AS manual_project_name,
               b.quota_library_ids,
               COALESCE(array_agg(l.mc ORDER BY l.id) FILTER (WHERE l.id IS NOT NULL), '{}') AS library_names,
               b.status, b.selected_count, b.completed_count, b.failed_count,
               b.created_at, b.started_at, b.finished_at, b.kb_version_id,
               (
                   SELECT ROUND(
                       SUM(COALESCE((r.evaluation->>'hit_count')::numeric, 0))
                       / NULLIF(SUM(COALESCE((r.evaluation->>'manual_count')::numeric, 0)), 0),
                       4
                   )
                   FROM pricing_task_batch_item_runs r
                   WHERE r.batch_id=b.id AND r.evaluation IS NOT NULL
               ) AS consistency_rate,
               b.concurrency_limit
        FROM pricing_task_batches b
        JOIN boq_projects p ON p.id = b.boq_project_id
        LEFT JOIN manual_boq_projects mp ON mp.id = b.manual_project_id
        LEFT JOIN LATERAL jsonb_array_elements_text(b.quota_library_ids) lib_id(value) ON TRUE
        LEFT JOIN tlibs l ON l.kb_version_id=pricing_kb_data_version(b.kb_version_id,'TLibs') AND l.id=lib_id.value::bigint
    """


@router.get("/pricing-task-batches")
def list_pricing_task_batches(user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                _batch_select_sql()
                + """
                WHERE b.status <> 'deleted' AND b.workflow_version='legacy' AND (%s OR b.owner_user_id=%s)
                GROUP BY b.id, p.project_name, mp.project_name
                ORDER BY b.created_at DESC
                """,
                (user.is_admin, user.id),
            )
            return [_row_to_batch(row) for row in cur.fetchall()]
    finally:
        conn.close()


@router.get("/pricing-task-new-batches")
def list_new_pricing_task_batches(user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                _batch_select_sql()
                + """
                WHERE b.status <> 'deleted' AND b.workflow_version='streamlined' AND (%s OR b.owner_user_id=%s)
                GROUP BY b.id, p.project_name, mp.project_name
                ORDER BY b.created_at DESC
                """,
                (user.is_admin, user.id),
            )
            return [_row_to_batch(row) for row in cur.fetchall()]
    finally:
        conn.close()


@router.get("/pricing-task-background-batches")
def list_background_pricing_task_batches(user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                _batch_select_sql()
                + """
                WHERE b.status <> 'deleted' AND b.workflow_version='background' AND (%s OR b.owner_user_id=%s)
                GROUP BY b.id, p.project_name, mp.project_name
                ORDER BY b.created_at DESC
                """,
                (user.is_admin, user.id),
            )
            return [_row_to_batch(row) for row in cur.fetchall()]
    finally:
        conn.close()


def _create_pricing_task_batch_record(
    body: PricingTaskBatchCreate | NewPricingTaskBatchCreate,
    workflow_version: str,
    *,
    require_manual_project: bool,
    user: CurrentUser,
    concurrency_limit: int | None = None,
):
    from db.connection import get_connection

    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="batch name required")
    if require_manual_project and not body.manual_project_id:
        raise HTTPException(status_code=400, detail="manual comparison project required")
    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            try:
                kb_version_id = resolve_version_id(conn, body.kb_version_id)
            except (ValueError, RuntimeError) as exc:
                raise HTTPException(status_code=409, detail=str(exc)) from exc
            require_project_owner(conn, user, body.boq_project_id)
            cur.execute("SELECT id FROM boq_projects WHERE id=%s", (body.boq_project_id,))
            if not cur.fetchone():
                raise HTTPException(status_code=404, detail="BOQ project not found")
            if body.manual_project_id:
                cur.execute("SELECT id FROM manual_boq_projects WHERE id=%s", (body.manual_project_id,))
                if not cur.fetchone():
                    raise HTTPException(status_code=404, detail="manual project not found")
            cur.execute(
                """
                INSERT INTO pricing_task_batches(
                    name, boq_project_id, quota_library_ids, manual_project_id,
                    kb_version_id, workflow_version, owner_user_id, concurrency_limit
                )
                VALUES (%s, %s, %s::jsonb, %s, %s, %s, %s, %s)
                RETURNING id
                """,
                (
                    name,
                    body.boq_project_id,
                    json.dumps(body.quota_library_ids or []),
                    body.manual_project_id,
                    kb_version_id,
                    workflow_version,
                    user.id,
                    concurrency_limit or 20,
                ),
            )
            batch_id = cur.fetchone()[0]
        conn.commit()
        return {"id": batch_id, "kb_version_id": kb_version_id}
    finally:
        conn.close()


@router.post("/pricing-task-batches")
def create_pricing_task_batch(body: PricingTaskBatchCreate, user: CurrentUser = Depends(current_user)):
    return _create_pricing_task_batch_record(
        body,
        "legacy",
        require_manual_project=False,
        user=user,
    )


@router.post("/pricing-task-new-batches")
def create_new_pricing_task_batch(body: NewPricingTaskBatchCreate, user: CurrentUser = Depends(current_user)):
    return _create_pricing_task_batch_record(
        body,
        "streamlined",
        require_manual_project=True,
        user=user,
    )


@router.post("/pricing-task-background-batches")
def create_background_pricing_task_batch(body: NewPricingTaskBatchCreate, user: CurrentUser = Depends(current_user)):
    return _create_pricing_task_batch_record(
        body,
        "background",
        require_manual_project=True,
        user=user,
        concurrency_limit=_background_worker_limit(),
    )


@router.get("/pricing-task-batches/{batch_id}")
def get_pricing_task_batch(batch_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, batch_id, batch=True)
        with conn.cursor() as cur:
            cur.execute(
                _batch_select_sql()
                + """
                WHERE b.id=%s AND b.status <> 'deleted'
                GROUP BY b.id, p.project_name, mp.project_name
                """,
                (batch_id,),
            )
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="batch not found")
            return _row_to_batch(row)
    finally:
        conn.close()


@router.delete("/pricing-task-batches/{batch_id}", status_code=204)
def delete_pricing_task_batch(batch_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, batch_id, batch=True)
        with conn.cursor() as cur:
            cur.execute("UPDATE pricing_task_batches SET status='deleted', updated_at=NOW() WHERE id=%s", (batch_id,))
        conn.commit()
    finally:
        conn.close()


@router.get("/pricing-task-batches/{batch_id}/items")
def get_pricing_task_batch_items(batch_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        batch = get_pricing_task_batch(batch_id, user)
        project_id = int(batch["boq_project_id"])
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, item_code, item_name, item_description, unit, quantity, item_seq
                FROM boq_items
                WHERE project_id=%s
                ORDER BY item_seq NULLS LAST, id
                """,
                (project_id,),
            )
            items = [
                {
                    "id": r[0],
                    "item_code": r[1],
                    "item_name": r[2],
                    "item_description": r[3],
                    "unit": r[4],
                    "quantity": float(r[5]) if r[5] is not None else None,
                    "item_seq": r[6],
                    "project_id": project_id,
                }
                for r in cur.fetchall()
            ]
            cur.execute(
                """
                SELECT boq_item_id, id, status, code_check, feature_check, chapter_rule_check, work_procedures,
                       quota_candidates, quota_match, evaluation, conversion_check,
                       coefficient_check, step_timings, error_message, created_at, finished_at,
                       reasoning_text, confirmed_results, kb_version_id, execution_id, attempt_count,
                       current_tool, current_tool_status, current_tool_output, updated_at
                FROM pricing_task_batch_item_runs
                WHERE batch_id=%s
                ORDER BY created_at DESC, id DESC
                """,
                (batch_id,),
            )
            seen: set[int] = set()
            runs = []
            for r in cur.fetchall():
                item_id = int(r[0])
                if item_id in seen:
                    continue
                seen.add(item_id)
                runs.append(
                    {
                        "boq_item_id": item_id,
                        "run": {
                            "id": r[1],
                            "status": r[2],
                            "code_check": r[3],
                            "feature_check": r[4],
                            "chapter_rule_check": r[5],
                            "work_procedures": r[6],
                            "quota_candidates": r[7],
                            "quota_match": r[8],
                            "evaluation": r[9],
                            "conversion_check": _hydrate_conversion_for_batch_run(conn, int(r[1]), r[10]),
                            "coefficient_check": r[11],
                            "step_timings": r[12],
                            "error_message": r[13],
                            "created_at": r[14],
                            "finished_at": r[15],
                            "reasoning_text": r[16],
                            "confirmed_results": r[17] or [],
                            "kb_version_id": int(r[18]) if r[18] is not None else None,
                            "execution_id": int(r[19]) if r[19] is not None else None,
                            "attempt_count": int(r[20] or 0),
                            "current_tool": r[21],
                            "current_tool_status": r[22],
                            "current_tool_output": r[23],
                            "updated_at": r[24],
                        },
                    }
                )
            execution = None
            if batch.get("status") in {"queued", "running", "stop_requested"} or batch.get("concurrency_limit"):
                cur.execute(
                    """SELECT id,status,concurrency_limit,selected_count,completed_count,failed_count,
                              created_at,started_at,finished_at
                       FROM pricing_task_batch_executions WHERE batch_id=%s
                       ORDER BY id DESC LIMIT 1""",
                    (batch_id,),
                )
                erow = cur.fetchone()
                if erow:
                    execution = {
                        "id": int(erow[0]), "status": erow[1], "concurrency_limit": int(erow[2]),
                        "selected_count": int(erow[3]), "completed_count": int(erow[4]),
                        "failed_count": int(erow[5]), "created_at": erow[6], "started_at": erow[7],
                        "finished_at": erow[8],
                    }
        return {"batch": batch, "items": items, "runs": runs, "execution": execution}
    finally:
        conn.close()


def _background_event(conn, batch_id: int, execution_id: int | None, item_run_id: int | None,
                      boq_item_id: int | None, event_type: str, payload: dict[str, Any], *, commit: bool = True) -> int:
    with conn.cursor() as cur:
        cur.execute(
            """INSERT INTO pricing_task_batch_events(
                   batch_id,execution_id,item_run_id,boq_item_id,event_type,payload
               ) VALUES(%s,%s,%s,%s,%s,%s) RETURNING id""",
            (batch_id, execution_id, item_run_id, boq_item_id, event_type, Json(payload, dumps=_json_dumps)),
        )
        event_id = int(cur.fetchone()[0])
    if commit:
        conn.commit()
    return event_id


@router.post("/pricing-task-background-batches/{batch_id}/executions")
def start_background_batch_execution(
    batch_id: int,
    body: BackgroundBatchExecutionCreate,
    user: CurrentUser = Depends(current_user),
):
    from db.connection import get_connection

    item_ids = list(dict.fromkeys(body.boq_item_ids))
    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, batch_id, batch=True)
        with conn.cursor() as cur:
            cur.execute(
                """SELECT boq_project_id,kb_version_id,concurrency_limit,workflow_version
                   FROM pricing_task_batches WHERE id=%s AND status<>'deleted' FOR UPDATE""",
                (batch_id,),
            )
            batch_row = cur.fetchone()
            if not batch_row or batch_row[3] != "background":
                raise HTTPException(status_code=404, detail="后台批次不存在")
            cur.execute(
                """SELECT 1 FROM pricing_task_batch_executions
                   WHERE batch_id=%s AND status IN ('queued','running','stop_requested') LIMIT 1""",
                (batch_id,),
            )
            if cur.fetchone():
                raise HTTPException(status_code=409, detail="当前批次已有正在执行的后台任务")
            cur.execute(
                "SELECT id FROM boq_items WHERE project_id=%s AND id=ANY(%s)",
                (batch_row[0], item_ids),
            )
            valid_ids = {int(row[0]) for row in cur.fetchall()}
            if valid_ids != set(item_ids):
                raise HTTPException(status_code=400, detail="部分清单不属于当前批次工程")
            cur.execute(
                """INSERT INTO pricing_task_batch_executions(
                       batch_id,status,concurrency_limit,selected_count
                   ) VALUES(%s,'queued',%s,%s) RETURNING id""",
                (batch_id, int(batch_row[2]), len(item_ids)),
            )
            execution_id = int(cur.fetchone()[0])
            for item_id in item_ids:
                cur.execute(
                    """INSERT INTO pricing_task_batch_item_runs(
                           batch_id,boq_item_id,boq_project_id,kb_version_id,status,execution_id,
                           attempt_count,step_timings,confirmed_results,updated_at
                       ) VALUES(%s,%s,%s,%s,'queued',%s,0,'{}'::jsonb,'[]'::jsonb,NOW())
                       ON CONFLICT(batch_id,boq_item_id) DO UPDATE SET
                           status='queued',execution_id=EXCLUDED.execution_id,attempt_count=0,
                           reasoning_text='',code_check=NULL,feature_check=NULL,chapter_rule_check=NULL,
                           work_procedures=NULL,quota_candidates=NULL,quota_match=NULL,evaluation=NULL,
                           confirmed_results='[]'::jsonb,conversion_check=NULL,coefficient_check=NULL,
                           step_timings='{}'::jsonb,error_message=NULL,current_tool=NULL,
                           current_tool_status=NULL,current_tool_output=NULL,next_attempt_at=NULL,
                           lease_owner=NULL,lease_expires_at=NULL,started_at=NULL,finished_at=NULL,updated_at=NOW()""",
                    (batch_id, item_id, batch_row[0], batch_row[1], execution_id),
                )
            cur.execute(
                """UPDATE pricing_task_batches SET status='queued',selected_count=(
                       SELECT COUNT(*) FROM pricing_task_batch_item_runs WHERE batch_id=%s
                   ),completed_count=(SELECT COUNT(*) FROM pricing_task_batch_item_runs
                       WHERE batch_id=%s AND status IN ('confirmed','completed','no_match')),
                   failed_count=(SELECT COUNT(*) FROM pricing_task_batch_item_runs WHERE batch_id=%s AND status='failed'),
                   started_at=COALESCE(started_at,NOW()),finished_at=NULL,updated_at=NOW() WHERE id=%s""",
                (batch_id, batch_id, batch_id, batch_id),
            )
        conn.commit()
        _background_event(conn, batch_id, execution_id, None, None, "execution_queued", {"selected_count": len(item_ids)})
        return {"id": execution_id, "status": "queued", "concurrency_limit": int(batch_row[2])}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _background_workspace_summary(
    items: list[dict[str, Any]],
    execution: dict[str, Any] | None,
) -> dict[str, Any]:
    status_counts = {key: 0 for key in ("waiting", "processing", "retrying", "completed", "failed")}
    consistency_counts = {
        key: 0 for key in ("exact", "partial", "inconsistent", "no_comparable", "unassessed")
    }
    hit_count = missed_count = extra_count = manual_count = ai_count = 0
    for item in items:
        status_counts[item["execution_status"]] = status_counts.get(item["execution_status"], 0) + 1
        consistency_counts[item["consistency_status"]] = consistency_counts.get(item["consistency_status"], 0) + 1
        hit_count += item["hit_count"]
        missed_count += item["missed_count"]
        extra_count += item["extra_count"]
        manual_count += item["manual_count"]
        ai_count += item["ai_count"]

    elapsed_seconds = 0.0
    if execution and execution.get("started_at"):
        finish = execution.get("finished_at") or datetime.now()
        elapsed_seconds = max(0.0, (finish - execution["started_at"]).total_seconds())
    completed_count = status_counts["completed"]
    throughput = completed_count / (elapsed_seconds / 60) if elapsed_seconds > 0 and completed_count else 0.0
    remaining = status_counts["waiting"] + status_counts["processing"] + status_counts["retrying"]
    eta_seconds = remaining / (throughput / 60) if throughput > 0 and remaining else None
    evaluated_count = len(items) - consistency_counts["unassessed"]
    return {
        "total_count": len(items),
        "selected_count": int(execution.get("selected_count") or 0) if execution else 0,
        "waiting_count": status_counts["waiting"],
        "running_count": status_counts["processing"],
        "retrying_count": status_counts["retrying"],
        "completed_count": completed_count,
        "failed_count": status_counts["failed"],
        "evaluated_count": evaluated_count,
        "exact_count": consistency_counts["exact"],
        "partial_count": consistency_counts["partial"],
        "inconsistent_count": consistency_counts["inconsistent"],
        "no_comparable_count": consistency_counts["no_comparable"],
        "unassessed_count": consistency_counts["unassessed"],
        "hit_count": hit_count,
        "missed_count": missed_count,
        "extra_count": extra_count,
        "manual_count": manual_count,
        "ai_count": ai_count,
        "hit_rate": round(hit_count / manual_count, 4) if manual_count else None,
        "elapsed_seconds": round(elapsed_seconds, 1),
        "throughput_per_minute": round(throughput, 2),
        "eta_seconds": round(eta_seconds, 1) if eta_seconds is not None else None,
    }


@router.get("/pricing-task-background-batches/{batch_id}/workspace")
def get_background_pricing_task_workspace(batch_id: int, user: CurrentUser = Depends(current_user)):
    """Return a compact monitoring projection; full tool payloads are loaded per focused item."""
    from db.connection import get_connection, get_connection_pool_metrics

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, batch_id, batch=True)
        with conn.cursor() as cur:
            # Keep the dashboard projection on one short pooled lease. Calling
            # the public batch endpoint here used to create a nested connection.
            cur.execute(
                _batch_select_sql()
                + """
                WHERE b.id=%s AND b.status <> 'deleted'
                GROUP BY b.id, p.project_name, mp.project_name
                """,
                (batch_id,),
            )
            batch_row = cur.fetchone()
            if not batch_row:
                raise HTTPException(status_code=404, detail="批次不存在")
            batch = _row_to_batch(batch_row)
            cur.execute("SELECT workflow_version FROM pricing_task_batches WHERE id=%s", (batch_id,))
            workflow_row = cur.fetchone()
            if not workflow_row or workflow_row[0] != "background":
                raise HTTPException(status_code=404, detail="后台批次不存在")
            cur.execute(
                """
                SELECT id,status,concurrency_limit,selected_count,completed_count,failed_count,
                       created_at,started_at,finished_at
                FROM pricing_task_batch_executions
                WHERE batch_id=%s ORDER BY id DESC LIMIT 1
                """,
                (batch_id,),
            )
            erow = cur.fetchone()
            execution = None
            if erow:
                execution = {
                    "id": int(erow[0]), "status": erow[1], "concurrency_limit": int(erow[2]),
                    "selected_count": int(erow[3]), "completed_count": int(erow[4]),
                    "failed_count": int(erow[5]), "created_at": erow[6], "started_at": erow[7],
                    "finished_at": erow[8],
                }
            cur.execute(
                """
                SELECT i.id,i.item_code,i.item_name,i.item_description,i.unit,i.quantity,i.item_seq,
                       r.id,r.status,r.execution_id,r.attempt_count,r.current_tool,r.current_tool_status,
                       r.current_tool_output,r.error_message,
                       COALESCE(
                           (r.quota_candidates->>'total')::integer,
                           CASE WHEN jsonb_typeof(r.quota_candidates->'candidates')='array'
                                THEN jsonb_array_length(r.quota_candidates->'candidates') END,
                           0
                       ),
                       CASE WHEN jsonb_typeof(r.quota_match->'matches')='array'
                            THEN jsonb_array_length(r.quota_match->'matches') ELSE 0 END,
                       r.evaluation IS NOT NULL,
                       COALESCE((r.evaluation->>'hit_count')::integer,0),
                       COALESCE((r.evaluation->>'missed_count')::integer,0),
                       COALESCE((r.evaluation->>'extra_count')::integer,0),
                       COALESCE((r.evaluation->>'manual_count')::integer,0),
                       COALESCE((r.evaluation->>'ai_count')::integer,0),
                       r.started_at,r.finished_at,r.updated_at,r.next_attempt_at,r.kb_version_id
                FROM boq_items i
                LEFT JOIN LATERAL (
                    SELECT * FROM pricing_task_batch_item_runs latest
                    WHERE latest.batch_id=%s AND latest.boq_item_id=i.id
                    ORDER BY latest.created_at DESC,latest.id DESC LIMIT 1
                ) r ON TRUE
                WHERE i.project_id=%s
                ORDER BY i.item_seq NULLS LAST,i.id
                """,
                (batch_id, int(batch["boq_project_id"])),
            )
            items = []
            now = datetime.now()
            for row in cur.fetchall():
                has_evaluation = bool(row[17])
                consistency_status = _background_consistency_status(
                    hit_count=int(row[18]), missed_count=int(row[19]), extra_count=int(row[20]),
                    manual_count=int(row[21]), ai_count=int(row[22]),
                ) if has_evaluation else "unassessed"
                end_at = row[24] or (now if row[8] in {"running", "retrying"} else None)
                duration_ms = None
                if row[23] and end_at:
                    duration_ms = max(0, int((end_at - row[23]).total_seconds() * 1000))
                items.append({
                    "id": int(row[0]), "item_code": row[1] or "", "item_name": row[2] or "",
                    "item_description": row[3] or "", "unit": row[4] or "",
                    "quantity": float(row[5]) if row[5] is not None else None, "item_seq": row[6],
                    "run_id": int(row[7]) if row[7] is not None else None,
                    "run_status": row[8] or "idle", "execution_status": _background_execution_status(row[8]),
                    "execution_id": int(row[9]) if row[9] is not None else None,
                    "attempt_count": int(row[10] or 0), "current_tool": row[11],
                    "current_tool_status": row[12], "current_tool_output": row[13],
                    "error_message": row[14], "candidate_count": int(row[15] or 0),
                    "match_count": int(row[16] or 0), "hit_count": int(row[18]),
                    "missed_count": int(row[19]), "extra_count": int(row[20]),
                    "manual_count": int(row[21]), "ai_count": int(row[22]),
                    "consistency_status": consistency_status, "duration_ms": duration_ms,
                    "updated_at": row[25], "next_attempt_at": row[26],
                    "kb_version_id": int(row[27]) if row[27] is not None else None,
                })
            cur.execute(
                """SELECT COUNT(*) FROM pricing_task_batch_item_runs
                   WHERE batch_id=%s AND status='failed'
                     AND error_message ILIKE '%%模型服务限流或配额不足%%'""",
                (batch_id,),
            )
            model_rate_limited_failed_count = int(cur.fetchone()[0] or 0)
        return {
            "batch": batch,
            "execution": execution,
            "summary": _background_workspace_summary(items, execution),
            "runtime_metrics": {
                "model": get_model_gate_metrics(),
                "database": get_connection_pool_metrics(),
                "model_rate_limited_failed_count": model_rate_limited_failed_count,
            },
            "items": items,
        }
    finally:
        conn.close()


@router.get("/pricing-task-background-batches/{batch_id}/items/{boq_item_id}")
def get_background_pricing_task_item_detail(
    batch_id: int,
    boq_item_id: int,
    user: CurrentUser = Depends(current_user),
):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, batch_id, batch=True)
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT b.workflow_version,b.boq_project_id,i.id,i.item_code,i.item_name,i.item_description,
                       i.unit,i.quantity,i.item_seq
                FROM pricing_task_batches b
                JOIN boq_items i ON i.project_id=b.boq_project_id AND i.id=%s
                WHERE b.id=%s AND b.status<>'deleted'
                """,
                (boq_item_id, batch_id),
            )
            item_row = cur.fetchone()
            if not item_row or item_row[0] != "background":
                raise HTTPException(status_code=404, detail="后台批次清单不存在")
            item = {
                "id": int(item_row[2]), "project_id": int(item_row[1]), "item_code": item_row[3] or "",
                "item_name": item_row[4] or "", "item_description": item_row[5] or "",
                "unit": item_row[6] or "", "quantity": float(item_row[7]) if item_row[7] is not None else None,
                "item_seq": item_row[8],
            }
            cur.execute(
                """
                SELECT id,status,code_check,feature_check,chapter_rule_check,work_procedures,
                       quota_candidates,quota_match,evaluation,conversion_check,coefficient_check,
                       step_timings,error_message,created_at,started_at,finished_at,confirmed_results,
                       kb_version_id,execution_id,attempt_count,current_tool,current_tool_status,
                       current_tool_output,updated_at,next_attempt_at
                FROM pricing_task_batch_item_runs
                WHERE batch_id=%s AND boq_item_id=%s
                ORDER BY created_at DESC,id DESC LIMIT 1
                """,
                (batch_id, boq_item_id),
            )
            row = cur.fetchone()
            run = None
            if row:
                run = {
                    "id": int(row[0]), "status": row[1], "code_check": row[2], "feature_check": row[3],
                    "chapter_rule_check": row[4], "work_procedures": row[5], "quota_candidates": row[6],
                    "quota_match": row[7], "evaluation": row[8],
                    "conversion_check": _hydrate_conversion_for_batch_run(conn, int(row[0]), row[9]),
                    "coefficient_check": row[10], "step_timings": row[11], "error_message": row[12],
                    "created_at": row[13], "started_at": row[14], "finished_at": row[15],
                    "confirmed_results": row[16] or [],
                    "kb_version_id": int(row[17]) if row[17] is not None else None,
                    "execution_id": int(row[18]) if row[18] is not None else None,
                    "attempt_count": int(row[19] or 0), "current_tool": row[20],
                    "current_tool_status": row[21], "current_tool_output": row[22],
                    "updated_at": row[23], "next_attempt_at": row[24],
                }
        return {"item": item, "run": run}
    finally:
        conn.close()


@router.post("/pricing-task-background-batches/{batch_id}/executions/{execution_id}/stop")
def stop_background_batch_execution(batch_id: int, execution_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, batch_id, batch=True)
        with conn.cursor() as cur:
            cur.execute(
                """UPDATE pricing_task_batch_executions
                   SET status='stop_requested',stop_requested_at=NOW(),updated_at=NOW()
                   WHERE id=%s AND batch_id=%s AND status IN ('queued','running') RETURNING id""",
                (execution_id, batch_id),
            )
            if not cur.fetchone():
                raise HTTPException(status_code=409, detail="当前执行无法停止")
            cur.execute("UPDATE pricing_task_batches SET status='stop_requested',updated_at=NOW() WHERE id=%s", (batch_id,))
        conn.commit()
        _background_event(conn, batch_id, execution_id, None, None, "execution_stop_requested", {})
        return {"id": execution_id, "status": "stop_requested"}
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@router.get("/pricing-task-background-batches/{batch_id}/events-stream")
def stream_background_batch_events(batch_id: int, after_id: int = 0, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    check_conn = get_connection()
    try:
        _ensure_schema(check_conn)
        require_task_owner(check_conn, user, batch_id, batch=True)
    finally:
        check_conn.close()

    def generate():
        cursor_id = max(0, after_id)
        idle_rounds = 0
        while idle_rounds < 30:
            events: list[str] = []
            try:
                rows = _read_background_event_chunk(batch_id, cursor_id)
                if rows:
                    idle_rounds = 0
                    for row in rows:
                        cursor_id = int(row["id"])
                        events.append(_sse(row))
                else:
                    idle_rounds += 1
                    events.append(": keep-alive\n\n")
            except DatabasePoolBusyError:
                idle_rounds += 1
                events.append(_sse({
                    "id": cursor_id,
                    "execution_id": None,
                    "item_run_id": None,
                    "boq_item_id": None,
                    "event_type": "database_busy",
                    "payload": {"retryable": True, "retry_after_seconds": 1},
                    "created_at": datetime.now(),
                }))
            # The physical connection is always returned before streaming any
            # bytes to a potentially slow or disconnected browser.
            yield from events
            sleep(1)

    return StreamingResponse(generate(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


def _read_background_event_chunk(batch_id: int, after_id: int) -> list[dict[str, Any]]:
    """Load and detach one SSE event chunk from its short database lease."""
    from db.connection import get_connection

    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """SELECT id,execution_id,item_run_id,boq_item_id,event_type,payload,created_at
                   FROM pricing_task_batch_events WHERE batch_id=%s AND id>%s ORDER BY id LIMIT 200""",
                (batch_id, after_id),
            )
            rows = cur.fetchall()
        return [
            {
                "id": int(row[0]), "execution_id": row[1], "item_run_id": row[2],
                "boq_item_id": row[3], "event_type": row[4], "payload": row[5] or {},
                "created_at": row[6],
            }
            for row in rows
        ]
    finally:
        conn.close()


def _load_batch_and_item(conn, batch_id: int, boq_item_id: int) -> tuple[dict[str, Any], dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT id, name, boq_project_id, quota_library_ids, manual_project_id, status, kb_version_id
            FROM pricing_task_batches
            WHERE id=%s AND status <> 'deleted'
            """,
            (batch_id,),
        )
        batch_row = cur.fetchone()
        if not batch_row:
            raise HTTPException(status_code=404, detail="batch not found")
        cur.execute(
            """
            SELECT id, item_code, item_name, item_description, unit, quantity, project_id
            FROM boq_items
            WHERE id=%s AND project_id=%s
            """,
            (boq_item_id, batch_row[2]),
        )
        item_row = cur.fetchone()
        if not item_row:
            raise HTTPException(status_code=404, detail="item not found in batch project")
    return (
        {
            "id": batch_row[0],
            "name": batch_row[1],
            "boq_project_id": batch_row[2],
            "quota_library_ids": batch_row[3] or [],
            "manual_project_id": batch_row[4],
            "status": batch_row[5],
            "kb_version_id": int(batch_row[6]),
        },
        {
            "id": item_row[0],
            "item_code": item_row[1],
            "item_name": item_row[2],
            "item_description": item_row[3],
            "unit": item_row[4],
            "quantity": float(item_row[5]) if item_row[5] is not None else None,
            "project_id": item_row[6],
        },
    )


def _create_or_reset_batch_item_run(
    conn, batch_id: int, boq_item: dict[str, Any], kb_version_id: int
) -> int:
    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO pricing_task_batch_item_runs(
                batch_id, boq_item_id, boq_project_id, kb_version_id, status, started_at,
                reasoning_text, code_check, feature_check, chapter_rule_check, work_procedures,
                quota_candidates, quota_match, evaluation, confirmed_results,
                conversion_check, coefficient_check, step_timings, error_message, finished_at
            )
            VALUES (%s, %s, %s, %s, 'running', NOW(), '', NULL, NULL, NULL, NULL, NULL, NULL, NULL, '[]'::jsonb, NULL, NULL, '{}'::jsonb, NULL, NULL)
            ON CONFLICT (batch_id, boq_item_id) DO UPDATE SET
                status='running',
                kb_version_id=EXCLUDED.kb_version_id,
                started_at=NOW(),
                finished_at=NULL,
                reasoning_text='',
                code_check=NULL,
                feature_check=NULL,
                chapter_rule_check=NULL,
                work_procedures=NULL,
                quota_candidates=NULL,
                quota_match=NULL,
                evaluation=NULL,
                confirmed_results='[]'::jsonb,
                conversion_check=NULL,
                coefficient_check=NULL,
                step_timings='{}'::jsonb,
                error_message=NULL
            RETURNING id
            """,
            (batch_id, boq_item["id"], boq_item["project_id"], kb_version_id),
        )
        item_run_id = cur.fetchone()[0]
        cur.execute(
            """
            UPDATE pricing_task_batches
            SET status='running',
                started_at=COALESCE(started_at, NOW()),
                selected_count=(SELECT COUNT(*) FROM pricing_task_batch_item_runs WHERE batch_id=%s),
                updated_at=NOW()
            WHERE id=%s
            """,
            (batch_id, batch_id),
        )
    conn.commit()
    return int(item_run_id)


def _refresh_batch_counts(conn, batch_id: int, finish_if_idle: bool = False) -> None:
    with conn.cursor() as cur:
        cur.execute(
            """
            UPDATE pricing_task_batches b
            SET selected_count = stats.total_count,
                completed_count = stats.completed_count,
                failed_count = stats.failed_count,
                status = CASE
                    WHEN %s AND stats.running_count = 0 THEN 'completed'
                    ELSE b.status
                END,
                finished_at = CASE
                    WHEN %s AND stats.running_count = 0 THEN COALESCE(b.finished_at, NOW())
                    ELSE b.finished_at
                END,
                updated_at = NOW()
            FROM (
                SELECT COUNT(*) AS total_count,
                       COUNT(*) FILTER (WHERE status IN ('confirmed', 'no_match', 'completed')) AS completed_count,
                       COUNT(*) FILTER (WHERE status = 'failed') AS failed_count,
                       COUNT(*) FILTER (WHERE status = 'running') AS running_count
                FROM pricing_task_batch_item_runs
                WHERE batch_id=%s
            ) stats
            WHERE b.id=%s
            """,
            (finish_if_idle, finish_if_idle, batch_id, batch_id),
        )
    conn.commit()


def _load_batch_item_run_context(conn, item_run_id: int) -> tuple[int, dict[str, Any], dict[str, Any], dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT r.id, r.batch_id, r.status, r.quota_match, r.confirmed_results,
                   r.conversion_check, i.id, i.item_code, i.item_name, i.item_description,
                   i.unit, i.quantity, i.project_id, r.kb_version_id, r.feature_check
            FROM pricing_task_batch_item_runs r
            JOIN boq_items i ON i.id = r.boq_item_id
            WHERE r.id=%s
            """,
            (item_run_id,),
        )
        row = cur.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="batch item run not found")
    batch_id = int(row[1])
    batch_run = {
        "id": row[0],
        "batch_id": batch_id,
        "status": row[2],
        "quota_match": row[3] or {},
        "confirmed_results": row[4] or [],
        "conversion_check": row[5],
        "kb_version_id": int(row[13]),
    }
    feature_check = row[14] if isinstance(row[14], dict) else {}
    effective_description = str(feature_check.get("effective_description") or "").strip()
    boq_item = {
        "run_id": row[0],
        "status": row[2],
        "id": row[6],
        "item_code": row[7],
        "item_name": row[8],
        "item_description": effective_description or row[9] or "",
        "original_item_description": row[9] or "",
        "unit": row[10] or "",
        "quantity": float(row[11]) if row[11] is not None else None,
        "project_id": row[12],
    }
    matches = batch_run["quota_match"].get("matches", []) if isinstance(batch_run["quota_match"], dict) else []
    confirmed_items = _confirmed_items_from_matches(conn, matches, int(row[13]))
    return batch_id, batch_run, boq_item, {"items": confirmed_items}


@router.post("/pricing-task-batches/{batch_id}/items/{boq_item_id}/run-stream")
def pricing_task_batch_run_item_stream(batch_id: int, boq_item_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    def generate():
        conn = get_connection()
        item_run_id = None
        try:
            _ensure_schema(conn)
            require_task_owner(conn, user, batch_id, batch=True)
            with conn.cursor() as cur:
                cur.execute("SELECT owner_user_id FROM pricing_task_batches WHERE id=%s", (batch_id,))
                owner = cur.fetchone()
            if not owner or owner[0] is None:
                raise HTTPException(status_code=404, detail="批量任务不存在")
            batch, boq_item = _load_batch_and_item(conn, batch_id, boq_item_id)
            item_run_id = _create_or_reset_batch_item_run(
                conn, batch_id, boq_item, batch["kb_version_id"]
            )
            yield _sse({
                "type": "run_started", "run_id": item_run_id,
                "kb_version_id": batch["kb_version_id"],
            })
            yield _sse({"type": "item_info", "item": boq_item})
            fields: dict[str, Any] = {}
            reasoning_text_parts: list[str] = []
            had_error = False
            for event_type, data in _stream_pricing_item(
                conn,
                boq_item,
                batch["quota_library_ids"],
                batch["manual_project_id"],
                None,
                None,
                batch["kb_version_id"],
                persist_run=False,
            ):
                if event_type == "reasoning_token":
                    reasoning_text_parts.append(data)
                    yield _sse({"type": "reasoning_token", "token": data})
                elif event_type == "error":
                    had_error = True
                    yield _sse({"type": "error", "error": data})
                else:
                    if event_type == "code_check":
                        fields["code_check"] = data
                    elif event_type == "feature_check":
                        fields["feature_check"] = data
                    elif event_type == "chapter_rule_check":
                        fields["chapter_rule_check"] = data
                    elif event_type == "quota_candidates":
                        fields["quota_candidates"] = data
                    elif event_type == "quota_match":
                        fields["quota_match"] = data
                    elif event_type == "evaluation":
                        fields["evaluation"] = data
                    elif event_type == "step_timing":
                        timings = dict(fields.get("step_timings") or {})
                        timings[str(data["step_no"])] = data
                        fields["step_timings"] = timings
                    yield _sse({"type": event_type, **data} if event_type != "evaluation" else {"type": "evaluation", "evaluation": data})
            quota_match = fields.get("quota_match") or {}
            matches = quota_match.get("matches", []) if isinstance(quota_match, dict) else []
            confirmed_results = _confirmed_results_from_matches(
                conn, matches, batch["kb_version_id"]
            )
            final_status = "no_match" if not matches else "completed"
            if had_error:
                final_status = "failed"
            _update_batch_item_run(
                conn,
                item_run_id,
                **fields,
                confirmed_results=confirmed_results,
                reasoning_text="".join(reasoning_text_parts),
                status=final_status,
                finished_at=datetime.now(),
            )
            _refresh_batch_counts(conn, batch_id, finish_if_idle=True)
            yield _sse({"type": "done", "run_id": item_run_id})
        except Exception as exc:
            if item_run_id:
                _update_batch_item_run(conn, item_run_id, status="failed", error_message=str(exc), finished_at=datetime.now())
                _refresh_batch_counts(conn, batch_id, finish_if_idle=True)
            print(f"[pricing-task] batch SSE error: {exc}", file=sys.stderr, flush=True)
            yield _sse({"type": "error", "error": str(exc)})
        finally:
            conn.close()

    return StreamingResponse(
        bind_default_profile_iterator(user.id, generate(), business_type="pricing_batch", batch_id=batch_id),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-store, no-transform",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


@router.get("/pricing-tasks/{task_id}/items/{boq_item_id}/runs")
def list_item_runs(task_id: int, boq_item_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, task_id)
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id, status, code_check, feature_check, chapter_rule_check, work_procedures, quota_candidates,
                       quota_match, evaluation, conversion_check, coefficient_check,
                       step_timings, error_message, created_at, finished_at, reasoning_text,
                       kb_version_id
                FROM pricing_task_runs
                WHERE task_id=%s AND boq_item_id=%s
                ORDER BY created_at DESC
                """,
                (task_id, boq_item_id),
            )
            rows = cur.fetchall()
        confirmed_results = _load_confirmed_results(conn, [int(r[0]) for r in rows])
        return [
            {
                "id": r[0],
                "status": r[1],
                "code_check": r[2],
                "feature_check": r[3],
                "chapter_rule_check": r[4],
                "work_procedures": r[5],
                "quota_candidates": r[6],
                "quota_match": r[7],
                "evaluation": r[8],
                "conversion_check": _hydrate_conversion_for_run(conn, int(r[0]), r[9]),
                "coefficient_check": r[10],
                "step_timings": r[11],
                "error_message": r[12],
                "created_at": r[13],
                "finished_at": r[14],
                "reasoning_text": r[15],
                "confirmed_results": confirmed_results.get(int(r[0]), []),
                "kb_version_id": int(r[16]) if r[16] is not None else None,
            }
            for r in rows
        ]
    finally:
        conn.close()


@router.get("/pricing-tasks/{task_id}/runs/latest")
def list_latest_task_runs(task_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, task_id)
        with conn.cursor() as cur:
            cur.execute(
                "SELECT manual_project_id FROM pricing_tasks WHERE id=%s AND status <> 'deleted'",
                (task_id,),
            )
            task_row = cur.fetchone()
            if not task_row:
                raise HTTPException(status_code=404, detail="task not found")
            manual_project_id = task_row[0]
            cur.execute(
                """
                SELECT DISTINCT ON (r.boq_item_id)
                       r.boq_item_id, r.id, r.status, r.code_check, r.feature_check, r.chapter_rule_check, r.work_procedures,
                       r.quota_candidates, r.quota_match, r.evaluation, r.conversion_check, r.coefficient_check,
                       r.step_timings, r.error_message, r.created_at, r.finished_at, r.reasoning_text,
                       r.kb_version_id, i.item_code
                FROM pricing_task_runs r
                JOIN boq_items i ON i.id = r.boq_item_id
                WHERE r.task_id=%s
                ORDER BY r.boq_item_id, r.created_at DESC, r.id DESC
                """,
                (task_id,),
            )
            rows = cur.fetchall()
        confirmed_results = _load_confirmed_results(conn, [int(row[1]) for row in rows])
        manual_by_code = _manual_quotas_by_code(conn, manual_project_id)
        result = []
        for row in rows:
            quota_match = row[8] if isinstance(row[8], dict) else {}
            item_code = str(row[18] or "").strip().replace(" ", "")
            current_evaluation = _evaluate(
                quota_match.get("matches", []),
                manual_by_code.get(item_code, []),
            )
            result.append(
                {
                    "boq_item_id": row[0],
                    "run": {
                        "id": row[1],
                        "status": row[2],
                        "code_check": row[3],
                        "feature_check": row[4],
                        "chapter_rule_check": row[5],
                        "work_procedures": row[6],
                        "quota_candidates": row[7],
                        "quota_match": quota_match,
                        "evaluation": current_evaluation,
                        "conversion_check": _hydrate_conversion_for_run(conn, int(row[1]), row[10]),
                        "coefficient_check": row[11],
                        "step_timings": row[12],
                        "error_message": row[13],
                        "created_at": row[14],
                        "finished_at": row[15],
                        "reasoning_text": row[16],
                        "confirmed_results": confirmed_results.get(int(row[1]), []),
                        "kb_version_id": int(row[17]) if row[17] is not None else None,
                    },
                }
            )
        return result
    finally:
        conn.close()


@router.post("/pricing-tasks/{task_id}/items/{boq_item_id}/run-stream")
def pricing_task_run_item_stream(task_id: int, boq_item_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    def generate():
        conn = get_connection()
        run_id = None
        try:
            _ensure_schema(conn)
            require_task_owner(conn, user, task_id)
            with conn.cursor() as cur:
                cur.execute("SELECT owner_user_id FROM pricing_tasks WHERE id=%s", (task_id,))
                owner = cur.fetchone()
            if not owner or owner[0] is None:
                raise HTTPException(status_code=404, detail="组价任务不存在")
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT id, name, boq_project_id, quota_library_ids, manual_project_id, kb_version_id FROM pricing_tasks WHERE id=%s",
                    (task_id,),
                )
                task = cur.fetchone()
                if not task:
                    yield _sse({"type": "error", "error": "task not found"})
                    return
                quota_library_ids = task[3] or []
                manual_project_id = task[4]
                kb_version_id = int(task[5])
                cur.execute(
                    """
                    SELECT id, item_code, item_name, item_description, unit, quantity, project_id
                    FROM boq_items WHERE id = %s AND project_id = %s
                    """,
                    (boq_item_id, task[2]),
                )
                row = cur.fetchone()
            if not row:
                yield _sse({"type": "error", "error": "item not found in task project"})
                return
            boq_item = {
                "id": row[0],
                "item_code": row[1],
                "item_name": row[2],
                "item_description": row[3],
                "unit": row[4],
                "quantity": float(row[5]) if row[5] is not None else None,
                "project_id": row[6],
            }
            run_id = _create_run(conn, task_id, boq_item, kb_version_id)
            yield _sse({
                "type": "run_started", "run_id": run_id,
                "kb_version_id": kb_version_id,
            })
            yield _sse({"type": "item_info", "item": boq_item})
            had_error = False
            reasoning_text_parts: list[str] = []
            for event_type, data in _stream_pricing_item(
                conn, boq_item, quota_library_ids, manual_project_id,
                task_id, run_id, kb_version_id
            ):
                if event_type == "reasoning_token":
                    reasoning_text_parts.append(data)
                    yield _sse({"type": "reasoning_token", "token": data})
                elif event_type == "error":
                    had_error = True
                    yield _sse({"type": "error", "error": data})
                else:
                    if event_type == "evaluation":
                        yield _sse({"type": "evaluation", "evaluation": data})
                    else:
                        yield _sse({"type": event_type, **data})
            if not had_error:
                _update_run(conn, run_id, status="completed", reasoning_text="".join(reasoning_text_parts), finished_at=datetime.now())
                yield _sse({"type": "done", "run_id": run_id})
        except Exception as exc:
            friendly_error = _user_facing_model_error(exc)
            if run_id:
                _update_run(conn, run_id, status="failed", error_message=friendly_error, finished_at=datetime.now())
            print(f"[pricing-task] SSE error: {exc}", file=sys.stderr, flush=True)
            yield _sse({"type": "error", "error": friendly_error})
        finally:
            conn.close()

    return StreamingResponse(
        bind_default_profile_iterator(user.id, generate(), business_type="pricing_task", task_id=task_id),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-store, no-transform",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


@router.post("/pricing-tasks/{task_id}/accuracy-report")
def generate_pricing_task_accuracy_report(task_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, task_id)
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT t.id, t.name, t.boq_project_id, p.project_name, t.manual_project_id, mp.project_name
                FROM pricing_tasks t
                JOIN boq_projects p ON p.id = t.boq_project_id
                LEFT JOIN manual_boq_projects mp ON mp.id = t.manual_project_id
                WHERE t.id=%s AND t.status <> 'deleted'
                """,
                (task_id,),
            )
            task_row = cur.fetchone()
            if not task_row:
                raise HTTPException(status_code=404, detail="task not found")
            cur.execute(
                """
                SELECT DISTINCT ON (r.boq_item_id)
                       r.id, r.boq_item_id, r.quota_match, r.evaluation, r.conversion_check, r.coefficient_check,
                       i.item_code, i.item_name, i.item_description, i.unit, i.quantity
                FROM pricing_task_runs r
                JOIN boq_items i ON i.id = r.boq_item_id
                WHERE r.task_id=%s
                ORDER BY r.boq_item_id, r.created_at DESC, r.id DESC
                """,
                (task_id,),
            )
            rows = cur.fetchall()
        task = {
            "id": task_row[0],
            "name": task_row[1],
            "boq_project_id": task_row[2],
            "project_name": task_row[3],
            "manual_project_id": task_row[4],
            "manual_project_name": task_row[5],
        }
        items = []
        metrics = {
            "total_items": len(rows),
            "evaluated_item_count": 0,
            "exact_item_count": 0,
            "hit_count": 0,
            "missed_count": 0,
            "extra_count": 0,
            "manual_count": 0,
            "ai_count": 0,
            "hit_rate": None,
        }
        for row in rows:
            quota_match = row[2] if isinstance(row[2], dict) else {}
            evaluation = _evaluate(
                quota_match.get("matches", []),
                _manual_quotas(conn, task_row[4], row[6]),
            )
            if "manual_quotas" in evaluation:
                metrics["evaluated_item_count"] += 1
                metrics["hit_count"] += int(evaluation.get("hit_count") or 0)
                metrics["missed_count"] += int(evaluation.get("missed_count") or 0)
                metrics["extra_count"] += int(evaluation.get("extra_count") or 0)
                metrics["manual_count"] += int(evaluation.get("manual_count") or 0)
                metrics["ai_count"] += int(evaluation.get("ai_count") or 0)
                if int(evaluation.get("manual_count") or 0) > 0 and int(evaluation.get("missed_count") or 0) == 0 and int(evaluation.get("extra_count") or 0) == 0:
                    metrics["exact_item_count"] += 1
            items.append(
                {
                    "run_id": row[0],
                    "boq_item_id": row[1],
                    "quota_match": quota_match,
                    "evaluation": evaluation,
                    "conversion_check": _hydrate_conversion_for_run(conn, int(row[0]), row[4]) if isinstance(row[4], dict) else row[4],
                    "coefficient_check": row[5] if isinstance(row[5], dict) else row[5],
                    "item_code": row[6],
                    "item_name": row[7],
                    "item_description": row[8] or "",
                    "unit": row[9] or "",
                    "quantity": float(row[10]) if row[10] is not None else None,
                }
            )
        if metrics["evaluated_item_count"] == 0:
            raise HTTPException(status_code=400, detail="task has no evaluated pricing runs")
        if metrics["manual_count"] == 0:
            raise HTTPException(status_code=400, detail="task has no manual comparison data")
        metrics["hit_rate"] = round(metrics["hit_count"] / metrics["manual_count"], 4) if metrics["manual_count"] else None
        report = _generate_task_accuracy_report(task, items, metrics)
        with conn.cursor() as cur:
            cur.execute("UPDATE pricing_tasks SET accuracy_report=%s, updated_at=NOW() WHERE id=%s", (Json(report, dumps=_json_dumps), task_id))
        conn.commit()
        return report
    finally:
        conn.close()


@router.get("/pricing-tasks/{task_id}/detail-report")
def get_pricing_task_detail_report(task_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        require_task_owner(conn, user, task_id)
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT t.id, t.name, t.boq_project_id, p.project_name, t.manual_project_id, mp.project_name,
                       t.quota_library_ids
                FROM pricing_tasks t
                JOIN boq_projects p ON p.id = t.boq_project_id
                LEFT JOIN manual_boq_projects mp ON mp.id = t.manual_project_id
                WHERE t.id=%s AND t.status <> 'deleted'
                """,
                (task_id,),
            )
            task_row = cur.fetchone()
            if not task_row:
                raise HTTPException(status_code=404, detail="task not found")
            cur.execute(
                """
                SELECT DISTINCT ON (r.boq_item_id)
                       r.id, r.boq_item_id, r.status, r.code_check, r.feature_check, r.chapter_rule_check,
                       r.quota_candidates, r.quota_match, r.evaluation, r.conversion_check,
                       r.coefficient_check, r.step_timings, r.reasoning_text, r.created_at, r.finished_at,
                       i.item_code, i.item_name, i.item_description, i.unit, i.quantity, i.item_seq
                FROM pricing_task_runs r
                JOIN boq_items i ON i.id = r.boq_item_id
                WHERE r.task_id=%s
                ORDER BY r.boq_item_id, r.created_at DESC, r.id DESC
                """,
                (task_id,),
            )
            rows = cur.fetchall()
        task = {
            "id": task_row[0],
            "name": task_row[1],
            "boq_project_id": task_row[2],
            "project_name": task_row[3],
            "manual_project_id": task_row[4],
            "manual_project_name": task_row[5],
            "quota_library_ids": task_row[6] or [],
        }
        return _build_pricing_task_detail_report(conn, task, rows)
    finally:
        conn.close()


@router.get("/pricing-tasks/{task_id}/detail-report/export")
def export_pricing_task_detail_report(task_id: int, user: CurrentUser = Depends(current_user)):
    report = get_pricing_task_detail_report(task_id, user)
    stream = _build_pricing_task_detail_report_excel(report)
    task_name = str((report.get("task") or {}).get("name") or f"task-{task_id}")
    safe_name = re.sub(r'[\\/:*?"<>|]+', "_", task_name).strip() or f"task-{task_id}"
    filename = f"单条组价明细报表-{safe_name}.xlsx"
    quoted = quote(filename)
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quoted}"}
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def _get_new_batch_detail_report(batch_id: int, workflow_version: str = "streamlined") -> dict[str, Any]:
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT b.id, b.name, b.boq_project_id, p.project_name,
                       b.manual_project_id, mp.project_name, b.quota_library_ids
                FROM pricing_task_batches b
                JOIN boq_projects p ON p.id = b.boq_project_id
                LEFT JOIN manual_boq_projects mp ON mp.id = b.manual_project_id
                WHERE b.id=%s
                  AND b.status <> 'deleted'
                  AND b.workflow_version=%s
                """,
                (batch_id, workflow_version),
            )
            batch_row = cur.fetchone()
            if not batch_row:
                raise HTTPException(status_code=404, detail="new pricing batch not found")
            cur.execute(
                """
                SELECT DISTINCT ON (r.boq_item_id)
                       r.id, r.boq_item_id, r.status, r.code_check, r.feature_check, r.chapter_rule_check,
                       r.quota_candidates, r.quota_match, r.evaluation, r.conversion_check,
                       r.coefficient_check, r.step_timings, r.reasoning_text, r.created_at, r.finished_at,
                       i.item_code, i.item_name, i.item_description, i.unit, i.quantity, i.item_seq
                FROM pricing_task_batch_item_runs r
                JOIN boq_items i ON i.id = r.boq_item_id
                WHERE r.batch_id=%s AND r.evaluation IS NOT NULL
                ORDER BY r.boq_item_id, r.created_at DESC, r.id DESC
                """,
                (batch_id,),
            )
            rows = cur.fetchall()
        if not rows:
            raise HTTPException(status_code=400, detail="batch has no evaluated pricing results")
        batch = {
            "id": batch_row[0],
            "name": batch_row[1],
            "boq_project_id": batch_row[2],
            "project_name": batch_row[3],
            "manual_project_id": batch_row[4],
            "manual_project_name": batch_row[5],
            "quota_library_ids": batch_row[6] or [],
            "report_kind": "batch",
        }
        report = _build_pricing_task_detail_report(conn, batch, rows, source_type="batch")
        for item in report.get("items") or []:
            consistency = item.get("consistency") or {}
            if consistency.get("status") == "一致":
                consistency["status"] = "完全一致"
        return report
    finally:
        conn.close()


@router.get("/pricing-task-new-batches/{batch_id}/detail-report/export")
def export_new_pricing_task_batch_detail_report(batch_id: int):
    report = _get_new_batch_detail_report(batch_id)
    stream = _build_pricing_task_detail_report_excel(report)
    batch_name = str((report.get("task") or {}).get("name") or f"batch-{batch_id}")
    safe_name = re.sub(r'[\\/:*?"<>|]+', "_", batch_name).strip() or f"batch-{batch_id}"
    filename = f"新批量组价明细报表-{safe_name}.xlsx"
    quoted = quote(filename)
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{quoted}"}
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.get("/pricing-task-background-batches/{batch_id}/detail-report/export")
def export_background_pricing_task_batch_detail_report(batch_id: int):
    report = _get_new_batch_detail_report(batch_id, "background")
    stream = _build_pricing_task_detail_report_excel(report)
    batch_name = str((report.get("task") or {}).get("name") or f"batch-{batch_id}")
    safe_name = re.sub(r'[\\/:*?"<>|]+', "_", batch_name).strip() or f"batch-{batch_id}"
    filename = f"后台批量组价明细报表-{safe_name}.xlsx"
    quoted = quote(filename)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quoted}"},
    )


@router.post("/pricing-task-runs/{run_id}/conversion-check-stream")
def pricing_task_conversion_check_stream(run_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    def generate():
        conn = get_connection()
        try:
            _ensure_schema(conn)
            with conn.cursor() as cur:
                cur.execute("SELECT t.id,t.owner_user_id FROM pricing_task_runs r JOIN pricing_tasks t ON t.id=r.task_id WHERE r.id=%s", (run_id,))
                owner = cur.fetchone()
            if not owner or owner[1] is None:
                raise HTTPException(status_code=404, detail="运行记录不存在")
            require_task_owner(conn, user, int(owner[0]))
            profile_context = use_default_profile(int(owner[1]), business_type="pricing_task_conversion", task_id=int(owner[0]), run_id=run_id)
            profile_context.__enter__()
            step_timings = _load_step_timings(conn, run_id)
            step_started_at = datetime.now()
            step_started_perf = perf_counter()
            boq_item, confirmed_items = _confirmed_conversion_context(conn, run_id)
            yield _sse({"type": "conversion_check_start", "run_id": run_id, "total": len(confirmed_items)})

            if not confirmed_items:
                result = {"items": [], "issues": ["未找到已确认定额，无法进行组合换算。"]}
                yield _sse({"type": "combo_adjustment_rules", "items": []})
                _update_run(conn, run_id, conversion_check=result)
                yield _sse({"type": "conversion_check", "conversion_check": result})
                yield _sse({"type": "step_timing", **_finish_step_timing(conn, run_id, step_timings, 7, "组合换算", step_started_at, step_started_perf)})
                yield _sse({"type": "done", "run_id": run_id})
                return

            input_payload = {
                "boq_item": boq_item,
                "confirmed_quotas": confirmed_items,
                "combo_adjustment_rule_guide": COMBO_ADJUSTMENT_RULE_GUIDE,
            }
            context_text = _json_dumps(input_payload)
            system_prompt = build_system_prompt()
            combo_preview = _default_conversion_check(confirmed_items)
            yield _sse({"type": "combo_adjustment_rules", "items": combo_preview["items"]})
            if not any(item.get("adjustment_rules") for item in confirmed_items):
                result = _default_conversion_check(
                    confirmed_items,
                    "已确认定额均未查询到 tdek_tzhhs 组合定额规则。",
                    boq_item=boq_item,
                )
                _update_run(conn, run_id, conversion_check=result)
                yield _sse({"type": "conversion_check", "conversion_check": result})
                yield _sse({"type": "step_timing", **_finish_step_timing(conn, run_id, step_timings, 7, "组合换算", step_started_at, step_started_perf)})
                yield _sse({"type": "done", "run_id": run_id})
                return
            analysis_prompt = (
                "请对已人工确认的定额进行第六轮组合定额换算分析。先进行分析，不要调用工具，不要输出 JSON。\n"
                "本轮只允许使用输入中的 adjustment_rules（来自 tdek_tzhhs 并关联 tdek_tdezm）。\n"
                "请逐条基础定额、逐条组合规则判断：项目特征中是否存在与 prompt、combo_name 增减指标匹配的数量特征。\n"
                "若匹配，请提取数量特征原文和数值；不要自行输出材料替换、工料机调整或工程量系数调整。\n"
                "第六轮结果只作为组合定额次数建议，不允许修改已确认定额，也不要改写工程量系数。\n\n"
                f"【输入数据】\n{context_text}"
            )
            conversion_analysis = ""
            for event_type, data in _stream_text_completion(
                [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": analysis_prompt},
                ],
                6000,
            ):
                if event_type == "reasoning_token":
                    yield _sse({"type": "reasoning_token", "token": data})
                elif event_type == "text_result":
                    conversion_analysis = data

            submit_prompt = (
                "请严格调用 submit_conversion_check 提交第六轮结构化换算建议。\n"
                "必须覆盖每一条已确认定额。\n"
                "每个 item 的 adjustment_rules 必须覆盖输入中该定额的每一条组合规则；rule_index 必须与输入保持一致。\n"
                "只判断项目特征数量特征是否匹配 prompt/combo_name 的增减指标；匹配时 matched=true，并填写 matched_feature 和 feature_value。\n"
                "calculated_times 可按 (feature_value - base_value) / increment_unit 先填写，但后端会重新计算；小数保留，负数按 0 理解。\n"
                "不匹配时 matched=false，feature_value 和 calculated_times 填 0，并说明原因。\n"
                "不要输出工料机调整、材料替换、工程量系数调整或其他非组合定额规则。\n"
                "reason 写整体判断说明，missing_inputs 写缺失的数量特征或单位信息。\n"
                "confidence 只能是 high、medium、low。\n\n"
                f"【第六轮分析】\n{conversion_analysis or '（无分析文本）'}\n\n"
                f"【输入数据】\n{context_text}"
            )
            raw_result = _run_submit_conversion_check(
                [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": submit_prompt},
                ]
            )
            result = _normalize_conversion_check(raw_result, confirmed_items, boq_item)
            _update_run(conn, run_id, conversion_check=result)
            yield _sse({"type": "conversion_check", "conversion_check": result})
            yield _sse({"type": "step_timing", **_finish_step_timing(conn, run_id, step_timings, 7, "组合换算", step_started_at, step_started_perf)})
            yield _sse({"type": "done", "run_id": run_id})
        except HTTPException as exc:
            yield _sse({"type": "error", "error": str(exc.detail)})
        except Exception as exc:
            print(f"[pricing-task] conversion check SSE error: {exc}", file=sys.stderr, flush=True)
            yield _sse({"type": "error", "error": str(exc)})
        finally:
            if profile_context:
                profile_context.__exit__(None, None, None)
            conn.close()

    return StreamingResponse(generate(), media_type="text/event-stream")


@router.post("/pricing-task-runs/{run_id}/coefficient-check-stream")
def pricing_task_coefficient_check_stream(run_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    def generate():
        conn = get_connection()
        try:
            _ensure_schema(conn)
            with conn.cursor() as cur:
                cur.execute("SELECT t.id,t.owner_user_id FROM pricing_task_runs r JOIN pricing_tasks t ON t.id=r.task_id WHERE r.id=%s", (run_id,))
                owner = cur.fetchone()
            if not owner or owner[1] is None:
                raise HTTPException(status_code=404, detail="运行记录不存在")
            require_task_owner(conn, user, int(owner[0]))
            profile_context = use_default_profile(int(owner[1]), business_type="pricing_task_coefficient", task_id=int(owner[0]), run_id=run_id)
            profile_context.__enter__()
            step_timings = _load_step_timings(conn, run_id)
            step_started_at = datetime.now()
            step_started_perf = perf_counter()
            boq_item, items = _coefficient_context(conn, run_id)
            yield _sse({"type": "coefficient_check_start", "run_id": run_id, "total": len(items)})
            preview = _default_coefficient_check(items)
            yield _sse({"type": "coefficient_rules", "items": preview["items"]})

            if not items:
                result = {"items": [], "issues": ["未找到已确认定额，无法进行系数换算。"]}
                _update_run(conn, run_id, coefficient_check=result)
                yield _sse({"type": "coefficient_check", "coefficient_check": result})
                yield _sse({"type": "step_timing", **_finish_step_timing(conn, run_id, step_timings, 8, "系数换算", step_started_at, step_started_perf)})
                yield _sse({"type": "done", "run_id": run_id})
                return

            if not any(item.get("coefficient_rules") for item in items):
                result = _default_coefficient_check(items, "所有定额均未查询到 tdek_tznhs 系数换算说明。")
                _update_run(conn, run_id, coefficient_check=result)
                yield _sse({"type": "coefficient_check", "coefficient_check": result})
                yield _sse({"type": "step_timing", **_finish_step_timing(conn, run_id, step_timings, 8, "系数换算", step_started_at, step_started_perf)})
                yield _sse({"type": "done", "run_id": run_id})
                return

            input_payload = {
                "boq_item": boq_item,
                "quotas": items,
                "rule_source": "tdek_tznhs.hssm",
                "target_resource_type_guide": {
                    "1": "人工费/人工消耗量",
                    "2": "材料费/材料消耗量",
                    "3": "机械费/机械消耗量",
                    "all": "子目或相应子目整体乘以系数",
                },
                "output_note": "本轮只保存系数和作用对象，不计算调整后含量。",
            }
            context_text = _json_dumps(input_payload)
            system_prompt = build_system_prompt()
            analysis_prompt = (
                "请进行第七轮系数换算分析。先分析，不要调用工具，不要输出 JSON。\n"
                "规则来源只允许使用输入中的 coefficient_rules（tdek_tznhs.hssm）。\n"
                "请逐条判断项目特征是否触发 hssm；触发时提取命中特征、特征值、系数和作用对象。\n"
                "作用对象按说明判断：人工费=1，材料费=2，机械消耗量/机械费=3，子目乘以系数/相应子目乘以系数=all。\n"
                "本轮不计算调整后含量，只说明哪些工料机行应显示乘以的系数。\n\n"
                f"【输入数据】\n{context_text}"
            )
            try:
                analysis = ""
                for event_type, data in _stream_text_completion(
                    [
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": analysis_prompt},
                    ],
                    6000,
                ):
                    if event_type == "reasoning_token":
                        yield _sse({"type": "reasoning_token", "token": data})
                    elif event_type == "text_result":
                        analysis = data

                submit_prompt = (
                    "请严格调用 submit_coefficient_check 提交第七轮结构化系数换算判断。\n"
                    "必须覆盖输入中的每一条定额、每一条 coefficient_rules；rule_index 必须保持一致。\n"
                    "matched=true 时填写 matched_feature、feature_value、factor、target_resource_types 和 reason。\n"
                    "factor 从 hssm 中的“系数X”提取；如果未命中，matched=false，factor 仍填写 hssm 中可识别的系数或 1。\n"
                    "target_resource_types 只能使用 1、2、3、all；不要输出调整后数量。\n\n"
                    f"【第七轮分析】\n{analysis or '（无分析文本）'}\n\n"
                    f"【输入数据】\n{context_text}"
                )
                raw_result = _run_submit_coefficient_check(
                    [
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": submit_prompt},
                    ]
                )
                result = _normalize_coefficient_check(raw_result, items)
            except Exception as exc:
                print(f"[pricing-task] coefficient model result fallback: {exc}", file=sys.stderr, flush=True)
                result = _default_coefficient_check(
                    items,
                    f"模型结构化输出解析失败，已保留系数换算说明待人工复核：{exc}",
                )
            _update_run(conn, run_id, coefficient_check=result)
            yield _sse({"type": "coefficient_check", "coefficient_check": result})
            yield _sse({"type": "step_timing", **_finish_step_timing(conn, run_id, step_timings, 8, "系数换算", step_started_at, step_started_perf)})
            yield _sse({"type": "done", "run_id": run_id})
        except HTTPException as exc:
            yield _sse({"type": "error", "error": str(exc.detail)})
        except Exception as exc:
            print(f"[pricing-task] coefficient check SSE error: {exc}", file=sys.stderr, flush=True)
            yield _sse({"type": "error", "error": str(exc)})
        finally:
            if profile_context:
                profile_context.__exit__(None, None, None)
            conn.close()

    return StreamingResponse(generate(), media_type="text/event-stream")


def _batch_conversion_business_events(confirmed_items: list[dict[str, Any]], boq_item: dict[str, Any]):
    if not confirmed_items:
        yield "combo_adjustment_rules", {"items": []}
        yield "conversion_check", {"conversion_check": {"items": [], "issues": ["未找到批量确认定额，无法进行组合换算。"]}}
        return
    preview = _default_conversion_check(confirmed_items)
    yield "combo_adjustment_rules", {"items": preview["items"]}
    if not any(item.get("adjustment_rules") for item in confirmed_items):
        result = _default_conversion_check(
            confirmed_items, "已确认定额均未查询到 tdek_tzhhs 组合定额规则。", boq_item=boq_item
        )
        yield "conversion_check", {"conversion_check": result}
        return
    input_payload = {
        "boq_item": boq_item,
        "confirmed_quotas": confirmed_items,
        "combo_adjustment_rule_guide": COMBO_ADJUSTMENT_RULE_GUIDE,
    }
    context_text = _json_dumps(input_payload)
    system_prompt = build_system_prompt()
    analysis_prompt = (
        "请对已确认的批量定额进行第六轮组合定额换算分析。先进行分析，不要调用工具，不要输出 JSON。\n"
        "本轮只允许使用输入中的 adjustment_rules（来自 tdek_tzhhs 并关联 tdek_tdezm）。\n"
        "请逐条基础定额、逐条组合规则判断：项目特征中是否存在与 prompt、combo_name 增减指标匹配的数量特征。\n"
        "若匹配，请提取数量特征原文和数值；不要自行输出材料替换、工料机调整或工程量系数调整。\n\n"
        f"【输入数据】\n{context_text}"
    )
    conversion_analysis = ""
    for event_type, data in _stream_text_completion(
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": analysis_prompt}], 6000
    ):
        if event_type == "reasoning_token":
            yield "reasoning_token", {"token": data}
        elif event_type == "text_result":
            conversion_analysis = data
    submit_prompt = (
        "请严格调用 submit_conversion_check 提交第六轮结构化换算建议。\n"
        "必须覆盖每一条已确认定额。\n"
        "每个 item 的 adjustment_rules 必须覆盖输入中该定额的每一条组合规则；rule_index 必须与输入保持一致。\n"
        "只判断项目特征数量特征是否匹配 prompt/combo_name 的增减指标；匹配时 matched=true，并填写 matched_feature 和 feature_value。\n"
        "不匹配时 matched=false，feature_value 和 calculated_times 填 0，并说明原因。\n\n"
        f"【第六轮分析】\n{conversion_analysis or '（无分析文本）'}\n\n"
        f"【输入数据】\n{context_text}"
    )
    raw_result = _run_submit_conversion_check(
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": submit_prompt}]
    )
    yield "conversion_check", {"conversion_check": _normalize_conversion_check(raw_result, confirmed_items, boq_item)}


def _batch_coefficient_business_events(items: list[dict[str, Any]], boq_item: dict[str, Any]):
    preview = _default_coefficient_check(items)
    yield "coefficient_rules", {"items": preview["items"]}
    if not items:
        yield "coefficient_check", {"coefficient_check": {"items": [], "issues": ["未找到批量确认定额，无法进行系数换算。"]}}
        return
    if not any(item.get("coefficient_rules") for item in items):
        yield "coefficient_check", {"coefficient_check": _default_coefficient_check(items, "所有定额均未查询到 tdek_tznhs 系数换算说明。")}
        return
    input_payload = {
        "boq_item": boq_item,
        "quotas": items,
        "rule_source": "tdek_tznhs.hssm",
        "target_resource_type_guide": {
            "1": "人工费/人工消耗量", "2": "材料费/材料消耗量",
            "3": "机械费/机械消耗量", "all": "子目或相应子目整体乘以系数",
        },
        "output_note": "本轮只保存系数和作用对象，不计算调整后含量。",
    }
    context_text = _json_dumps(input_payload)
    system_prompt = build_system_prompt()
    analysis_prompt = (
        "请进行第七轮系数换算分析。先分析，不要调用工具，不要输出 JSON。\n"
        "规则来源只允许使用输入中的 coefficient_rules（tdek_tznhs.hssm）。\n"
        "请逐条判断项目特征是否触发 hssm；触发时提取命中特征、特征值、系数和作用对象。\n"
        "本轮不计算调整后含量，只说明哪些工料机行应显示乘以的系数。\n\n"
        f"【输入数据】\n{context_text}"
    )
    try:
        analysis = ""
        for event_type, data in _stream_text_completion(
            [{"role": "system", "content": system_prompt}, {"role": "user", "content": analysis_prompt}], 6000
        ):
            if event_type == "reasoning_token":
                yield "reasoning_token", {"token": data}
            elif event_type == "text_result":
                analysis = data
        submit_prompt = (
            "请严格调用 submit_coefficient_check 提交第七轮结构化系数换算判断。\n"
            "必须覆盖输入中的每一条定额、每一条 coefficient_rules；rule_index 必须保持一致。\n"
            "matched=true 时填写 matched_feature、feature_value、factor、target_resource_types 和 reason。\n"
            "target_resource_types 只能使用 1、2、3、all；不要输出调整后数量。\n\n"
            f"【第七轮分析】\n{analysis or '（无分析文本）'}\n\n"
            f"【输入数据】\n{context_text}"
        )
        raw_result = _run_submit_coefficient_check(
            [{"role": "system", "content": system_prompt}, {"role": "user", "content": submit_prompt}]
        )
        result = _normalize_coefficient_check(raw_result, items)
    except Exception as exc:
        if isinstance(exc, ModelRateLimitError):
            raise
        print(f"[pricing-task] batch coefficient model result fallback: {exc}", file=sys.stderr, flush=True)
        result = _default_coefficient_check(items, f"模型结构化输出解析失败，已保留系数换算说明待人工复核：{exc}")
    yield "coefficient_check", {"coefficient_check": result}


@router.post("/pricing-task-batch-item-runs/{item_run_id}/conversion-check-stream")
def pricing_task_batch_conversion_check_stream(item_run_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    def generate():
        conn = get_connection()
        profile_context = None
        try:
            _ensure_schema(conn)
            _, _, boq_item, context = _load_batch_item_run_context(conn, item_run_id)
            with conn.cursor() as cur:
                cur.execute("SELECT b.owner_user_id FROM pricing_task_batch_item_runs r JOIN pricing_task_batches b ON b.id=r.batch_id WHERE r.id=%s", (item_run_id,))
                owner = cur.fetchone()
            if not owner or owner[0] is None:
                raise HTTPException(status_code=404, detail="批量运行记录不存在")
            require_task_owner(conn, user, _load_batch_item_run_context(conn, item_run_id)[0], batch=True)
            confirmed_items = context["items"]
            step_timings = _load_batch_step_timings(conn, item_run_id)
            step_started_at = datetime.now()
            step_started_perf = perf_counter()
            yield _sse({"type": "conversion_check_start", "run_id": item_run_id, "total": len(confirmed_items)})

            result = None
            for event_type, payload in _batch_conversion_business_events(confirmed_items, boq_item):
                if event_type == "conversion_check":
                    result = payload["conversion_check"]
                yield _sse({"type": event_type, **payload})
            if result is None:
                raise RuntimeError("组合换算未生成结果")
            _update_batch_item_run(conn, item_run_id, conversion_check=result)
            yield _sse({"type": "step_timing", **_finish_batch_step_timing(
                conn, item_run_id, step_timings, 6, "组合换算", step_started_at, step_started_perf
            )})
            yield _sse({"type": "done", "run_id": item_run_id})
        except HTTPException as exc:
            yield _sse({"type": "error", "error": str(exc.detail)})
        except Exception as exc:
            print(f"[pricing-task] batch conversion check SSE error: {exc}", file=sys.stderr, flush=True)
            yield _sse({"type": "error", "error": str(exc)})
        finally:
            conn.close()

    return StreamingResponse(bind_default_profile_iterator(user.id, generate(), business_type="pricing_batch_conversion", item_run_id=item_run_id), media_type="text/event-stream")


@router.post("/pricing-task-batch-item-runs/{item_run_id}/coefficient-check-stream")
def pricing_task_batch_coefficient_check_stream(item_run_id: int, user: CurrentUser = Depends(current_user)):
    from db.connection import get_connection

    def generate():
        conn = get_connection()
        try:
            _ensure_schema(conn)
            batch_id, batch_run, boq_item, context = _load_batch_item_run_context(conn, item_run_id)
            require_task_owner(conn, user, batch_id, batch=True)
            with conn.cursor() as cur:
                cur.execute("SELECT owner_user_id FROM pricing_task_batches WHERE id=%s", (batch_id,))
                owner = cur.fetchone()
            if not owner or owner[0] is None:
                raise HTTPException(status_code=404, detail="批量任务不存在")
            kb_version_id = int(batch_run["kb_version_id"])
            conversion_check = _hydrate_conversion_combo_resources(
                conn, batch_run.get("conversion_check") or {}, kb_version_id
            ) or {}
            items = _coefficient_items_from_context(conn, context["items"], conversion_check, kb_version_id)
            step_timings = _load_batch_step_timings(conn, item_run_id)
            step_started_at = datetime.now()
            step_started_perf = perf_counter()
            yield _sse({"type": "coefficient_check_start", "run_id": item_run_id, "total": len(items)})

            result = None
            for event_type, payload in _batch_coefficient_business_events(items, boq_item):
                if event_type == "coefficient_check":
                    result = payload["coefficient_check"]
                yield _sse({"type": event_type, **payload})
            if result is None:
                raise RuntimeError("系数换算未生成结果")
            _update_batch_item_run(
                conn, item_run_id, coefficient_check=result, status="confirmed", finished_at=datetime.now()
            )
            _refresh_batch_counts(conn, batch_id, finish_if_idle=True)
            yield _sse({"type": "step_timing", **_finish_batch_step_timing(
                conn, item_run_id, step_timings, 7, "系数换算", step_started_at, step_started_perf
            )})
            yield _sse({"type": "done", "run_id": item_run_id})
        except HTTPException as exc:
            yield _sse({"type": "error", "error": str(exc.detail)})
        except Exception as exc:
            print(f"[pricing-task] batch coefficient check SSE error: {exc}", file=sys.stderr, flush=True)
            yield _sse({"type": "error", "error": str(exc)})
        finally:
            conn.close()

    return StreamingResponse(bind_default_profile_iterator(user.id, generate(), business_type="pricing_batch_coefficient", item_run_id=item_run_id), media_type="text/event-stream")


_BACKGROUND_TOOL_NEXT = {
    "code_check": "feature_check",
    "feature_check": "chapter_rule_check",
    "chapter_rule_check": "quota_candidates",
    "quota_candidates": "quota_match",
    "quota_match": "evaluation",
}


def _background_tool_output(event_type: str, data: dict[str, Any]) -> str:
    if event_type == "code_check":
        return "编码与名称一致" if data.get("is_consistent") else "编码核查完成"
    if event_type == "feature_check":
        return "项目特征完整" if data.get("is_complete") else "项目特征存在缺失"
    if event_type == "chapter_rule_check":
        return f"命中 {sum(1 for rule in data.get('rules', []) if rule.get('matched'))} 条章节规则"
    if event_type == "quota_candidates":
        return f"检索到 {data.get('total', 0)} 条候选定额"
    if event_type == "quota_match":
        return f"输出 {len(data.get('matches', []))} 条定额"
    if event_type == "evaluation":
        return f"命中 {data.get('hit_count', 0)}，遗漏 {data.get('missed_count', 0)}，额外 {data.get('extra_count', 0)}"
    return "执行完成"


def _set_background_tool(conn, batch_id: int, execution_id: int, item_run_id: int, boq_item_id: int,
                         tool_id: str, status: str, output: str | None = None) -> None:
    _update_batch_item_run(
        conn, item_run_id, current_tool=tool_id, current_tool_status=status,
        current_tool_output=output, updated_at=datetime.now(),
        lease_expires_at=datetime.now() + timedelta(minutes=15),
        commit=False,
    )
    _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id,
                      "tool_started" if status == "running" else "tool_completed",
                      {"tool_id": tool_id, "status": status, "output": output})


def _execute_background_conversion(conn, batch_id: int, execution_id: int, item_run_id: int, boq_item_id: int) -> None:
    _set_background_tool(conn, batch_id, execution_id, item_run_id, boq_item_id, "conversion_check", "running")
    _, _, boq_item, context = _load_batch_item_run_context(conn, item_run_id)
    started_at, started_perf = datetime.now(), perf_counter()
    result = None
    for event_type, payload in _batch_conversion_business_events(context["items"], boq_item):
        if event_type == "reasoning_token":
            continue
        if event_type == "conversion_check":
            result = payload["conversion_check"]
        else:
            _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, event_type, payload)
    if result is None:
        raise RuntimeError("组合换算未生成结果")
    timings = _load_batch_step_timings(conn, item_run_id)
    timing = _finish_step_timing(conn, None, timings, 6, "组合换算", started_at, started_perf)
    _update_batch_item_run(
        conn, item_run_id, conversion_check=result, step_timings=timings, updated_at=datetime.now(), commit=False
    )
    _set_background_tool(
        conn, batch_id, execution_id, item_run_id, boq_item_id, "conversion_check", "success",
        f"完成 {len(result.get('items', []))} 条定额检查",
    )
    _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, "step_timing", timing)


def _execute_background_coefficient(conn, batch_id: int, execution_id: int, item_run_id: int, boq_item_id: int) -> None:
    _set_background_tool(conn, batch_id, execution_id, item_run_id, boq_item_id, "coefficient_check", "running")
    _, batch_run, boq_item, context = _load_batch_item_run_context(conn, item_run_id)
    kb_version_id = int(batch_run["kb_version_id"])
    conversion = _hydrate_conversion_combo_resources(
        conn, batch_run.get("conversion_check") or {}, kb_version_id
    ) or {}
    items = _coefficient_items_from_context(conn, context["items"], conversion, kb_version_id)
    started_at, started_perf = datetime.now(), perf_counter()
    result = None
    for event_type, payload in _batch_coefficient_business_events(items, boq_item):
        if event_type == "reasoning_token":
            continue
        if event_type == "coefficient_check":
            result = payload["coefficient_check"]
        else:
            _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, event_type, payload)
    if result is None:
        raise RuntimeError("系数换算未生成结果")
    timings = _load_batch_step_timings(conn, item_run_id)
    timing = _finish_step_timing(conn, None, timings, 7, "系数换算", started_at, started_perf)
    _update_batch_item_run(
        conn, item_run_id, coefficient_check=result, step_timings=timings, updated_at=datetime.now(), commit=False
    )
    _set_background_tool(
        conn, batch_id, execution_id, item_run_id, boq_item_id, "coefficient_check", "success",
        f"完成 {len(result.get('items', []))} 条定额检查",
    )
    _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, "step_timing", timing)


def _refresh_background_execution(conn, execution_id: int) -> None:
    with conn.cursor() as cur:
        cur.execute("SELECT batch_id,status,selected_count FROM pricing_task_batch_executions WHERE id=%s FOR UPDATE", (execution_id,))
        row = cur.fetchone()
        if not row:
            conn.rollback()
            return
        batch_id, execution_status, selected = int(row[0]), row[1], int(row[2])
        cur.execute(
            """SELECT COUNT(*) FILTER (WHERE status IN ('confirmed','completed','no_match')),
                      COUNT(*) FILTER (WHERE status='failed'),
                      COUNT(*) FILTER (WHERE status IN ('queued','running','retrying'))
               FROM pricing_task_batch_item_runs WHERE execution_id=%s""",
            (execution_id,),
        )
        completed, failed, pending = (int(v or 0) for v in cur.fetchone())
        final_status = None
        if pending == 0:
            final_status = "stopped" if execution_status == "stop_requested" else "completed"
        cur.execute(
            """UPDATE pricing_task_batch_executions SET completed_count=%s,failed_count=%s,
                      status=COALESCE(%s,status),finished_at=CASE WHEN %s IS NULL THEN finished_at ELSE NOW() END,
                      updated_at=NOW() WHERE id=%s""",
            (completed, failed, final_status, final_status, execution_id),
        )
        cur.execute(
            """UPDATE pricing_task_batches SET
                 selected_count=(SELECT COUNT(*) FROM pricing_task_batch_item_runs WHERE batch_id=%s),
                 completed_count=(SELECT COUNT(*) FROM pricing_task_batch_item_runs WHERE batch_id=%s AND status IN ('confirmed','completed','no_match')),
                 failed_count=(SELECT COUNT(*) FROM pricing_task_batch_item_runs WHERE batch_id=%s AND status='failed'),
                 status=COALESCE(%s,status),finished_at=CASE WHEN %s IS NULL THEN finished_at ELSE NOW() END,
                 updated_at=NOW() WHERE id=%s""",
            (batch_id, batch_id, batch_id, final_status, final_status, batch_id),
        )
    conn.commit()
    if final_status:
        _background_event(conn, batch_id, execution_id, None, None, "execution_finished",
                          {"status": final_status, "selected_count": selected, "completed_count": completed, "failed_count": failed})


def _defer_background_item_after_pool_busy(item_run_id: int, message: str) -> bool:
    """Return infrastructure-starved work to the queue without consuming an attempt."""
    from db.connection import get_connection

    delay = 0.5
    while True:
        sleep(delay)
        retry_conn = get_connection()
        try:
            with retry_conn.cursor() as cur:
                cur.execute(
                    """UPDATE pricing_task_batch_item_runs
                       SET status='retrying',attempt_count=GREATEST(attempt_count-1,0),
                           error_message=NULL,current_tool_status='retrying',
                           current_tool_output='数据库连接暂时繁忙，已重新排队',
                           next_attempt_at=NOW()+INTERVAL '2 seconds',lease_owner=NULL,
                           lease_expires_at=NULL,finished_at=NULL,updated_at=NOW()
                       WHERE id=%s
                       RETURNING batch_id,execution_id,boq_item_id""",
                    (item_run_id,),
                )
                row = cur.fetchone()
                if row:
                    # One event per busy episode is enough for observability;
                    # hundreds of duplicate events only add more write pressure.
                    cur.execute(
                        """INSERT INTO pricing_task_batch_events(
                               batch_id,execution_id,item_run_id,boq_item_id,event_type,payload
                           )
                           SELECT %s,%s,%s,%s,'item_deferred',%s
                           WHERE NOT EXISTS (
                               SELECT 1 FROM pricing_task_batch_events
                               WHERE item_run_id=%s AND event_type='item_deferred'
                                 AND created_at>NOW()-INTERVAL '1 minute'
                           )""",
                        (
                            int(row[0]), int(row[1]), item_run_id, int(row[2]),
                            Json({"reason": "database_pool_busy", "retry_after_seconds": 2}, dumps=_json_dumps),
                            item_run_id,
                        ),
                    )
            retry_conn.commit()
            if row:
                return True
            return False
        except DatabasePoolBusyError:
            retry_conn.rollback()
            delay = min(delay * 2, 15.0)
            continue
        except Exception as exc:
            retry_conn.rollback()
            print(f"[background-pricing] failed to defer item {item_run_id}: {exc}", file=sys.stderr, flush=True)
            return False
        finally:
            retry_conn.close()


def _run_background_item(item_run_id: int, worker_id: str) -> None:
    from db.connection import (
        background_database_connections,
        get_connection,
        release_connections_during_model_calls,
    )

    database_context = background_database_connections()
    database_context.__enter__()
    conn = get_connection()
    batch_id = execution_id = boq_item_id = None
    profile_context = None
    usage_context = None
    try:
      with release_connections_during_model_calls():
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """SELECT r.batch_id,r.execution_id,r.boq_item_id,b.quota_library_ids,b.manual_project_id,b.kb_version_id,
                          i.item_code,i.item_name,i.item_description,i.unit,i.quantity,i.project_id,r.attempt_count,
                          b.owner_user_id
                   FROM pricing_task_batch_item_runs r
                   JOIN pricing_task_batches b ON b.id=r.batch_id JOIN boq_items i ON i.id=r.boq_item_id
                   WHERE r.id=%s""",
                (item_run_id,),
            )
            row = cur.fetchone()
        if not row:
            return
        # Model profile resolution opens its own short database connection.
        # Release the item-context transaction first so 99 workers cannot each
        # hold one pool slot while waiting for a second slot.
        conn.commit()
        batch_id, execution_id, boq_item_id = int(row[0]), int(row[1]), int(row[2])
        if row[13] is None:
            raise RuntimeError("后台批量任务没有归属用户，无法解析模型配置")
        profile_context = use_default_profile(int(row[13]))
        profile_context.__enter__()
        usage_context = use_usage_context(business_type="background_pricing_batch", batch_id=batch_id, item_run_id=item_run_id, execution_id=execution_id)
        usage_context.__enter__()
        boq_item = {"id": boq_item_id, "item_code": row[6], "item_name": row[7], "item_description": row[8],
                    "unit": row[9], "quantity": float(row[10]) if row[10] is not None else None, "project_id": int(row[11])}
        _set_background_tool(conn, batch_id, execution_id, item_run_id, boq_item_id, "code_check", "running")
        fields: dict[str, Any] = {}
        chapter_seen = False
        for event_type, data in _stream_pricing_item(conn, boq_item, row[3] or [], row[4], None, None, int(row[5]), persist_run=False):
            if event_type in {"reasoning_token", "judgment"}:
                continue
            if event_type == "step_timing":
                timings = dict(fields.get("step_timings") or {})
                timings[str(data["step_no"])] = data
                fields["step_timings"] = timings
                _update_batch_item_run(conn, item_run_id, step_timings=timings, updated_at=datetime.now(), commit=False)
                _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, "step_timing", data)
                continue
            if event_type not in _BACKGROUND_TOOL_NEXT and event_type != "evaluation":
                continue
            if event_type == "chapter_rule_check" and chapter_seen:
                fields[event_type] = data
                _update_batch_item_run(conn, item_run_id, chapter_rule_check=data, updated_at=datetime.now())
                continue
            chapter_seen = chapter_seen or event_type == "chapter_rule_check"
            fields[event_type] = data
            _update_batch_item_run(conn, item_run_id, commit=False, **{event_type: data, "updated_at": datetime.now()})
            _set_background_tool(conn, batch_id, execution_id, item_run_id, boq_item_id, event_type, "success",
                                 _background_tool_output(event_type, data))
            next_tool = _BACKGROUND_TOOL_NEXT.get(event_type)
            if next_tool:
                _set_background_tool(conn, batch_id, execution_id, item_run_id, boq_item_id, next_tool, "running")
        matches = (fields.get("quota_match") or {}).get("matches", [])
        confirmed = _confirmed_results_from_matches(conn, matches, int(row[5]))
        _update_batch_item_run(conn, item_run_id, confirmed_results=confirmed, commit=False)
        if not matches:
            _update_batch_item_run(conn, item_run_id, status="no_match", current_tool_status="success",
                                   error_message=None,
                                   finished_at=datetime.now(), lease_owner=None, lease_expires_at=None, updated_at=datetime.now(), commit=False)
        else:
            _execute_background_conversion(conn, batch_id, execution_id, item_run_id, boq_item_id)
            _execute_background_coefficient(conn, batch_id, execution_id, item_run_id, boq_item_id)
            _update_batch_item_run(conn, item_run_id, status="confirmed", current_tool_status="success",
                                   error_message=None,
                                   finished_at=datetime.now(), lease_owner=None, lease_expires_at=None, updated_at=datetime.now(), commit=False)
        _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, "item_succeeded", {"status": "confirmed" if matches else "no_match"})
    except DatabasePoolBusyError as exc:
        conn.rollback()
        _defer_background_item_after_pool_busy(item_run_id, str(exc))
    except ModelRateLimitError as exc:
        conn.rollback()
        if item_run_id and batch_id and execution_id and boq_item_id:
            _update_batch_item_run(conn, item_run_id, status="failed", error_message=str(exc),
                                   current_tool_status="error", current_tool_output=str(exc),
                                   finished_at=datetime.now(), lease_owner=None, lease_expires_at=None, updated_at=datetime.now(), commit=False)
            _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, "item_failed",
                              {"error": str(exc), "failure_type": "model_rate_limited"})
    except Exception as exc:
        conn.rollback()
        if item_run_id:
            with conn.cursor() as cur:
                cur.execute("SELECT attempt_count FROM pricing_task_batch_item_runs WHERE id=%s", (item_run_id,))
                attempt_row = cur.fetchone()
            attempt = int(attempt_row[0] or 1) if attempt_row else 3
            if attempt < 3:
                delay = 5 if attempt == 1 else 20
                with conn.cursor() as cur:
                    cur.execute(
                        """UPDATE pricing_task_batch_item_runs SET status='retrying',error_message=%s,
                           current_tool_status='retrying',current_tool_output=%s,
                           next_attempt_at=NOW()+(%s * INTERVAL '1 second'),lease_owner=NULL,
                           lease_expires_at=NULL,updated_at=NOW() WHERE id=%s""",
                        (str(exc), f"第 {attempt} 次执行失败，{delay} 秒后重试", delay, item_run_id),
                    )
                if batch_id and execution_id and boq_item_id:
                    _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, "item_retrying",
                                      {"attempt": attempt, "delay_seconds": delay, "error": str(exc)})
                else:
                    conn.commit()
            else:
                _update_batch_item_run(conn, item_run_id, status="failed", error_message=str(exc),
                                       current_tool_status="error", current_tool_output=str(exc),
                                       finished_at=datetime.now(), lease_owner=None, lease_expires_at=None, updated_at=datetime.now(), commit=False)
                if batch_id and execution_id and boq_item_id:
                    _background_event(conn, batch_id, execution_id, item_run_id, boq_item_id, "item_failed", {"error": str(exc)})
                else:
                    conn.commit()
    finally:
        if usage_context:
            usage_context.__exit__(None, None, None)
        if profile_context:
            profile_context.__exit__(None, None, None)
        # The dispatcher refreshes all active execution counters once per tick.
        # Avoid 99 completed items contending for the same batch row lock.
        try:
            conn.close()
        finally:
            database_context.__exit__(None, None, None)


def _background_worker_limit() -> int:
    """Maximum in-flight background items for this API process.

    Model calls have an additional process-wide gate.  Keeping this limit at
    the same ceiling prevents a large executor from multiplying work across
    several batches before that gate is reached.
    """
    try:
        configured = int(os.getenv("PRICING_BACKGROUND_BATCH_CONCURRENCY", "99"))
    except ValueError:
        configured = 99
    return max(1, min(configured, 99))


def _claim_background_items(worker_id: str, global_limit: int) -> list[int]:
    from db.connection import get_connection

    conn = get_connection()
    claimed: list[int] = []
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                """UPDATE pricing_task_batch_item_runs SET status='retrying',lease_owner=NULL,lease_expires_at=NULL,
                          next_attempt_at=NOW(),updated_at=NOW()
                   WHERE status='running' AND execution_id IS NOT NULL AND lease_expires_at<NOW()"""
            )
        conn.commit()
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM pricing_task_batch_executions WHERE status IN ('queued','running','stop_requested') ORDER BY id")
            execution_ids = [int(row[0]) for row in cur.fetchall()]
            cur.execute(
                """SELECT COUNT(*) FROM pricing_task_batch_item_runs r
                   JOIN pricing_task_batch_executions e ON e.id=r.execution_id
                   WHERE e.status IN ('queued','running') AND r.status='running'"""
            )
            remaining_capacity = max(0, global_limit - int(cur.fetchone()[0] or 0))
        for execution_id in execution_ids:
            with conn.cursor() as cur:
                cur.execute("SELECT batch_id,status,concurrency_limit FROM pricing_task_batch_executions WHERE id=%s FOR UPDATE", (execution_id,))
                execution = cur.fetchone()
                if not execution:
                    conn.rollback()
                    continue
                batch_id, status, limit = int(execution[0]), execution[1], int(execution[2])
                cur.execute("SELECT COUNT(*) FROM pricing_task_batch_item_runs WHERE execution_id=%s AND status='running'", (execution_id,))
                running = int(cur.fetchone()[0])
                if status == "stop_requested":
                    if running == 0:
                        cur.execute("UPDATE pricing_task_batch_item_runs SET status='idle',execution_id=NULL,current_tool=NULL,current_tool_status=NULL,updated_at=NOW() WHERE execution_id=%s AND status IN ('queued','retrying')", (execution_id,))
                        cur.execute("UPDATE pricing_task_batch_executions SET status='stopped',finished_at=NOW(),updated_at=NOW() WHERE id=%s", (execution_id,))
                        cur.execute("UPDATE pricing_task_batches SET status='stopped',finished_at=NOW(),updated_at=NOW() WHERE id=%s", (batch_id,))
                    conn.commit()
                    continue
                available = min(max(0, limit - running), remaining_capacity)
                if available:
                    cur.execute(
                        """SELECT id FROM pricing_task_batch_item_runs
                           WHERE execution_id=%s AND status IN ('queued','retrying')
                             AND (next_attempt_at IS NULL OR next_attempt_at<=NOW())
                           ORDER BY id FOR UPDATE SKIP LOCKED LIMIT %s""",
                        (execution_id, available),
                    )
                    ids = [int(row[0]) for row in cur.fetchall()]
                    if ids:
                        cur.execute(
                            """UPDATE pricing_task_batch_item_runs SET status='running',attempt_count=attempt_count+1,
                               started_at=COALESCE(started_at,NOW()),lease_owner=%s,
                               lease_expires_at=NOW()+INTERVAL '1 minute',next_attempt_at=NULL,
                               error_message=NULL,updated_at=NOW()
                               WHERE id=ANY(%s)""",
                            (worker_id, ids),
                        )
                        claimed.extend(ids)
                        remaining_capacity -= len(ids)
                        cur.execute("UPDATE pricing_task_batch_executions SET status='running',started_at=COALESCE(started_at,NOW()),updated_at=NOW() WHERE id=%s", (execution_id,))
                        cur.execute("UPDATE pricing_task_batches SET status='running',updated_at=NOW() WHERE id=%s", (batch_id,))
                conn.commit()
            _refresh_background_execution(conn, execution_id)
        return claimed
    finally:
        conn.close()


def run_background_pricing_worker(stop_event: Any) -> None:
    from db.connection import background_database_connections

    worker_id = f"background:{os.getpid()}"
    worker_limit = _background_worker_limit()
    executor = ThreadPoolExecutor(max_workers=worker_limit, thread_name_prefix="background-pricing")
    futures: set[Any] = set()
    try:
        while not stop_event.is_set():
            futures = {future for future in futures if not future.done()}
            try:
                with background_database_connections():
                    claimed = _claim_background_items(worker_id, worker_limit)
                for item_run_id in claimed:
                    futures.add(executor.submit(_run_background_item, item_run_id, worker_id))
            except DatabasePoolBusyError:
                stop_event.wait(2.0)
                continue
            except Exception as exc:
                print(f"[background-pricing] dispatcher error: {exc}", file=sys.stderr, flush=True)
            stop_event.wait(0.8)
    finally:
        executor.shutdown(wait=False, cancel_futures=False)


def _manual_comparison_review_payload(row: Any) -> dict[str, Any]:
    return {
        "id": int(row[0]),
        "source_type": row[1],
        "source_run_id": int(row[2]),
        "manual_project_id": int(row[3]),
        "manual_item_id": int(row[4]),
        "before_manual_quotas": row[5] or [],
        "after_manual_quotas": row[6] or [],
        "before_evaluation": row[7] or {},
        "after_evaluation": row[8] or {},
        "retained_manual_quota_ids": row[9] or [],
        "accepted_ai_quotas": row[10] or [],
        "operator_name": row[11] or "系统",
        "created_at": row[12],
    }


def _load_manual_comparison_context(conn, source_type: str, source_run_id: int) -> dict[str, Any]:
    with conn.cursor() as cur:
        if source_type == "single":
            cur.execute(
                """
                SELECT r.id, r.task_id, r.quota_match, t.manual_project_id,
                       i.item_code, r.boq_item_id
                FROM pricing_task_runs r
                JOIN pricing_tasks t ON t.id = r.task_id
                JOIN boq_items i ON i.id = r.boq_item_id
                WHERE r.id=%s AND t.status <> 'deleted'
                FOR UPDATE OF r
                """,
                (source_run_id,),
            )
        elif source_type == "batch":
            cur.execute(
                """
                SELECT r.id, NULL, r.quota_match, b.manual_project_id,
                       i.item_code, r.boq_item_id
                FROM pricing_task_batch_item_runs r
                JOIN pricing_task_batches b ON b.id = r.batch_id
                JOIN boq_items i ON i.id = r.boq_item_id
                WHERE r.id=%s AND b.status <> 'deleted'
                FOR UPDATE OF r
                """,
                (source_run_id,),
            )
        else:
            raise ValueError(f"unsupported manual comparison source: {source_type}")
        row = cur.fetchone()
    if not row:
        label = "run" if source_type == "single" else "batch item run"
        raise HTTPException(status_code=404, detail=f"{label} not found")
    if not row[3]:
        label = "task" if source_type == "single" else "batch"
        raise HTTPException(status_code=400, detail=f"{label} has no manual comparison project")
    return {
        "source_type": source_type,
        "source_run_id": int(row[0]),
        "task_id": int(row[1]) if row[1] is not None else None,
        "quota_match": row[2] if isinstance(row[2], dict) else {},
        "manual_project_id": int(row[3]),
        "item_code": row[4],
        "boq_item_id": int(row[5]),
    }


def _list_manual_comparison_reviews(conn, source_type: str, source_run_id: int) -> list[dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT id, source_type, source_run_id, manual_project_id, manual_item_id,
                   before_manual_quotas, after_manual_quotas,
                   before_evaluation, after_evaluation,
                   retained_manual_quota_ids, accepted_ai_quotas,
                   operator_name, created_at
            FROM pricing_task_manual_comparison_reviews
            WHERE source_type=%s AND source_run_id=%s
            ORDER BY created_at DESC, id DESC
            """,
            (source_type, source_run_id),
        )
        return [_manual_comparison_review_payload(row) for row in cur.fetchall()]


def _apply_manual_comparison_review(
    conn,
    context: dict[str, Any],
    body: ManualComparisonUpdateRequest,
) -> dict[str, Any]:
    source_type = str(context["source_type"])
    source_run_id = int(context["source_run_id"])
    manual_project_id = int(context["manual_project_id"])
    item_code = context["item_code"]
    quota_match = context["quota_match"]
    matches = quota_match.get("matches", []) if isinstance(quota_match, dict) else []
    ai_by_key = {
        (int(item.get("dekid")), int(item.get("dezmid"))): item
        for item in matches
        if item.get("dekid") is not None and item.get("dezmid") is not None
    }

    normalized_code = str(item_code or "").strip().replace(" ", "")
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT id, quantity
            FROM manual_boq_items
            WHERE project_id=%s
              AND replace(trim(COALESCE(item_code, '')), ' ', '')=%s
            ORDER BY id
            FOR UPDATE
            """,
            (manual_project_id, normalized_code),
        )
        manual_items = cur.fetchall()
        if not manual_items:
            raise HTTPException(status_code=404, detail="manual comparison item not found")
        if len(manual_items) > 1:
            raise HTTPException(status_code=409, detail="duplicate item codes in manual comparison project")
        manual_item_id, manual_quantity = manual_items[0]

        current_manual = _manual_quotas(conn, manual_project_id, item_code)
        before_evaluation = _evaluate(matches, current_manual)
        current_by_id = {int(item["id"]): item for item in current_manual}
        retained_ids = {int(value) for value in body.retained_manual_quota_ids}
        unknown_manual_ids = retained_ids - set(current_by_id)
        if unknown_manual_ids:
            raise HTTPException(status_code=400, detail="manual quota does not belong to this item")

        requested_ai_keys = {(item.dekid, item.dezmid) for item in body.accepted_ai_quotas}
        unknown_ai_keys = requested_ai_keys - set(ai_by_key)
        if unknown_ai_keys:
            raise HTTPException(status_code=400, detail="AI quota does not belong to this run")

        ai_codes = [str(item.get("zmbh") or "").strip() for item in matches]
        locked_manual_ids = {
            int(item["id"])
            for item in current_manual
            if any(ai_code and ai_code in str(item.get("quota_code") or "") for ai_code in ai_codes)
        }
        if not locked_manual_ids.issubset(retained_ids):
            raise HTTPException(status_code=400, detail="consistent manual quotas must be retained")

        final_codes = {
            str(current_by_id[item_id].get("quota_code") or "").strip()
            for item_id in retained_ids
        }
        for key in requested_ai_keys:
            code = str(ai_by_key[key].get("zmbh") or "").strip()
            if code:
                final_codes.add(code)
        if not final_codes:
            raise HTTPException(status_code=400, detail="at least one final quota is required")

        if retained_ids:
            cur.execute(
                "DELETE FROM manual_boq_quotas WHERE boq_item_id=%s AND NOT (id = ANY(%s))",
                (manual_item_id, list(retained_ids)),
            )
        else:
            cur.execute("DELETE FROM manual_boq_quotas WHERE boq_item_id=%s", (manual_item_id,))

        retained_codes = [
            str(current_by_id[item_id].get("quota_code") or "").strip()
            for item_id in retained_ids
        ]
        for key in requested_ai_keys:
            item = ai_by_key[key]
            quota_code = str(item.get("zmbh") or "").strip()
            if not quota_code or any(quota_code in manual_code for manual_code in retained_codes):
                continue
            try:
                raw_qty_factor = item.get("qty_factor", 1)
                qty_factor = Decimal(str(1 if raw_qty_factor is None else raw_qty_factor))
            except (InvalidOperation, ValueError) as exc:
                raise HTTPException(status_code=400, detail="invalid AI quota factor") from exc
            if not qty_factor.is_finite() or qty_factor <= 0:
                raise HTTPException(status_code=400, detail="AI quota factor must be positive")
            quantity = Decimal(manual_quantity) * qty_factor if manual_quantity is not None else None
            cur.execute(
                """
                INSERT INTO manual_boq_quotas(
                    boq_item_id, quota_code, quota_name, quota_unit,
                    quantity, unit_price, total_price, qty_factor, quota_item_id
                )
                VALUES (%s,%s,%s,%s,%s,NULL,NULL,%s,NULL)
                """,
                (manual_item_id, quota_code, item.get("zmmc"), item.get("dw"), quantity, qty_factor),
            )
            retained_codes.append(quota_code)

        refreshed_manual = _manual_quotas(conn, manual_project_id, item_code)
        after_evaluation = _evaluate(matches, refreshed_manual)
        if source_type == "single":
            cur.execute(
                "UPDATE pricing_task_runs SET evaluation=%s, accuracy_report=NULL WHERE id=%s",
                (Json(after_evaluation, dumps=_json_dumps), source_run_id),
            )
            cur.execute(
                "UPDATE pricing_task_results SET evaluation=%s, updated_at=NOW() WHERE run_id=%s",
                (Json(after_evaluation, dumps=_json_dumps), source_run_id),
            )
            if context.get("task_id") is not None:
                cur.execute(
                    "UPDATE pricing_tasks SET accuracy_report=NULL, updated_at=NOW() WHERE id=%s",
                    (context["task_id"],),
                )
        else:
            cur.execute(
                "UPDATE pricing_task_batch_item_runs SET evaluation=%s WHERE id=%s",
                (Json(after_evaluation, dumps=_json_dumps), source_run_id),
            )

        accepted_ai_quotas = [
            {"dekid": int(key[0]), "dezmid": int(key[1])}
            for key in sorted(requested_ai_keys)
        ]
        cur.execute(
            """
            INSERT INTO pricing_task_manual_comparison_reviews(
                source_type, source_run_id, manual_project_id, manual_item_id,
                before_manual_quotas, after_manual_quotas,
                before_evaluation, after_evaluation,
                retained_manual_quota_ids, accepted_ai_quotas, operator_name
            )
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'系统')
            RETURNING id, source_type, source_run_id, manual_project_id, manual_item_id,
                      before_manual_quotas, after_manual_quotas,
                      before_evaluation, after_evaluation,
                      retained_manual_quota_ids, accepted_ai_quotas,
                      operator_name, created_at
            """,
            (
                source_type,
                source_run_id,
                manual_project_id,
                manual_item_id,
                Json(current_manual, dumps=_json_dumps),
                Json(refreshed_manual, dumps=_json_dumps),
                Json(before_evaluation, dumps=_json_dumps),
                Json(after_evaluation, dumps=_json_dumps),
                Json(sorted(retained_ids), dumps=_json_dumps),
                Json(accepted_ai_quotas, dumps=_json_dumps),
            ),
        )
        review = _manual_comparison_review_payload(cur.fetchone())

    return {
        "ok": True,
        "evaluation": after_evaluation,
        "manual_quotas": refreshed_manual,
        "review": review,
        "invalidated_task_id": context.get("task_id"),
    }


def _update_manual_comparison(source_type: str, source_run_id: int, body: ManualComparisonUpdateRequest):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        context = _load_manual_comparison_context(conn, source_type, source_run_id)
        result = _apply_manual_comparison_review(conn, context, body)
        conn.commit()
        return result
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _get_manual_comparison_history(source_type: str, source_run_id: int):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        _load_manual_comparison_context(conn, source_type, source_run_id)
        return _list_manual_comparison_reviews(conn, source_type, source_run_id)
    finally:
        conn.rollback()
        conn.close()


@router.put("/pricing-task-runs/{run_id}/manual-comparison")
def update_pricing_task_manual_comparison(run_id: int, body: ManualComparisonUpdateRequest):
    return _update_manual_comparison("single", run_id, body)


@router.get("/pricing-task-runs/{run_id}/manual-comparison-history")
def get_pricing_task_manual_comparison_history(run_id: int):
    return _get_manual_comparison_history("single", run_id)


@router.put("/pricing-task-batch-item-runs/{item_run_id}/manual-comparison")
def update_pricing_task_batch_manual_comparison(item_run_id: int, body: ManualComparisonUpdateRequest):
    return _update_manual_comparison("batch", item_run_id, body)


@router.get("/pricing-task-batch-item-runs/{item_run_id}/manual-comparison-history")
def get_pricing_task_batch_manual_comparison_history(item_run_id: int):
    return _get_manual_comparison_history("batch", item_run_id)


@router.post("/pricing-task-runs/{run_id}/confirm")
def confirm_pricing_task_run(run_id: int, body: ConfirmRunRequest):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute(
                "SELECT task_id, boq_item_id, boq_project_id, kb_version_id FROM pricing_task_runs WHERE id=%s",
                (run_id,),
            )
            run_row = cur.fetchone()
            if not run_row:
                raise HTTPException(status_code=404, detail="run not found")
            if body.results is not None:
                cur.execute("DELETE FROM pricing_task_results WHERE run_id=%s", (run_id,))
                task_id, boq_item_id, project_id, kb_version_id = run_row
                for m in body.results:
                    cur.execute(
                        """
                        SELECT zmbh, zmmc FROM tdek_tdezm
                        WHERE kb_version_id=pricing_kb_data_version(%s,'TDEK_TDEZM') AND dekid=%s AND id=%s
                        """,
                        (kb_version_id, m["dekid"], m["dezmid"]),
                    )
                    qrow = cur.fetchone()
                    if not qrow:
                        raise HTTPException(status_code=400, detail="quota item not found")
                    cur.execute(
                        """
                        INSERT INTO pricing_task_results(
                            task_id, run_id, boq_item_id, boq_project_id, dekid, dezmid,
                            subitem_code, subitem_name, qty_factor, confidence, match_reason, ai_reasoning, status
                        )
                        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'confirmed')
                        """,
                        (
                            task_id,
                            run_id,
                            boq_item_id,
                            project_id,
                            m["dekid"],
                            m["dezmid"],
                            qrow[0],
                            qrow[1],
                            m.get("qty_factor", 1),
                            m.get("confidence", "medium"),
                            m.get("match_reason", ""),
                            m.get("match_reason", ""),
                        ),
                    )
            else:
                cur.execute("UPDATE pricing_task_results SET status='confirmed', updated_at=NOW() WHERE run_id=%s", (run_id,))
            cur.execute("UPDATE pricing_task_runs SET status='confirmed', finished_at=COALESCE(finished_at, NOW()) WHERE id=%s", (run_id,))
        conn.commit()
        return {"ok": True}
    finally:
        conn.close()


@router.post("/pricing-task-runs/{run_id}/conversion-confirm")
def confirm_pricing_task_conversion(run_id: int, body: ConversionConfirmRequest):
    from db.connection import get_connection

    if not body.items:
        raise HTTPException(status_code=400, detail="conversion items required")

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute("SELECT status FROM pricing_task_runs WHERE id=%s", (run_id,))
            run_row = cur.fetchone()
            if not run_row:
                raise HTTPException(status_code=404, detail="run not found")
            if run_row[0] != "confirmed":
                raise HTTPException(status_code=400, detail="run must be confirmed before conversion writeback")

            for item in body.items:
                try:
                    dekid = int(item.get("dekid"))
                    dezmid = int(item.get("dezmid"))
                except Exception as exc:
                    raise HTTPException(status_code=400, detail="invalid dekid/dezmid") from exc
                try:
                    qty_factor = float(item.get("confirmed_qty_factor", item.get("qty_factor", 1)) or 1)
                except Exception as exc:
                    raise HTTPException(status_code=400, detail="invalid confirmed_qty_factor") from exc
                if qty_factor <= 0:
                    raise HTTPException(status_code=400, detail="confirmed_qty_factor must be positive")

                resources = item.get("conversion_resources", item.get("resources", []))
                resource_changes = item.get("conversion_resource_changes", item.get("resource_changes", []))
                if not isinstance(resources, list):
                    raise HTTPException(status_code=400, detail="conversion_resources must be array")
                if not isinstance(resource_changes, list):
                    raise HTTPException(status_code=400, detail="conversion_resource_changes must be array")

                cur.execute(
                    """
                    UPDATE pricing_task_results
                    SET qty_factor=%s,
                        conversion_confirmed=TRUE,
                        conversion_note=%s,
                        conversion_confirmed_at=NOW(),
                        conversion_resources=%s,
                        conversion_resource_changes=%s,
                        updated_at=NOW()
                    WHERE run_id=%s AND dekid=%s AND dezmid=%s AND status='confirmed'
                    """,
                    (
                        qty_factor,
                        item.get("conversion_note", "") or "",
                        Json(resources, dumps=_json_dumps),
                        Json(resource_changes, dumps=_json_dumps),
                        run_id,
                        dekid,
                        dezmid,
                    ),
                )
                if cur.rowcount == 0:
                    raise HTTPException(status_code=400, detail=f"confirmed quota not found: {dekid}/{dezmid}")
        conn.commit()
        return {"ok": True}
    finally:
        conn.close()


@router.post("/pricing-task-runs/{run_id}/reject")
def reject_pricing_task_run(run_id: int):
    from db.connection import get_connection

    conn = get_connection()
    try:
        _ensure_schema(conn)
        with conn.cursor() as cur:
            cur.execute("UPDATE pricing_task_results SET status='rejected', updated_at=NOW() WHERE run_id=%s", (run_id,))
            cur.execute("UPDATE pricing_task_runs SET status='rejected', finished_at=COALESCE(finished_at, NOW()) WHERE id=%s", (run_id,))
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="run not found")
        conn.commit()
        return {"ok": True}
    finally:
        conn.close()


@router.post("/pricing-task/match-item-stream")
def pricing_task_match_item_stream(req: dict[str, Any], user: CurrentUser = Depends(current_user)):
    """Compatibility endpoint used by older frontend code."""
    from db.connection import get_connection

    def generate():
        conn = get_connection()
        run_id = None
        profile_context = None
        try:
            _ensure_schema(conn)
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT id, item_code, item_name, item_description, unit, quantity, project_id FROM boq_items WHERE id = %s",
                    (req.get("boq_item_id"),),
                )
                row = cur.fetchone()
            if not row:
                yield _sse({"type": "error", "error": "Item not found"})
                return
            require_project_owner(conn, user, int(row[6]))
            profile_context = use_default_profile(user.id, business_type="pricing_match", run_id=run_id)
            profile_context.__enter__()
            boq_item = {
                "id": row[0],
                "item_code": row[1],
                "item_name": row[2],
                "item_description": row[3],
                "unit": row[4],
                "quantity": float(row[5]) if row[5] is not None else None,
                "project_id": row[6],
            }
            try:
                kb_version_id = resolve_version_id(conn, req.get("kb_version_id"))
            except ValueError as exc:
                yield _sse({"type": "error", "error": str(exc)})
                return
            run_id = _create_run(conn, None, boq_item, kb_version_id)
            yield _sse({
                "type": "run_started", "run_id": run_id,
                "kb_version_id": kb_version_id,
            })
            yield _sse({"type": "item_info", "item": boq_item})
            reasoning_text_parts: list[str] = []
            for event_type, data in _stream_pricing_item(
                conn, boq_item, req.get("quota_library_ids") or [],
                req.get("manual_project_id"), None, run_id, kb_version_id
            ):
                if event_type == "reasoning_token":
                    reasoning_text_parts.append(data)
                    yield _sse({"type": "reasoning_token", "token": data})
                else:
                    if event_type == "evaluation":
                        yield _sse({"type": "evaluation", "evaluation": data})
                    else:
                        yield _sse({"type": event_type, **data})
            _update_run(conn, run_id, status="completed", reasoning_text="".join(reasoning_text_parts), finished_at=datetime.now())
            yield _sse({"type": "done", "run_id": run_id})
        except Exception as exc:
            if run_id:
                _update_run(conn, run_id, status="failed", error_message=str(exc), finished_at=datetime.now())
            yield _sse({"type": "error", "error": str(exc)})
        finally:
            if profile_context:
                profile_context.__exit__(None, None, None)
            conn.close()

    return StreamingResponse(generate(), media_type="text/event-stream")
