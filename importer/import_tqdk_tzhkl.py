from __future__ import annotations

import argparse
import hashlib
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from psycopg2.extras import execute_values

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from db.connection import get_connection


DEFAULT_QDKID = 1020025
DEFAULT_FILE_GLOB = "*特征默认值*6.18.xlsx"


@dataclass(frozen=True)
class FeatureDefaultRow:
    qdkid: int
    qdzmid: int | None
    zmbh: str
    zmmc: str | None
    chapter_name: str
    feature_name: str
    feature_value: str
    default_value: str
    source_file_sha256: str
    source_sheet: str
    source_rowid: int


def create_table(conn) -> None:
    with conn.cursor() as cur:
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS tqdk_tzhkl (
                id BIGSERIAL PRIMARY KEY,
                qdkid BIGINT NOT NULL,
                qdzmid BIGINT NULL,
                zmbh VARCHAR(64) NOT NULL,
                zmmc TEXT NULL,
                chapter_name TEXT NOT NULL,
                feature_name TEXT NOT NULL,
                feature_value TEXT NOT NULL,
                default_value TEXT NOT NULL,
                source_file_sha256 VARCHAR(64) NOT NULL,
                source_sheet VARCHAR(128) NOT NULL,
                source_rowid INTEGER NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """
        )
        cur.execute(
            """
            CREATE UNIQUE INDEX IF NOT EXISTS uq_tqdk_tzhkl_source_seq
            ON tqdk_tzhkl(source_file_sha256, source_sheet, source_rowid)
            """
        )
        cur.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_tqdk_tzhkl_code
            ON tqdk_tzhkl(qdkid, zmbh)
            """
        )
        cur.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_tqdk_tzhkl_qdzm
            ON tqdk_tzhkl(qdkid, qdzmid)
            """
        )
        cur.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_tqdk_tzhkl_feature
            ON tqdk_tzhkl USING gin (
                to_tsvector('simple', COALESCE(feature_name, '') || ' ' || COALESCE(default_value, ''))
            )
            """
        )
    conn.commit()


def find_source_file(path_arg: str | None) -> Path:
    if path_arg:
        path = Path(path_arg)
        if not path.is_absolute():
            path = ROOT_DIR / path
        if not path.exists():
            raise FileNotFoundError(path)
        return path

    rule_dir = ROOT_DIR / "mydoc" / "rule"
    matches = sorted(rule_dir.glob(DEFAULT_FILE_GLOB))
    if not matches:
        raise FileNotFoundError(rule_dir / DEFAULT_FILE_GLOB)
    if len(matches) > 1:
        names = ", ".join(str(p) for p in matches)
        raise RuntimeError(f"匹配到多个特征默认值文件，请显式指定: {names}")
    return matches[0]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def normalize_code(value) -> str:
    if isinstance(value, int):
        return str(value).zfill(9)
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).zfill(9)
    text = str(value).strip()
    if text.isdigit():
        return text.zfill(9)
    return text


def normalize_text(value) -> str:
    return str(value).strip()


def load_workbook_rows(path: Path, file_hash: str, qdkid: int) -> tuple[list[dict], list[str]]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook["Sheet1"]
    expected_header = ["9位章节编码", "章节名称", "特征名称", "特征值", "默认值"]
    header = [normalize_text(cell.value).replace(" ", "") for cell in sheet[1][:5]]
    if header != expected_header:
        raise RuntimeError(f"Excel 表头不符合预期: {header}")

    item_rows: list[dict] = []
    codes: list[str] = []
    for rowid, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = row[:5]
        if all(value is None or normalize_text(value) == "" for value in values):
            continue
        if any(value is None or normalize_text(value) == "" for value in values):
            raise RuntimeError(f"第 {rowid} 行存在空字段: {values}")

        zmbh = normalize_code(values[0])
        item_rows.append(
            {
                "qdkid": qdkid,
                "qdzmid": None,
                "zmbh": zmbh,
                "zmmc": None,
                "chapter_name": normalize_text(values[1]),
                "feature_name": normalize_text(values[2]),
                "feature_value": normalize_text(values[3]),
                "default_value": normalize_text(values[4]),
                "source_file_sha256": file_hash,
                "source_sheet": sheet.title,
                "source_rowid": rowid,
            }
        )
        codes.append(zmbh)

    return item_rows, sorted(set(codes))


def get_qdzm_map(conn, qdkid: int, codes: list[str]) -> dict[str, tuple[int, str]]:
    if not codes:
        return {}
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT DISTINCT ON (zmbh) zmbh, id, zmmc
            FROM tqdk_tqdzm
            WHERE qdkid = %s AND zmbh = ANY(%s)
            ORDER BY zmbh, id
            """,
            (qdkid, codes),
        )
        return {zmbh: (item_id, item_name) for zmbh, item_id, item_name in cur.fetchall()}


def build_feature_rows(
    item_rows: list[dict], qdzm_map: dict[str, tuple[int, str]]
) -> list[FeatureDefaultRow]:
    rows: list[FeatureDefaultRow] = []
    for item in item_rows:
        qdzm = qdzm_map.get(item["zmbh"])
        rows.append(
            FeatureDefaultRow(
                qdkid=item["qdkid"],
                qdzmid=qdzm[0] if qdzm else None,
                zmbh=item["zmbh"],
                zmmc=qdzm[1] if qdzm else None,
                chapter_name=item["chapter_name"],
                feature_name=item["feature_name"],
                feature_value=item["feature_value"],
                default_value=item["default_value"],
                source_file_sha256=item["source_file_sha256"],
                source_sheet=item["source_sheet"],
                source_rowid=item["source_rowid"],
            )
        )
    return rows


def import_rows(conn, file_hash: str, rows: list[FeatureDefaultRow]) -> int:
    with conn.cursor() as cur:
        cur.execute("DELETE FROM tqdk_tzhkl WHERE source_file_sha256 = %s", (file_hash,))
        if rows:
            execute_values(
                cur,
                """
                INSERT INTO tqdk_tzhkl (
                    qdkid, qdzmid, zmbh, zmmc, chapter_name, feature_name,
                    feature_value, default_value, source_file_sha256, source_sheet, source_rowid
                )
                VALUES %s
                """,
                [
                    (
                        row.qdkid,
                        row.qdzmid,
                        row.zmbh,
                        row.zmmc,
                        row.chapter_name,
                        row.feature_name,
                        row.feature_value,
                        row.default_value,
                        row.source_file_sha256,
                        row.source_sheet,
                        row.source_rowid,
                    )
                    for row in rows
                ],
            )
    conn.commit()
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="导入清单特征综合考虑默认值到 tqdk_tzhkl")
    parser.add_argument("--file", help="Excel 文件路径，默认匹配 mydoc/rule/*特征默认值*6.18.xlsx")
    parser.add_argument("--qdkid", type=int, default=DEFAULT_QDKID)
    args = parser.parse_args()

    source_file = find_source_file(args.file)
    file_hash = sha256_file(source_file)
    item_rows, codes = load_workbook_rows(source_file, file_hash, args.qdkid)

    conn = get_connection()
    try:
        create_table(conn)
        qdzm_map = get_qdzm_map(conn, args.qdkid, codes)
        feature_rows = build_feature_rows(item_rows, qdzm_map)
        inserted = import_rows(conn, file_hash, feature_rows)
    finally:
        conn.close()

    unmatched = [code for code in codes if code not in qdzm_map]
    print(f"source_file={source_file}")
    print(f"source_file_sha256={file_hash}")
    print(f"qdkid={args.qdkid}")
    print(f"feature_rows={len(item_rows)}")
    print(f"inserted_rows={inserted}")
    print(f"unique_codes={len(codes)}")
    print(f"matched_codes={len(qdzm_map)}")
    print(f"unmatched_codes={len(unmatched)}")
    if unmatched:
        print("unmatched_code_list=" + ",".join(unmatched))


if __name__ == "__main__":
    main()
