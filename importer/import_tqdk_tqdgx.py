from __future__ import annotations

import argparse
import hashlib
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from psycopg2.extras import execute_values

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from db.connection import get_connection


DEFAULT_QDKID = 1020025
DEFAULT_FILE_GLOB = "施工工序-*6.12.xlsx"


@dataclass(frozen=True)
class ProcessRow:
    qdkid: int
    qdzmid: int | None
    zmbh: str
    zmmc: str
    appendix_code: str | None
    appendix_name: str | None
    procedure_text: str
    source_file_sha256: str
    source_sheet: str
    source_rowid: int


def create_table(conn) -> None:
    with conn.cursor() as cur:
        cur.execute("DROP TABLE IF EXISTS tqdk_tqdgx")
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS tqdk_tqdgx (
                id BIGSERIAL PRIMARY KEY,
                qdkid BIGINT NOT NULL,
                qdzmid BIGINT NULL,
                zmbh VARCHAR(64) NOT NULL,
                zmmc TEXT NOT NULL,
                appendix_code VARCHAR(16) NULL,
                appendix_name TEXT NULL,
                procedure_text TEXT NOT NULL,
                source_file_sha256 VARCHAR(64) NOT NULL,
                source_sheet VARCHAR(128) NOT NULL,
                source_rowid INTEGER NOT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
            """
        )
        cur.execute(
            """
            CREATE UNIQUE INDEX IF NOT EXISTS uq_tqdk_tqdgx_source_seq
            ON tqdk_tqdgx(source_file_sha256, source_sheet, source_rowid)
            """
        )
        cur.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_tqdk_tqdgx_code
            ON tqdk_tqdgx(qdkid, zmbh)
            """
        )
        cur.execute(
            """
            CREATE INDEX IF NOT EXISTS idx_tqdk_tqdgx_qdzm
            ON tqdk_tqdgx(qdkid, qdzmid)
            """
        )
    conn.commit()


def find_source_file(path_arg: str | None) -> Path:
    if path_arg:
        path = Path(path_arg)
        if not path.is_absolute():
            path = ROOT_DIR / path
        if not path.exists():
            raise FileNotFoundError(path)
        return path

    rule_dir = ROOT_DIR / "mydoc" / "rule"
    matches = sorted(rule_dir.glob(DEFAULT_FILE_GLOB))
    if not matches:
        raise FileNotFoundError(rule_dir / DEFAULT_FILE_GLOB)
    if len(matches) > 1:
        names = ", ".join(str(p) for p in matches)
        raise RuntimeError(f"匹配到多个施工工序文件，请显式指定: {names}")
    return matches[0]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def normalize_code(value) -> str:
    if isinstance(value, int):
        return str(value).zfill(9)
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).zfill(9)
    text = str(value).strip()
    if text.isdigit():
        return text.zfill(9)
    return text


def parse_appendix(section_text: str | None) -> tuple[str | None, str | None]:
    if not section_text:
        return None, None
    match = re.match(r"^附录([A-Z])\s*(.+)$", section_text.strip())
    if match:
        return match.group(1), match.group(2).strip()
    return None, section_text.strip()


def load_workbook_rows(path: Path, file_hash: str, qdkid: int) -> tuple[list[dict], list[str]]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook["Sheet1"]
    header = tuple(cell.value for cell in sheet[2])

    item_rows: list[dict] = []
    codes: list[str] = []
    current_section: str | None = None

    for rowid, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        code_value, name_value, process_value = row[:3]

        if row == header:
            continue
        if isinstance(code_value, str) and name_value is None and process_value is None:
            current_section = code_value.strip()
            continue
        if code_value is None or name_value is None or process_value is None:
            continue

        zmbh = normalize_code(code_value)
        procedure_text = str(process_value).strip()
        appendix_code, appendix_name = parse_appendix(current_section)
        item_rows.append(
            {
                "qdkid": qdkid,
                "qdzmid": None,
                "zmbh": zmbh,
                "zmmc": str(name_value).strip(),
                "appendix_code": appendix_code,
                "appendix_name": appendix_name,
                "procedure_text": procedure_text,
                "source_file_sha256": file_hash,
                "source_sheet": sheet.title,
                "source_rowid": rowid,
            }
        )
        codes.append(zmbh)

    return item_rows, sorted(set(codes))


def get_qdzm_map(conn, qdkid: int, codes: list[str]) -> dict[str, int]:
    if not codes:
        return {}
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT zmbh, id
            FROM tqdk_tqdzm
            WHERE qdkid = %s AND zmbh = ANY(%s)
            """,
            (qdkid, codes),
        )
        return {zmbh: item_id for zmbh, item_id in cur.fetchall()}


def build_process_rows(
    item_rows: list[dict], qdzm_map: dict[str, int]
) -> list[ProcessRow]:
    rows: list[ProcessRow] = []
    for item in item_rows:
        qdzmid = qdzm_map.get(item["zmbh"])
        rows.append(
            ProcessRow(
                qdkid=item["qdkid"],
                qdzmid=qdzmid,
                zmbh=item["zmbh"],
                zmmc=item["zmmc"],
                appendix_code=item["appendix_code"],
                appendix_name=item["appendix_name"],
                procedure_text=item["procedure_text"],
                source_file_sha256=item["source_file_sha256"],
                source_sheet=item["source_sheet"],
                source_rowid=item["source_rowid"],
            )
        )
    return rows


def import_rows(conn, file_hash: str, rows: list[ProcessRow]) -> int:
    with conn.cursor() as cur:
        cur.execute(
            "DELETE FROM tqdk_tqdgx WHERE source_file_sha256 = %s",
            (file_hash,),
        )
        if rows:
            execute_values(
                cur,
                """
                INSERT INTO tqdk_tqdgx (
                    qdkid, qdzmid, zmbh, zmmc, appendix_code, appendix_name,
                    procedure_text, source_file_sha256, source_sheet, source_rowid
                )
                VALUES %s
                """,
                [
                    (
                        row.qdkid,
                        row.qdzmid,
                        row.zmbh,
                        row.zmmc,
                        row.appendix_code,
                        row.appendix_name,
                        row.procedure_text,
                        row.source_file_sha256,
                        row.source_sheet,
                        row.source_rowid,
                    )
                    for row in rows
                ],
            )
    conn.commit()
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="导入施工工序到 tqdk_tqdgx")
    parser.add_argument("--file", help="Excel 文件路径，默认匹配 mydoc/rule/施工工序-*6.12.xlsx")
    parser.add_argument("--qdkid", type=int, default=DEFAULT_QDKID)
    args = parser.parse_args()

    source_file = find_source_file(args.file)
    file_hash = sha256_file(source_file)
    item_rows, codes = load_workbook_rows(source_file, file_hash, args.qdkid)

    conn = get_connection()
    try:
        create_table(conn)
        qdzm_map = get_qdzm_map(conn, args.qdkid, codes)
        process_rows = build_process_rows(item_rows, qdzm_map)
        inserted = import_rows(conn, file_hash, process_rows)
    finally:
        conn.close()

    unmatched = [code for code in codes if code not in qdzm_map]
    print(f"source_file={source_file}")
    print(f"source_file_sha256={file_hash}")
    print(f"qdkid={args.qdkid}")
    print(f"item_rows={len(item_rows)}")
    print(f"process_chain_rows={inserted}")
    print(f"matched_codes={len(qdzm_map)}")
    print(f"unmatched_codes={len(unmatched)}")
    if unmatched:
        print("unmatched_code_list=" + ",".join(unmatched))


if __name__ == "__main__":
    main()
